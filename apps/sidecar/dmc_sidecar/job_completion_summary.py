from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]


SUCCESS_NOTES = {"filled", "dry_run", "default_missing_filled", "default_missing_dry_run", "submitted"}
SUMMARY_COLUMNS = (
    ("row_index", "Row"),
    ("record_id", "Record ID"),
    ("student_no", "Student No"),
    ("citizen_id", "Citizen ID"),
    ("full_name", "Name"),
    ("classroom", "Room"),
    ("status", "Status"),
    ("note", "Note"),
    ("message", "Message"),
    ("applied", "Applied"),
)


def build_completion_summary(results: list[dict[str, Any]]) -> dict[str, Any]:
    success_items: list[dict[str, Any]] = []
    failure_items: list[dict[str, Any]] = []
    for result in results:
        item = _completion_item(result)
        if _is_success_result(result):
            success_items.append(item)
        else:
            failure_items.append(item)

    return {
        "total": len(success_items) + len(failure_items),
        "succeeded": len(success_items),
        "failed": len(failure_items),
        "success_items": success_items,
        "failure_items": failure_items,
    }


def write_completion_summary_workbook(
    path: Path,
    *,
    job_id: str,
    module: str,
    source_file: str,
    results: list[dict[str, Any]],
) -> Path:
    summary = build_completion_summary(results)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _write_summary_sheet(
        summary_sheet,
        job_id=job_id,
        module=module,
        source_file=source_file,
        summary=summary,
    )
    _write_items_sheet(workbook.create_sheet("Success"), summary["success_items"])
    _write_items_sheet(workbook.create_sheet("Failed"), summary["failure_items"])
    _write_items_sheet(
        workbook.create_sheet("All Records"),
        [*summary["success_items"], *summary["failure_items"]],
    )
    workbook.save(path)
    return path


def _write_summary_sheet(
    sheet: Any,
    *,
    job_id: str,
    module: str,
    source_file: str,
    summary: dict[str, Any],
) -> None:
    sheet.append(["Metric", "Value"])
    sheet.append(["Job ID", job_id])
    sheet.append(["Module", module])
    sheet.append(["Source file", source_file])
    sheet.append(["Total", summary["total"]])
    sheet.append(["Succeeded", summary["succeeded"]])
    sheet.append(["Failed", summary["failed"]])
    _style_header_row(sheet)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 72


def _write_items_sheet(sheet: Any, items: list[dict[str, Any]]) -> None:
    sheet.append([label for _key, label in SUMMARY_COLUMNS])
    for item in items:
        sheet.append([item.get(key) for key, _label in SUMMARY_COLUMNS])
    _style_header_row(sheet)
    sheet.freeze_panes = "A2"
    for column_index, (_key, label) in enumerate(SUMMARY_COLUMNS, start=1):
        header_cell = sheet.cell(row=1, column=column_index)
        sheet.column_dimensions[header_cell.column_letter].width = max(12, min(len(label) + 8, 32))


def _style_header_row(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="E2E8F0")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill


def _completion_item(result: dict[str, Any]) -> dict[str, Any]:
    success = _is_success_result(result)
    status = _first_text(result, "status") or ("success" if success else "review")
    return {
        "row_index": _first_int(result, "row_index", "record_number", "matched_order", "portal_seq_no", "portal_row_index"),
        "record_id": _first_text(result, "record_id"),
        "student_no": _first_text(result, "student_no", "matched_student_no", "portal_student_no"),
        "citizen_id": _first_text(result, "citizen_id"),
        "full_name": _first_text(result, "full_name", "matched_name", "portal_name"),
        "classroom": _first_text(result, "classroom", "matched_room", "portal_room"),
        "status": status,
        "note": _first_text(result, "note"),
        "message": _first_text(result, "message"),
        "applied": result.get("applied") is True,
    }


def _is_success_result(result: dict[str, Any]) -> bool:
    status = (_first_text(result, "status") or "").lower()
    note = _first_text(result, "note") or ""
    return status == "success" or note in SUCCESS_NOTES or result.get("applied") is True


def _first_text(result: dict[str, Any], *keys: str) -> str | None:
    for key in keys:
        value = result.get(key)
        if value is None:
            continue
        text = str(value).strip()
        if text:
            return text
    return None


def _first_int(result: dict[str, Any], *keys: str) -> int | None:
    for key in keys:
        value = result.get(key)
        if isinstance(value, bool) or value is None:
            continue
        try:
            return int(value)
        except (TypeError, ValueError):
            continue
    return None
