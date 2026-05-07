from __future__ import annotations

import re
from collections.abc import Sequence
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from .config import bundled_resources_root, repo_root, reports_dir
from .errors import DomainError
from .schemas import ExportStudentBasicInfoFormResponse, StudentBasicInfoClassSummary


DATA_START_ROW = 5
FORM_MAX_COLUMN = 9
SOURCE_YEAR_TERM_PATTERN = re.compile(r"(?P<year>[0-9]{4})[-_ ](?P<term>[0-9]+)")
INVALID_SHEET_TITLE_CHARS = re.compile(r"[:\\/?*\[\]]")
EMPTY_MARKERS = {"", "-"}
DEFAULT_TERM = "2"
DOCUMENT_FONT_NAME = "TH Sarabun New"
DOCUMENT_FONT_SIZE = 16
DOCUMENT_TITLE = "แบบฟอร์มบันทึกข้อมูลพื้นฐานนักเรียน(งานพยาบาล)"


@dataclass(frozen=True)
class ColumnLookup:
    school_name: int
    citizen_id: int
    level: int
    room: int
    prefix: int
    first_name: int
    last_name: int
    birth_date: int
    weight: int
    height: int
    house_no: int
    moo: int
    road: int
    subdistrict: int
    district: int
    province: int
    father_first_name: int
    father_last_name: int
    mother_first_name: int
    mother_last_name: int


@dataclass(frozen=True)
class StudentBasicInfoRow:
    school_name: str
    citizen_id: str
    level: str
    room: str
    full_name: str
    birth_date: str
    address: str
    father_full_name: str
    mother_full_name: str
    weight: str | int | float
    height: str | int | float


def export_student_basic_info_form(
    *,
    excel_path: Path,
    template_path: Path | None = None,
) -> ExportStudentBasicInfoFormResponse:
    source_path = _validated_excel_path(excel_path)
    resolved_template_path = _resolved_template_path(template_path)
    students, rows_total = _read_students(source_path)
    if not students:
        raise DomainError("STUDENT_BASIC_INFO_NO_ROWS", "No student rows were found in the DMC Excel file.")

    groups = _group_students(students)
    output_path = _build_output_path(source_path)
    workbook = load_workbook(resolved_template_path)
    template_sheet = workbook.active
    for sheet in list(workbook.worksheets):
        if sheet is not template_sheet:
            workbook.remove(sheet)

    school_year, _source_term = _year_and_term_from_file_name(source_path.name)
    term = DEFAULT_TERM
    summaries: list[StudentBasicInfoClassSummary] = []
    existing_sheet_names: set[str] = set()
    for index, (level, room, class_students) in enumerate(_sorted_groups(groups)):
        sheet = template_sheet if index == 0 else workbook.copy_worksheet(template_sheet)
        sheet_name = _safe_sheet_title(level=level, room=room, existing=existing_sheet_names)
        sheet.title = sheet_name
        existing_sheet_names.add(sheet_name)
        _fill_class_sheet(
            sheet,
            students=class_students,
            school_name=class_students[0].school_name,
            school_year=school_year,
            term=term,
            level=level,
            room=room,
        )
        _apply_document_style(sheet)
        summaries.append(
            StudentBasicInfoClassSummary(
                level=level,
                room=room,
                sheet_name=sheet_name,
                students=len(class_students),
            )
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return ExportStudentBasicInfoFormResponse(
        module="studentBasicInfo",
        source_path=str(source_path),
        output_path=str(output_path),
        school_name=students[0].school_name or None,
        school_year=school_year,
        term=term,
        rows_total=rows_total,
        students_exported=len(students),
        classes_exported=len(summaries),
        classes=summaries,
    )


def _validated_excel_path(path: Path) -> Path:
    if not path.exists():
        raise DomainError("STUDENT_BASIC_INFO_INPUT_NOT_FOUND", "DMC Excel file was not found.")
    if not path.is_file():
        raise DomainError("STUDENT_BASIC_INFO_INPUT_NOT_FILE", "Selected path is not a file.")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise DomainError("STUDENT_BASIC_INFO_UNSUPPORTED_TYPE", "Only .xlsx and .xlsm files are supported.")
    return path


def _resolved_template_path(template_path: Path | None) -> Path:
    if template_path is not None:
        candidate = template_path
    else:
        bundled_root = bundled_resources_root()
        candidate = (
            bundled_root / "templates" / "student_basic_info_form.xlsx"
            if bundled_root is not None
            else repo_root() / "student_basic_info_form.xlsx"
        )
    if not candidate.exists():
        raise DomainError("STUDENT_BASIC_INFO_TEMPLATE_NOT_FOUND", "Student basic info template was not found.")
    if not candidate.is_file():
        raise DomainError("STUDENT_BASIC_INFO_TEMPLATE_NOT_FILE", "Student basic info template path is not a file.")
    return candidate


def _read_students(path: Path) -> tuple[list[StudentBasicInfoRow], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        max_column = max(int(sheet.max_column or 1), 1)
        header_row, headers = _find_header_row(sheet, max_column=max_column)
        columns = _build_column_lookup(headers)
        max_column = max(max_column, max(int(value) for value in vars(columns).values()))
        rows_total = 0
        students: list[StudentBasicInfoRow] = []
        for row_values in sheet.iter_rows(min_row=header_row + 1, max_col=max_column, values_only=True):
            rows_total += 1
            if _row_is_blank(row_values):
                continue
            students.append(_student_from_row(row_values, columns))
        return students, rows_total
    finally:
        workbook.close()


def _find_header_row(sheet: Any, *, max_column: int) -> tuple[int, dict[str, list[int]]]:
    max_row = min(int(sheet.max_row or 0), 20)
    for row_index, row_values in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_column, values_only=True),
        start=1,
    ):
        headers: dict[str, list[int]] = {}
        for column_index, value in enumerate(row_values, start=1):
            header = _clean_header(value)
            if header:
                headers.setdefault(header, []).append(column_index)
        if {"ชื่อโรงเรียน", "ชั้น", "ห้อง", "ชื่อ", "นามสกุล"}.issubset(headers):
            return row_index, headers
    raise DomainError(
        "STUDENT_BASIC_INFO_HEADERS_NOT_FOUND",
        "Could not find the DMC student header row in the Excel file.",
    )


def _build_column_lookup(headers: dict[str, list[int]]) -> ColumnLookup:
    return ColumnLookup(
        school_name=_column(headers, [("ชื่อโรงเรียน", 0)]),
        citizen_id=_column(
            headers,
            [
                ("เลขบัตรประชาชน", 0),
                ("เลขประจำตัวประชาชน", 0),
                ("เลขประจำตัวนักเรียน", 0),
            ],
        ),
        level=_column(headers, [("ชั้น", 0)]),
        room=_column(headers, [("ห้อง", 0)]),
        prefix=_column(headers, [("คำนำหน้าชื่อ", 0)]),
        first_name=_column(headers, [("ชื่อ", 0)]),
        last_name=_column(headers, [("นามสกุล", 0)]),
        birth_date=_column(headers, [("วันเกิด", 0)]),
        weight=_column(headers, [("น้ำหนัก", 0)]),
        height=_column(headers, [("ส่วนสูง", 0)]),
        house_no=_column(headers, [("บ้านเลขที่", 0)]),
        moo=_column(headers, [("หมู่", 0)]),
        road=_column(headers, [("ถนน/ซอย", 0)]),
        subdistrict=_column(headers, [("ตำบล", 0)]),
        district=_column(headers, [("อำเภอ", 0)]),
        province=_column(headers, [("จังหวัด", 0)]),
        father_first_name=_column(headers, [("ชื่อบิดา", 0)]),
        father_last_name=_column(headers, [("นามสกุลบิดา", 0)]),
        mother_first_name=_column(headers, [("ชื่อมารดา", 0)]),
        mother_last_name=_column(headers, [("นามสกุลมารดา", 0)]),
    )


def _column(headers: dict[str, list[int]], candidates: Sequence[tuple[str, int]]) -> int:
    for name, occurrence in candidates:
        columns = headers.get(name, [])
        if occurrence < len(columns):
            return columns[occurrence]
    raise DomainError(
        "STUDENT_BASIC_INFO_MISSING_COLUMN",
        f"Missing required DMC column: {candidates[0][0]}",
    )


def _student_from_row(row_values: Sequence[Any], columns: ColumnLookup) -> StudentBasicInfoRow:
    return StudentBasicInfoRow(
        school_name=_text(_cell(row_values, columns.school_name)),
        citizen_id=_text(_cell(row_values, columns.citizen_id)),
        level=_text(_cell(row_values, columns.level)),
        room=_text(_cell(row_values, columns.room)),
        full_name=_full_name(
            _text(_cell(row_values, columns.prefix)),
            _text(_cell(row_values, columns.first_name)),
            _text(_cell(row_values, columns.last_name)),
        ),
        birth_date=_text(_cell(row_values, columns.birth_date)),
        address=_address(
            house_no=_text(_cell(row_values, columns.house_no)),
            moo=_text(_cell(row_values, columns.moo)),
            road=_text(_cell(row_values, columns.road)),
            subdistrict=_text(_cell(row_values, columns.subdistrict)),
            district=_text(_cell(row_values, columns.district)),
            province=_text(_cell(row_values, columns.province)),
        ),
        father_full_name=_full_name(
            "",
            _text(_cell(row_values, columns.father_first_name)),
            _text(_cell(row_values, columns.father_last_name)),
        ),
        mother_full_name=_full_name(
            "",
            _text(_cell(row_values, columns.mother_first_name)),
            _text(_cell(row_values, columns.mother_last_name)),
        ),
        weight=_number_or_text(_cell(row_values, columns.weight)),
        height=_number_or_text(_cell(row_values, columns.height)),
    )


def _cell(row_values: Sequence[Any], column_index: int) -> Any:
    value_index = column_index - 1
    if value_index < 0 or value_index >= len(row_values):
        return None
    return row_values[value_index]


def _row_is_blank(row_values: Sequence[Any]) -> bool:
    return all(_text(value) == "" for value in row_values)


def _group_students(
    students: list[StudentBasicInfoRow],
) -> dict[tuple[str, str], list[StudentBasicInfoRow]]:
    groups: dict[tuple[str, str], list[StudentBasicInfoRow]] = {}
    for student in students:
        groups.setdefault((student.level, student.room), []).append(student)
    return groups


def _sorted_groups(
    groups: dict[tuple[str, str], list[StudentBasicInfoRow]],
) -> list[tuple[str, str, list[StudentBasicInfoRow]]]:
    return [
        (level, room, groups[(level, room)])
        for level, room in sorted(groups.keys(), key=lambda item: _class_sort_key(item[0], item[1]))
    ]


def _class_sort_key(level: str, room: str) -> tuple[int, int, int, str, str]:
    match = re.match(r"^\s*([^\d]+)\s*(\d+)", level)
    prefix = match.group(1).strip() if match else level
    level_number = int(match.group(2)) if match else 0
    prefix_order = {
        "อ.": 1,
        "อนุบาล": 1,
        "ป.": 2,
        "ประถม": 2,
        "ม.": 3,
        "มัธยม": 3,
    }.get(prefix, 99)
    return (prefix_order, level_number, _safe_int(room), level, room)


def _fill_class_sheet(
    sheet: Any,
    *,
    students: list[StudentBasicInfoRow],
    school_name: str,
    school_year: str | None,
    term: str | None,
    level: str,
    room: str,
) -> None:
    capacity = max(int(sheet.max_row) - DATA_START_ROW + 1, 0)
    _ensure_student_row_capacity(sheet, capacity=capacity, needed=len(students))

    if school_name:
        sheet["A1"] = DOCUMENT_TITLE
        sheet["A2"] = f"รายชื่อนักเรียนโรงเรียน{school_name}"
    if school_year:
        sheet["F2"] = f"ปีการศึกษา {school_year}"
    if term:
        sheet["H2"] = f"เทอม {term}"
    sheet["A3"] = f"ชั้น {level}/{room}" if room else f"ชั้น {level}"

    for offset, student in enumerate(students):
        row_index = DATA_START_ROW + offset
        values: list[Any] = [
            offset + 1,
            student.full_name,
            student.citizen_id,
            student.birth_date,
            student.address,
            student.father_full_name,
            student.mother_full_name,
            student.weight,
            student.height,
        ]
        for column_index, value in enumerate(values, 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.value = value
        sheet.cell(row=row_index, column=3).number_format = "@"
        sheet.cell(row=row_index, column=4).number_format = "@"

    for row_index in range(DATA_START_ROW + len(students), int(sheet.max_row) + 1):
        for column_index in range(1, FORM_MAX_COLUMN + 1):
            sheet.cell(row=row_index, column=column_index).value = None


def _ensure_student_row_capacity(sheet: Any, *, capacity: int, needed: int) -> None:
    if needed <= capacity:
        return
    extra_rows = needed - capacity
    insert_at = DATA_START_ROW + capacity
    source_row = max(insert_at - 1, DATA_START_ROW)
    sheet.insert_rows(insert_at, amount=extra_rows)
    for row_index in range(insert_at, insert_at + extra_rows):
        _copy_row_format(sheet, source_row=source_row, target_row=row_index)


def _copy_row_format(sheet: Any, *, source_row: int, target_row: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column_index in range(1, FORM_MAX_COLUMN + 1):
        source = sheet.cell(row=source_row, column=column_index)
        target = sheet.cell(row=target_row, column=column_index)
        if source.has_style:
            target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def _apply_document_style(sheet: Any) -> None:
    for row in sheet.iter_rows(
        min_row=1,
        max_row=int(sheet.max_row),
        min_col=1,
        max_col=max(FORM_MAX_COLUMN, int(sheet.max_column)),
    ):
        for cell in row:
            font = copy(cell.font)
            font.name = DOCUMENT_FONT_NAME
            font.sz = DOCUMENT_FONT_SIZE
            cell.font = font

            alignment = copy(cell.alignment)
            alignment.wrap_text = False
            cell.alignment = alignment


def _build_output_path(source_path: Path) -> Path:
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    output_dir = reports_dir() / f"student-basic-info-{timestamp}"
    stem = re.sub(r"[^\w.-]+", "-", source_path.stem, flags=re.UNICODE).strip("-") or "dmc-student"
    candidate = output_dir / f"{stem}-student-basic-info.xlsx"
    index = 2
    while candidate.exists():
        candidate = output_dir / f"{stem}-student-basic-info-{index}.xlsx"
        index += 1
    return candidate


def _year_and_term_from_file_name(file_name: str) -> tuple[str | None, str | None]:
    match = SOURCE_YEAR_TERM_PATTERN.search(file_name)
    if match is None:
        return None, None
    return match.group("year"), match.group("term")


def _safe_sheet_title(*, level: str, room: str, existing: set[str]) -> str:
    base = INVALID_SHEET_TITLE_CHARS.sub("-", f"{level}-{room}".strip("-")).strip() or "class"
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate in existing:
        marker = f"-{suffix}"
        candidate = f"{base[:31 - len(marker)]}{marker}"
        suffix += 1
    return candidate


def _address(
    *,
    house_no: str,
    moo: str,
    road: str,
    subdistrict: str,
    district: str,
    province: str,
) -> str:
    parts: list[str] = []
    if _has_value(house_no):
        parts.append(f"บ้านเลขที่ {house_no}")
    if _has_value(moo):
        parts.append(f"หมู่ {moo}")
    if _has_value(road):
        parts.append(f"ถนน/ซอย {road}")
    if _has_value(subdistrict):
        parts.append(f"ต.{subdistrict}")
    if _has_value(district):
        parts.append(f"อ.{district}")
    if _has_value(province):
        parts.append(f"จ.{province}")
    return " ".join(parts)


def _full_name(prefix: str, first_name: str, last_name: str) -> str:
    parts = [part for part in [prefix, first_name, last_name] if _has_value(part)]
    return " ".join(parts)


def _has_value(value: str) -> bool:
    return value.strip() not in EMPTY_MARKERS


def _clean_header(value: Any) -> str:
    return _text(value).replace("\n", "").replace("\r", "").strip()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number_or_text(value: Any) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, int | float):
        return value
    return _text(value)


def _safe_int(value: str) -> int:
    try:
        return int(float(value))
    except ValueError:
        return 0
