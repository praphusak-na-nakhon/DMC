from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
from collections import Counter, defaultdict
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Literal

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator

from .config import reports_dir
from .errors import DomainError


CURRENT_STUDENTS_MODULE: Literal["currentStudents"] = "currentStudents"
CSV_MIN_COLUMNS = 25
ROSTER_STUDENT_START_ROW = 4
IMPORT_HEADER_ROW = 1
IMPORT_FIELD_KEY_ROW = 2
IMPORT_DATA_START_ROW = 3
DMC_TRANSFER_IN_URL = "https://portal.bopp-obec.info/obec69/studentin/add_cif"
TITLE_PREFIXES = (
    "เด็กชาย",
    "เด็กหญิง",
    "ด.ช.",
    "ด.ญ.",
    "นาย",
    "นางสาว",
    "นาง",
    "Master",
    "Miss",
    "Mr",
    "Mrs",
    "Ms",
)
EMPTY_MARKERS = {"", "-", "—", "–", "........................", "..................."}
DMC_DEFAULT_MISSING_SIBLING_COUNT = "0"
DMC_DEFAULT_MISSING_MARRIAGE_STATUS_CODE = "01"
DMC_DEFAULT_MISSING_OCCUPATION_CODE = "5"
DMC_DEFAULT_MISSING_COMMUTE_MINUTES = "10.0"
DMC_DEFAULT_MISSING_JOURNEY_TYPE_CODE = "02"
DMC_DEFAULT_MISSING_RACE_CODE = "099"
DMC_OTHER_OCCUPATION_CODE = "99"

IMPORT_FIELD_DEFINITIONS: tuple[tuple[str, str, bool], ...] = (
    ("operation_type", "ประเภทงาน", True),
    ("school_year", "ปีการศึกษา", True),
    ("student_no", "เลขประจำตัวนักเรียน", False),
    ("citizen_id", "เลขประจำตัวประชาชน", True),
    ("grade", "ชั้น", False),
    ("room", "ห้อง", False),
    ("seat_no", "เลขที่", False),
    ("prefix", "คำนำหน้า", True),
    ("first_name", "ชื่อ", True),
    ("last_name", "นามสกุล", True),
    ("sex", "เพศ", False),
    ("birth_date", "วันเกิด", False),
    ("first_name_en", "ชื่ออังกฤษ", False),
    ("last_name_en", "นามสกุลอังกฤษ", False),
    ("birth_province", "จังหวัดเกิด", False),
    ("birth_hospital", "สถานพยาบาลเกิด", False),
    ("blood_type", "กลุ่มเลือด", False),
    ("race", "เชื้อชาติ", False),
    ("nationality", "สัญชาติ", False),
    ("religion", "ศาสนา", False),
    ("registered_address.house_id", "รหัสประจำบ้านตามทะเบียน", False),
    ("registered_address.house_no", "บ้านเลขที่ตามทะเบียน", False),
    ("registered_address.moo", "หมู่ที่ตามทะเบียน", False),
    ("registered_address.road", "ถนนตามทะเบียน", False),
    ("registered_address.subdistrict", "ตำบลตามทะเบียน", False),
    ("registered_address.district", "อำเภอตามทะเบียน", False),
    ("registered_address.province", "จังหวัดตามทะเบียน", False),
    ("registered_address.postal_code", "รหัสไปรษณีย์ตามทะเบียน", False),
    ("current_address.house_id", "รหัสประจำบ้านปัจจุบัน", False),
    ("current_address.house_no", "บ้านเลขที่ปัจจุบัน", False),
    ("current_address.moo", "หมู่ที่ปัจจุบัน", False),
    ("current_address.road", "ถนนปัจจุบัน", False),
    ("current_address.subdistrict", "ตำบลปัจจุบัน", False),
    ("current_address.district", "อำเภอปัจจุบัน", False),
    ("current_address.province", "จังหวัดปัจจุบัน", False),
    ("current_address.postal_code", "รหัสไปรษณีย์ปัจจุบัน", False),
    ("commute_method", "การเดินทาง", False),
    ("distance_paved_road_km", "ระยะทางถนนลาดยาง กม.", False),
    ("commute_minutes", "เวลาเดินทาง นาที", False),
    ("weight_kg", "น้ำหนัก กก.", False),
    ("height_cm", "ส่วนสูง ซม.", False),
    ("parents_marital_status", "สถานภาพบิดามารดา", False),
    ("older_brothers", "จำนวนพี่ชาย", False),
    ("younger_brothers", "จำนวนน้องชาย", False),
    ("older_sisters", "จำนวนพี่สาว", False),
    ("younger_sisters", "จำนวนน้องสาว", False),
    ("siblings_studying_count", "จำนวนพี่น้องที่ศึกษาอยู่", False),
    ("child_order", "นักเรียนเป็นบุตรคนที่", False),
    ("father.citizen_id", "เลขบัตรบิดา", False),
    ("father.first_name", "ชื่อบิดา", False),
    ("father.last_name", "นามสกุลบิดา", False),
    ("father.occupation", "อาชีพบิดา", False),
    ("father.income_text", "รายได้บิดา", False),
    ("father.phone", "โทรศัพท์บิดา", False),
    ("mother.citizen_id", "เลขบัตรมารดา", False),
    ("mother.first_name", "ชื่อมารดา", False),
    ("mother.last_name", "นามสกุลมารดา", False),
    ("mother.occupation", "อาชีพมารดา", False),
    ("mother.income_text", "รายได้มารดา", False),
    ("mother.phone", "โทรศัพท์มารดา", False),
    ("guardian.citizen_id", "เลขบัตรผู้ปกครอง", False),
    ("guardian.first_name", "ชื่อผู้ปกครอง", False),
    ("guardian.last_name", "นามสกุลผู้ปกครอง", False),
    ("guardian.occupation", "อาชีพผู้ปกครอง", False),
    ("guardian.income_text", "รายได้ผู้ปกครอง", False),
    ("guardian.phone", "โทรศัพท์ผู้ปกครอง", False),
    ("guardian_relationship", "ความเกี่ยวข้องผู้ปกครอง", False),
)
IMPORT_FIELD_NAMES = tuple(field_name for field_name, _label, _required in IMPORT_FIELD_DEFINITIONS)
IMPORT_FIELD_LABELS = {field_name: label for field_name, label, _required in IMPORT_FIELD_DEFINITIONS}
REQUIRED_IMPORT_FIELDS = tuple(field_name for field_name, _label, required in IMPORT_FIELD_DEFINITIONS if required)
FORM_JSON_FIELD_NAMES = tuple(field_name for field_name in IMPORT_FIELD_NAMES if field_name != "operation_type")
ADDRESS_FIELD_SUFFIXES: tuple[str, ...] = (
    "house_id",
    "house_no",
    "moo",
    "road",
    "subdistrict",
    "district",
    "province",
    "postal_code",
)
ADDRESS_COPY_MIN_COMPARABLE_FIELDS = 3
ADDRESS_COPY_MIN_MATCH_RATIO = 0.6
GUARDIAN_PARENT_FIELD_SUFFIXES: tuple[str, ...] = (
    "citizen_id",
    "card_type",
    "first_name",
    "last_name",
    "blood_type",
    "occupation",
    "income_text",
    "phone",
)
GUARDIAN_FATHER_MARKER = "\u0e1a\u0e34\u0e14\u0e32"
GUARDIAN_MOTHER_MARKER = "\u0e21\u0e32\u0e23\u0e14\u0e32"
DMC_BLOCKING_MISSING_FIELDS: tuple[str, ...] = (
    "student_no",
    "citizen_id",
    "grade",
    "room",
    "sex",
    "prefix",
    "first_name",
    "last_name",
    "birth_date",
    "birth_province",
    "weight_kg",
    "height_cm",
    "religion",
    "race",
    "nationality",
    "registered_address.house_id",
    "registered_address.house_no",
    "registered_address.subdistrict",
    "registered_address.district",
    "registered_address.province",
    "registered_address.postal_code",
    "father.first_name",
    "father.last_name",
    "mother.first_name",
    "mother.last_name",
    "guardian.first_name",
    "guardian.last_name",
)
DMC_TRANSFER_IN_LEVEL_CODES: dict[int, str] = {
    1: "10",
    2: "11",
    3: "12",
    4: "13",
    5: "14",
    6: "15",
    10: "10",
    11: "11",
    12: "12",
    13: "13",
    14: "14",
    15: "15",
}
DMC_TRANSFER_IN_VALID_LEVEL_CODES = {f"{index:02d}" for index in range(1, 19)}
OCR_FUZZY_MATCH_THRESHOLD = 0.9
DMC_TITLE_CODE_ITEMS: tuple[tuple[str, str], ...] = (
    ("เด็กชาย", "001"),
    ("ด.ช.", "001"),
    ("เด็กหญิง", "002"),
    ("ด.ญ.", "002"),
    ("นาย", "003"),
    ("นางสาว", "004"),
    ("นาง", "005"),
    ("สามเณร", "832"),
)
DMC_GENDER_CODES = {"ชาย": "M", "หญิง": "F"}
DMC_GENDER_BY_TITLE_CODE = {
    "001": "M",
    "003": "M",
    "832": "M",
    "002": "F",
    "004": "F",
    "005": "F",
}
DMC_PROVINCE_CODES = {
    "กรุงเทพมหานคร": "10000000",
    "กระบี่": "81000000",
    "กาญจนบุรี": "71000000",
    "กาฬสินธุ์": "46000000",
    "กำแพงเพชร": "62000000",
    "ขอนแก่น": "40000000",
    "จันทบุรี": "22000000",
    "ฉะเชิงเทรา": "24000000",
    "ชลบุรี": "20000000",
    "ชัยนาท": "18000000",
    "ชัยภูมิ": "36000000",
    "ชุมพร": "86000000",
    "ตรัง": "92000000",
    "ตราด": "23000000",
    "ตาก": "63000000",
    "นครนายก": "26000000",
    "นครปฐม": "73000000",
    "นครพนม": "48000000",
    "นครราชสีมา": "30000000",
    "นครศรีธรรมราช": "80000000",
    "นครสวรรค์": "60000000",
    "นนทบุรี": "12000000",
    "นราธิวาส": "96000000",
    "น่าน": "55000000",
    "บึงกาฬ": "38000000",
    "บุรีรัมย์": "31000000",
    "ปทุมธานี": "13000000",
    "ประจวบคีรีขันธ์": "77000000",
    "ปราจีนบุรี": "25000000",
    "ปัตตานี": "94000000",
    "พระนครศรีอยุธยา": "14000000",
    "พะเยา": "56000000",
    "พังงา": "82000000",
    "พัทลุง": "93000000",
    "พิจิตร": "66000000",
    "พิษณุโลก": "65000000",
    "ภูเก็ต": "83000000",
    "มหาสารคาม": "44000000",
    "มุกดาหาร": "49000000",
    "ยะลา": "95000000",
    "ยโสธร": "35000000",
    "ระนอง": "85000000",
    "ระยอง": "21000000",
    "ราชบุรี": "70000000",
    "ร้อยเอ็ด": "45000000",
    "ลพบุรี": "16000000",
    "ลำปาง": "52000000",
    "ลำพูน": "51000000",
    "ศรีสะเกษ": "33000000",
    "สกลนคร": "47000000",
    "สงขลา": "90000000",
    "สตูล": "91000000",
    "สมุทรปราการ": "11000000",
    "สมุทรสงคราม": "75000000",
    "สมุทรสาคร": "74000000",
    "สระบุรี": "19000000",
    "สระแก้ว": "27000000",
    "สิงห์บุรี": "17000000",
    "สุพรรณบุรี": "72000000",
    "สุราษฎร์ธานี": "84000000",
    "สุรินทร์": "32000000",
    "สุโขทัย": "64000000",
    "หนองคาย": "43000000",
    "หนองบัวลำภู": "39000000",
    "อำนาจเจริญ": "37000000",
    "อุดรธานี": "41000000",
    "อุตรดิตถ์": "53000000",
    "อุทัยธานี": "61000000",
    "อุบลราชธานี": "34000000",
    "อ่างทอง": "15000000",
    "เชียงราย": "57000000",
    "เชียงใหม่": "50000000",
    "เพชรบุรี": "76000000",
    "เพชรบูรณ์": "67000000",
    "เลย": "42000000",
    "แพร่": "54000000",
    "แม่ฮ่องสอน": "58000000",
}
DMC_KRABI_DISTRICT_CODES = {
    "เมืองกระบี่": "81010000",
    "เขาพนม": "81020000",
    "เกาะลันตา": "81030000",
    "คลองท่อม": "81040000",
    "อ่าวลึก": "81050000",
    "ปลายพระยา": "81060000",
    "ลำทับ": "81070000",
    "เหนือคลอง": "81080000",
}
DMC_KRABI_SUBDISTRICT_CODES = {
    "เหนือคลอง": "81080100",
    "เกาะศรีบอยา": "81080200",
    "คลองขนาน": "81080300",
    "คลองเขม้า": "81080400",
    "โคกยาง": "81080500",
    "ตลิ่งชัน": "81080600",
    "ปกาไส": "81080700",
    "ปกาสัย": "81080700",
    "ห้วยยูง": "81080800",
}
DMC_NATION_CODES = {"ไทย": "099"}
DMC_RACE_CODES = {"ไทย": "099"}
DMC_RELIGION_CODES = {
    "พุทธ": "001",
    "อิสลาม": "002",
    "คริสต์": "003",
    "ซิกส์": "004",
    "พราหมณ์/ฮินดู": "005",
    "อื่น": "006",
    "อื่นๆ": "006",
}
DMC_BLOOD_CODES = {
    "ไม่ทราบ": "N",
    "A": "A",
    "B": "B",
    "O": "O",
    "AB": "AB",
    "ARh+": "ARh+",
    "ARh-": "ARh-",
    "BRh+": "BRh+",
    "BRh-": "BRh-",
    "ABRh+": "ABRh+",
    "ABRh-": "ABRh-",
    "ORh+": "ORh+",
    "ORh-": "ORh-",
}
DMC_PARENT_RELATION_CODES: tuple[tuple[str, str], ...] = (
    ("บิดา", "01"),
    ("มารดา", "02"),
    ("พี่", "03"),
    ("น้อง", "04"),
    ("ปู่", "05"),
    ("ย่า", "06"),
    ("ตา", "07"),
    ("ยาย", "08"),
    ("ทวด", "09"),
    ("ลุง", "10"),
    ("ป้า", "11"),
    ("น้า", "12"),
    ("อา", "13"),
    ("สามี", "16"),
    ("ภรรยา", "17"),
    ("ผู้ปกครอง", "20"),
)
DMC_MARRIAGE_STATUS_CODES: tuple[tuple[str, str], ...] = (
    ("สมรส", "01"),
    ("อยู่ด้วยกัน", "05"),
    ("หม้าย", "04"),
    ("บิดาและมารดาถึงแก่กรรม", "09"),
    ("บิดาถึงแก่กรรมมารดาแต่งงานใหม่", "10"),
    ("มารดาถึงแก่กรรมบิดาแต่งงานใหม่", "11"),
    ("อยู่ด้วยกันไม่ได้จดทะเบียนสมรส", "05"),
    ("อยู่ด้วยกันจดทะเบียนสมรส", "01"),
    ("หย่าร้าง", "04"),
    ("แยกกันอยู่", "06"),
    ("บิดาถึงแก่กรรม", "07"),
    ("มารดาถึงแก่กรรม", "08"),
    ("โสด", "02"),
)
DMC_OCCUPATION_CODES: tuple[tuple[str, str], ...] = (
    ("ไม่ได้ประกอบอาชีพ", "0"),
    ("แม่บ้าน", "0"),
    ("รับราชการ", "1"),
    ("รัฐวิสาหกิจ", "2"),
    ("ค้าขาย", "3"),
    ("ธุรกิจ", "3"),
    ("ทำสวน", "4"),
    ("ชาวสวน", "4"),
    ("เกษตร", "4"),
    ("รับจ้าง", "5"),
    ("ช่าง", "5"),
    ("ลูกจ้าง", "6"),
    ("พนักงาน", "6"),
    ("พยาบาล", "6"),
    ("เกษียณ", "7"),
    ("พระ", "8"),
    ("นักบวช", "8"),
)
DMC_CARD_TYPE_CODES: tuple[tuple[str, str], ...] = (
    ("บัตรประชาชน", "I"),
    ("ประชาชน", "I"),
    ("พาสปอร์ต", "P"),
    ("passport", "P"),
    ("ต่างด้าว", "A"),
    ("ไม่มีเอกสาร", "O"),
    ("อื่น", "O"),
)
DMC_JOURNEY_TYPE_CODES: tuple[tuple[str, str], ...] = (
    ("เดินเท้า", "01"),
    ("พาหนะไม่เสียค่าโดยสาร", "02"),
    ("พาหนะเสียค่าโดยสาร", "03"),
    ("จักรยานยืมเรียน", "04"),
)
THAI_MONTH_CODES = {
    "มกราคม": "01",
    "ม.ค.": "01",
    "กุมภาพันธ์": "02",
    "ก.พ.": "02",
    "มีนาคม": "03",
    "มี.ค.": "03",
    "เมษายน": "04",
    "เม.ย.": "04",
    "พฤษภาคม": "05",
    "พ.ค.": "05",
    "มิถุนายน": "06",
    "มิ.ย.": "06",
    "กรกฎาคม": "07",
    "ก.ค.": "07",
    "สิงหาคม": "08",
    "ส.ค.": "08",
    "กันยายน": "09",
    "ก.ย.": "09",
    "ตุลาคม": "10",
    "ต.ค.": "10",
    "พฤศจิกายน": "11",
    "พ.ย.": "11",
    "ธันวาคม": "12",
    "ธ.ค.": "12",
}
MISSING_BLOCKER_BASIS = (
    "ตรวจจากบัญชีรายชื่อ, CSV เครื่องสแกนบัตร และ OCR แบบฟอร์มแล้วไม่พบข้อมูล "
    "ต้องเติมข้อมูลนี้ก่อนนำเข้า DMC"
)

SourceType = Literal["roster", "thai_id_scan", "ocr_form", "civil_registration", "derived", "manual"]
FieldConfidence = Literal["authoritative", "high", "review", "missing"]
OperationType = Literal["current", "transfer_in", "add_new"]
MatchStatus = Literal["auto_matched", "needs_review", "duplicate", "invalid_id", "new_or_transfer_candidate"]
DmcFormValue = str | int | float | bool | list[str] | None


@dataclass(frozen=True)
class ThaiAdminCodeIndex:
    province_codes: dict[str, str]
    district_codes: dict[tuple[str, str], str]
    subdistrict_codes: dict[tuple[str, str, str], str]
    unique_subdistrict_codes_by_province: dict[tuple[str, str], str]


THAI_ADMIN_CODES_PATH = Path(__file__).resolve().parent / "data" / "thai_admin_codes.json"
_THAI_ADMIN_CODE_INDEX: ThaiAdminCodeIndex | None = None


class CurrentStudentsWarning(BaseModel):
    model_config = ConfigDict(extra="forbid")

    code: str
    message: str
    source: SourceType | None = None
    source_path: str | None = None
    row_index: int | None = None
    sheet_name: str | None = None


class SourceReference(BaseModel):
    model_config = ConfigDict(extra="forbid")

    source: SourceType
    source_path: str
    row_index: int | None = None
    sheet_name: str | None = None


class CurrentStudentField(BaseModel):
    model_config = ConfigDict(extra="forbid")

    value: str | int | float | bool | None
    source: SourceType
    confidence: FieldConfidence
    raw_value: str | None = None


class CurrentStudentsConflictValue(BaseModel):
    model_config = ConfigDict(extra="forbid")

    source: SourceType
    value: str | int | float | bool | None
    confidence: FieldConfidence
    raw_value: str | None = None
    source_path: str | None = None
    row_index: int | None = None
    sheet_name: str | None = None


class CurrentStudentsFieldConflict(BaseModel):
    model_config = ConfigDict(extra="forbid")

    record_id: str
    full_name: str | None
    student_no: str | None
    citizen_id: str | None
    field_name: str
    field_label: str
    selected_value: str | int | float | bool | None
    selected_source: SourceType | None
    selected_basis: str
    reason: str
    source_values: list[CurrentStudentsConflictValue]


class MatchSuggestion(BaseModel):
    model_config = ConfigDict(extra="forbid")

    student_no: str
    full_name: str
    grade: int
    room: int
    score: float = Field(ge=0, le=1)
    source_path: str
    sheet_name: str
    row_index: int


class DmcFormMatchConfirmation(BaseModel):
    model_config = ConfigDict(extra="forbid")

    record_id: str
    student_no: str


class RosterStudent(BaseModel):
    model_config = ConfigDict(extra="forbid")

    student_no: str
    school_year: int
    grade: int
    room: int
    seat_no: int | None
    citizen_id: str | None = None
    citizen_id_valid: bool = False
    birth_date: str | None = None
    prefix: str
    first_name: str
    last_name: str
    full_name: str
    name_key: str
    source_path: str
    sheet_name: str
    row_index: int


class ThaiIdScanRecord(BaseModel):
    model_config = ConfigDict(extra="forbid")

    citizen_id: str | None
    citizen_id_valid: bool
    prefix: str
    first_name: str
    last_name: str
    full_name: str
    name_key: str
    first_name_en: str | None
    last_name_en: str | None
    sex: str | None
    birth_date: str | None
    house_no: str | None
    moo: str | None
    subdistrict: str | None
    district: str | None
    province: str | None
    source_path: str
    row_index: int
    raw_citizen_id: str


class OcrFormRecord(BaseModel):
    model_config = ConfigDict(extra="forbid")

    record_id: str
    citizen_id: str | None
    citizen_id_valid: bool
    student_no: str | None
    prefix: str | None
    first_name: str | None
    last_name: str | None
    full_name: str | None
    name_key: str | None
    fields: dict[str, CurrentStudentField]
    source_path: str
    row_index: int | None = None


class CivilRegistrationRecord(BaseModel):
    model_config = ConfigDict(extra="forbid")

    record_id: str
    citizen_id: str | None
    citizen_id_valid: bool
    prefix: str | None
    first_name: str | None
    last_name: str | None
    full_name: str | None
    name_key: str | None
    fields: dict[str, CurrentStudentField]
    source_path: str


class _StructuredOcrRecordPayload(BaseModel):
    model_config = ConfigDict(extra="ignore")

    record_type: str | None = None
    page_start: int | None = None
    page_end: int | None = None
    fields: dict[str, Any] = Field(default_factory=dict)
    needs_review: list[str] = Field(default_factory=list)

    @field_validator("needs_review", mode="before")
    @classmethod
    def _normalize_needs_review(cls, value: Any) -> list[str]:
        if value is None:
            return []
        if isinstance(value, list):
            return [str(item) for item in value if item is not None]
        return [str(value)]


class _StructuredOcrPayload(BaseModel):
    model_config = ConfigDict(extra="ignore")

    schema_version: str | None = None
    document_type: str | None = None
    records: list[_StructuredOcrRecordPayload] = Field(default_factory=list)


class CanonicalStudentRecord(BaseModel):
    model_config = ConfigDict(extra="forbid")

    record_id: str
    operation_type: OperationType
    match_status: MatchStatus
    student_no: str | None
    citizen_id: str | None
    grade: int | None
    room: int | None
    seat_no: int | None
    prefix: str | None
    first_name: str | None
    last_name: str | None
    full_name: str | None
    match_score: float | None = Field(default=None, ge=0, le=1)
    review_reasons: list[str] = Field(default_factory=list)
    suggestions: list[MatchSuggestion] = Field(default_factory=list)
    dmc_fields: dict[str, CurrentStudentField] = Field(default_factory=dict)
    sources: list[SourceReference] = Field(default_factory=list)
    conflicts: list[CurrentStudentsFieldConflict] = Field(default_factory=list)


class CurrentStudentsSummary(BaseModel):
    model_config = ConfigDict(extra="forbid")

    roster_records: int
    thai_id_scan_records: int
    ocr_form_records: int
    civil_registration_records: int = 0
    records_total: int
    auto_matched: int
    needs_review: int
    duplicate_records: int
    duplicate_scan_records: int
    invalid_id_records: int
    new_or_transfer_candidates: int
    roster_without_thai_id: int
    ocr_attached_records: int
    ocr_unmatched_records: int
    review_queue_records: int
    warnings_total: int


class ReconcileCurrentStudentsRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    roster_excel_path: str
    thai_id_csv_path: str | None = None
    ocr_markdown_paths: list[str] = Field(default_factory=list)
    civil_registration_markdown_paths: list[str] = Field(default_factory=list)
    school_year: int
    grade_levels: list[int] | None = None
    operation_type: OperationType = "current"
    fuzzy_match_threshold: float = Field(default=0.9, ge=0, le=1)
    confirmed_matches: list[DmcFormMatchConfirmation] = Field(default_factory=list)


class CurrentStudentsReconciliationResponse(BaseModel):
    model_config = ConfigDict(extra="forbid")

    module: Literal["currentStudents"]
    school_year: int
    grade_levels: list[int] | None
    operation_type: OperationType
    roster_excel_path: str
    thai_id_csv_path: str | None
    ocr_markdown_paths: list[str]
    civil_registration_markdown_paths: list[str]
    summary: CurrentStudentsSummary
    records: list[CanonicalStudentRecord]
    review_queue: list[CanonicalStudentRecord]
    warnings: list[CurrentStudentsWarning]


class ExportCurrentStudentsBlankFormRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    output_path: str


class ExportCurrentStudentsBlankFormResponse(BaseModel):
    model_config = ConfigDict(extra="forbid")

    module: Literal["currentStudents"]
    output_path: str
    field_count: int
    required_fields: list[str]


class ExportCurrentStudentsImportExcelRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    roster_excel_path: str
    thai_id_csv_path: str | None = None
    ocr_markdown_paths: list[str] = Field(default_factory=list)
    civil_registration_markdown_paths: list[str] = Field(default_factory=list)
    school_year: int
    grade_levels: list[int] | None = None
    operation_type: OperationType = "current"
    output_path: str | None = None


class ExportCurrentStudentsImportExcelResponse(BaseModel):
    model_config = ConfigDict(extra="forbid")

    module: Literal["formConverter"]
    output_path: str
    school_year: int
    operation_type: OperationType
    rows_exported: int
    summary: CurrentStudentsSummary
    warnings: list[CurrentStudentsWarning]
    conflicts: list[CurrentStudentsFieldConflict]


class DmcFormJsonRecord(BaseModel):
    model_config = ConfigDict(extra="forbid")

    record_id: str
    match_status: MatchStatus
    student_no: str | None
    citizen_id: str | None
    grade: int | None
    room: int | None
    seat_no: int | None
    prefix: str | None
    first_name: str | None
    last_name: str | None
    full_name: str | None
    match_score: float | None = Field(default=None, ge=0, le=1)
    review_reasons: list[str] = Field(default_factory=list)
    suggestions: list[MatchSuggestion] = Field(default_factory=list)
    sources: list[SourceReference] = Field(default_factory=list)
    fields: dict[str, str | int | float | bool | None]
    dmc_form_values: dict[str, DmcFormValue] = Field(default_factory=dict)
    field_details: dict[str, CurrentStudentField]


class PreviewDmcFormJsonRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    roster_excel_path: str
    thai_id_csv_path: str | None = None
    ocr_markdown_paths: list[str] = Field(default_factory=list)
    civil_registration_markdown_paths: list[str] = Field(default_factory=list)
    school_year: int
    grade_levels: list[int] | None = None
    admission_date: str | None = None
    confirmed_matches: list[DmcFormMatchConfirmation] = Field(default_factory=list)


class PreviewDmcFormJsonResponse(BaseModel):
    model_config = ConfigDict(extra="forbid")

    module: Literal["formConverter"]
    schema_version: Literal["dmc_form_json.v1"]
    generated_at: str
    school_year: int
    grade_levels: list[int] | None
    records_previewed: int
    field_labels: dict[str, str]
    summary: CurrentStudentsSummary
    warnings: list[CurrentStudentsWarning]
    conflicts: list[CurrentStudentsFieldConflict]
    records: list[DmcFormJsonRecord]


class ExportDmcFormJsonRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    roster_excel_path: str
    thai_id_csv_path: str | None = None
    ocr_markdown_paths: list[str] = Field(default_factory=list)
    civil_registration_markdown_paths: list[str] = Field(default_factory=list)
    school_year: int
    grade_levels: list[int] | None = None
    admission_date: str | None = None
    output_path: str | None = None
    confirmed_matches: list[DmcFormMatchConfirmation] = Field(default_factory=list)
    excluded_record_ids: list[str] = Field(default_factory=list)


class ExportDmcFormJsonResponse(BaseModel):
    model_config = ConfigDict(extra="forbid")

    module: Literal["formConverter"]
    schema_version: Literal["dmc_form_json.v1"]
    generated_at: str
    output_path: str
    school_year: int
    grade_levels: list[int] | None
    records_exported: int
    field_labels: dict[str, str]
    summary: CurrentStudentsSummary
    warnings: list[CurrentStudentsWarning]
    conflicts: list[CurrentStudentsFieldConflict]
    records: list[DmcFormJsonRecord]


class CurrentStudentsImportRowPreview(BaseModel):
    model_config = ConfigDict(extra="forbid")

    row_index: int
    status: Literal["ready", "needs_review", "invalid"]
    operation_type: str | None
    student_no: str | None
    citizen_id: str | None
    full_name: str
    issues: list[str] = Field(default_factory=list)


class CurrentStudentsImportSummary(BaseModel):
    model_config = ConfigDict(extra="forbid")

    rows_total: int
    ready_rows: int
    needs_review_rows: int
    invalid_rows: int
    duplicate_citizen_ids: int
    warnings_total: int


class DmcTransferInImportRecord(BaseModel):
    model_config = ConfigDict(extra="forbid")

    row_index: int
    record_id: str
    student_no: str
    level_dtl_code: str
    classroom: str
    citizen_id: str
    full_name: str
    dmc_form_values: dict[str, DmcFormValue] = Field(default_factory=dict)


class ValidateCurrentStudentsImportFormRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    excel_path: str


class ValidateCurrentStudentsImportFormResponse(BaseModel):
    model_config = ConfigDict(extra="forbid")

    module: Literal["currentStudents"]
    excel_path: str
    summary: CurrentStudentsImportSummary
    preview: list[CurrentStudentsImportRowPreview]
    warnings: list[CurrentStudentsWarning]


class _DmcFormJsonImportPayload(BaseModel):
    model_config = ConfigDict(extra="ignore")

    schema_version: Literal["dmc_form_json.v1"]
    records: list[DmcFormJsonRecord]


def reconcile_current_students(request: ReconcileCurrentStudentsRequest) -> CurrentStudentsReconciliationResponse:
    roster_path = _validated_file(Path(request.roster_excel_path), suffixes={".xlsx", ".xlsm"})
    roster, warnings = read_student_roster(
        roster_path,
        school_year=request.school_year,
        grade_levels=request.grade_levels,
    )
    if not roster:
        raise DomainError("CURRENT_STUDENTS_ROSTER_EMPTY", "No roster rows were found for the requested year/grade.")

    scans: list[ThaiIdScanRecord] = []
    if request.thai_id_csv_path:
        scan_path = _validated_file(Path(request.thai_id_csv_path), suffixes={".csv"})
        parsed_scans, scan_warnings = read_thai_id_scan_csv(scan_path)
        scans = parsed_scans
        warnings.extend(scan_warnings)

    ocr_records: list[OcrFormRecord] = []
    for markdown_path in request.ocr_markdown_paths:
        path = _validated_file(Path(markdown_path), suffixes={".md", ".txt", ".csv", ".json"})
        parsed_ocr_records, ocr_warnings = read_ocr_records(path, record_index_start=len(ocr_records) + 1)
        ocr_records.extend(parsed_ocr_records)
        warnings.extend(ocr_warnings)

    civil_records: list[CivilRegistrationRecord] = []
    for markdown_path in request.civil_registration_markdown_paths:
        path = _validated_file(Path(markdown_path), suffixes={".md", ".txt", ".csv", ".json"})
        parsed_civil_records, civil_warnings = read_civil_registration_markdown_records(
            path,
            record_index_start=len(civil_records) + 1,
        )
        civil_records.extend(parsed_civil_records)
        warnings.extend(civil_warnings)

    records, ocr_attached, ocr_unmatched = _build_canonical_records(
        roster=roster,
        scans=scans,
        ocr_records=ocr_records,
        civil_records=civil_records,
        operation_type=request.operation_type,
        fuzzy_match_threshold=request.fuzzy_match_threshold,
        confirmed_ocr_matches=_confirmed_match_map(request.confirmed_matches),
    )
    records = sorted(records, key=_record_sort_key)
    review_queue = [record for record in records if record.match_status != "auto_matched"]
    summary = _summary(
        roster=roster,
        scans=scans,
        ocr_records=ocr_records,
        civil_records=civil_records,
        records=records,
        ocr_attached=ocr_attached,
        ocr_unmatched=ocr_unmatched,
        warnings=warnings,
    )
    return CurrentStudentsReconciliationResponse(
        module=CURRENT_STUDENTS_MODULE,
        school_year=request.school_year,
        grade_levels=request.grade_levels,
        operation_type=request.operation_type,
        roster_excel_path=str(roster_path),
        thai_id_csv_path=request.thai_id_csv_path,
        ocr_markdown_paths=request.ocr_markdown_paths,
        civil_registration_markdown_paths=request.civil_registration_markdown_paths,
        summary=summary,
        records=records,
        review_queue=review_queue,
        warnings=warnings,
    )


def export_current_students_blank_form(
    request: ExportCurrentStudentsBlankFormRequest,
) -> ExportCurrentStudentsBlankFormResponse:
    output_path = Path(request.output_path)
    if output_path.suffix.lower() != ".xlsx":
        raise DomainError("CURRENT_STUDENTS_TEMPLATE_UNSUPPORTED_TYPE", "Blank current student form must be .xlsx.")
    _write_import_workbook(output_path, records=[])
    return ExportCurrentStudentsBlankFormResponse(
        module=CURRENT_STUDENTS_MODULE,
        output_path=str(output_path),
        field_count=len(IMPORT_FIELD_DEFINITIONS),
        required_fields=list(REQUIRED_IMPORT_FIELDS),
    )


def export_current_students_import_excel(
    request: ExportCurrentStudentsImportExcelRequest,
) -> ExportCurrentStudentsImportExcelResponse:
    if not request.ocr_markdown_paths:
        raise DomainError(
            "CURRENT_STUDENTS_OCR_REQUIRED",
            "At least one OCR markdown form is required for DMC form conversion export.",
        )
    reconciliation = reconcile_current_students(
        ReconcileCurrentStudentsRequest(
            roster_excel_path=request.roster_excel_path,
            thai_id_csv_path=request.thai_id_csv_path,
            ocr_markdown_paths=request.ocr_markdown_paths,
            civil_registration_markdown_paths=request.civil_registration_markdown_paths,
            school_year=request.school_year,
            grade_levels=request.grade_levels,
            operation_type=request.operation_type,
        )
    )
    export_records = [record for record in reconciliation.records if _has_source(record, "ocr_form")]
    if not export_records:
        raise DomainError(
            "CURRENT_STUDENTS_OCR_NO_RECORDS",
            "No records from OCR markdown forms were available for export.",
        )
    output_path = Path(request.output_path) if request.output_path else _default_import_output_path(request.school_year)
    if output_path.suffix.lower() != ".xlsx":
        raise DomainError("CURRENT_STUDENTS_EXPORT_UNSUPPORTED_TYPE", "Current student import file must be .xlsx.")
    _write_import_workbook(output_path, records=export_records, default_school_year=request.school_year)
    return ExportCurrentStudentsImportExcelResponse(
        module="formConverter",
        output_path=str(output_path),
        school_year=request.school_year,
        operation_type=request.operation_type,
        rows_exported=len(export_records),
        summary=_export_summary(reconciliation.summary, export_records, warnings=reconciliation.warnings),
        warnings=reconciliation.warnings,
        conflicts=_export_missing_blockers(export_records, default_school_year=request.school_year),
    )


def preview_dmc_form_json(request: PreviewDmcFormJsonRequest) -> PreviewDmcFormJsonResponse:
    payload = _build_dmc_form_json_payload(
        roster_excel_path=request.roster_excel_path,
        thai_id_csv_path=request.thai_id_csv_path,
        ocr_markdown_paths=request.ocr_markdown_paths,
        civil_registration_markdown_paths=request.civil_registration_markdown_paths,
        school_year=request.school_year,
        grade_levels=request.grade_levels,
        admission_date=request.admission_date,
        confirmed_matches=request.confirmed_matches,
        skip_unconfirmed_ocr_records=False,
    )
    return PreviewDmcFormJsonResponse(
        module="formConverter",
        schema_version="dmc_form_json.v1",
        generated_at=datetime.now(timezone.utc).isoformat(),
        school_year=request.school_year,
        grade_levels=request.grade_levels,
        records_previewed=len(payload.records),
        field_labels=payload.field_labels,
        summary=payload.summary,
        warnings=payload.warnings,
        conflicts=payload.conflicts,
        records=payload.records,
    )


def export_dmc_form_json(request: ExportDmcFormJsonRequest) -> ExportDmcFormJsonResponse:
    payload = _build_dmc_form_json_payload(
        roster_excel_path=request.roster_excel_path,
        thai_id_csv_path=request.thai_id_csv_path,
        ocr_markdown_paths=request.ocr_markdown_paths,
        civil_registration_markdown_paths=request.civil_registration_markdown_paths,
        school_year=request.school_year,
        grade_levels=request.grade_levels,
        admission_date=request.admission_date,
        confirmed_matches=request.confirmed_matches,
        skip_unconfirmed_ocr_records=True,
        excluded_record_ids=request.excluded_record_ids,
    )
    output_path = Path(request.output_path) if request.output_path else _default_form_json_output_path(request.school_year)
    if output_path.suffix.lower() != ".json":
        raise DomainError("DMC_FORM_JSON_EXPORT_UNSUPPORTED_TYPE", "DMC form conversion output file must be .json.")

    response = ExportDmcFormJsonResponse(
        module="formConverter",
        schema_version="dmc_form_json.v1",
        generated_at=datetime.now(timezone.utc).isoformat(),
        output_path=str(output_path),
        school_year=request.school_year,
        grade_levels=request.grade_levels,
        records_exported=len(payload.records),
        field_labels=payload.field_labels,
        summary=payload.summary,
        warnings=payload.warnings,
        conflicts=payload.conflicts,
        records=payload.records,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(response.model_dump(mode="json"), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return response


class _DmcFormJsonPayload(BaseModel):
    model_config = ConfigDict(extra="forbid")

    field_labels: dict[str, str]
    summary: CurrentStudentsSummary
    warnings: list[CurrentStudentsWarning]
    conflicts: list[CurrentStudentsFieldConflict]
    records: list[DmcFormJsonRecord]


def _build_dmc_form_json_payload(
    *,
    roster_excel_path: str,
    thai_id_csv_path: str | None,
    ocr_markdown_paths: list[str],
    civil_registration_markdown_paths: list[str],
    school_year: int,
    grade_levels: list[int] | None,
    admission_date: str | None,
    confirmed_matches: list[DmcFormMatchConfirmation],
    skip_unconfirmed_ocr_records: bool,
    excluded_record_ids: Sequence[str] = (),
) -> _DmcFormJsonPayload:
    if not ocr_markdown_paths:
        raise DomainError(
            "CURRENT_STUDENTS_OCR_REQUIRED",
            "At least one OCR markdown form is required for DMC form conversion export.",
        )
    reconciliation = reconcile_current_students(
        ReconcileCurrentStudentsRequest(
            roster_excel_path=roster_excel_path,
            thai_id_csv_path=thai_id_csv_path,
            ocr_markdown_paths=ocr_markdown_paths,
            civil_registration_markdown_paths=civil_registration_markdown_paths,
            school_year=school_year,
            grade_levels=grade_levels,
            operation_type="current",
            confirmed_matches=confirmed_matches,
        )
    )
    ocr_records = [record for record in reconciliation.records if _has_source(record, "ocr_form")]
    if not ocr_records:
        raise DomainError(
            "CURRENT_STUDENTS_OCR_NO_RECORDS",
            "No OCR markdown student records were found for DMC form conversion export.",
        )
    excluded_record_id_set = set(excluded_record_ids)
    export_records = [
        record
        for record in ocr_records
        if not skip_unconfirmed_ocr_records or not _is_unconfirmed_ocr_record(record)
        if record.record_id not in excluded_record_id_set
    ]

    return _DmcFormJsonPayload(
        field_labels={field_name: IMPORT_FIELD_LABELS[field_name] for field_name in FORM_JSON_FIELD_NAMES},
        summary=_export_summary(reconciliation.summary, export_records, warnings=reconciliation.warnings),
        warnings=reconciliation.warnings,
        conflicts=_export_missing_blockers(export_records, default_school_year=school_year),
        records=[
            _json_record(record, default_school_year=school_year, default_admission_date=admission_date)
            for record in export_records
        ],
    )


def validate_current_students_import_form(
    request: ValidateCurrentStudentsImportFormRequest,
) -> ValidateCurrentStudentsImportFormResponse:
    input_path = _validated_file(Path(request.excel_path), suffixes={".xlsx", ".xlsm", ".json"})
    if input_path.suffix.lower() == ".json":
        return _validate_current_students_json_import_form(input_path)
    return _validate_current_students_excel_import_form(input_path)


def _validate_current_students_excel_import_form(excel_path: Path) -> ValidateCurrentStudentsImportFormResponse:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    sheet = workbook.active
    field_columns = _read_import_field_columns(sheet)
    if not field_columns:
        raise DomainError(
            "CURRENT_STUDENTS_IMPORT_HEADERS_NOT_FOUND",
            "Current student import form headers were not found.",
        )

    row_values: list[tuple[int, dict[str, str]]] = []
    blank_streak = 0
    for row_index, row in enumerate(sheet.iter_rows(min_row=IMPORT_DATA_START_ROW, values_only=True), start=IMPORT_DATA_START_ROW):
        values = {
            field_name: _clean_text(row[column_index - 1] if column_index - 1 < len(row) else None)
            for field_name, column_index in field_columns.items()
        }
        if not any(values.values()):
            blank_streak += 1
            if blank_streak >= 20:
                break
            continue
        blank_streak = 0
        row_values.append((row_index, values))

    citizen_counts = Counter(values.get("citizen_id", "") for _row_index, values in row_values if values.get("citizen_id"))
    duplicate_citizen_ids = {citizen_id for citizen_id, count in citizen_counts.items() if count > 1}
    warnings: list[CurrentStudentsWarning] = []
    preview: list[CurrentStudentsImportRowPreview] = []
    status_counts: Counter[str] = Counter()
    for row_index, values in row_values:
        issues = _validate_import_row(values, duplicate_citizen_ids)
        for issue in issues:
            warnings.append(
                CurrentStudentsWarning(
                    code=issue,
                    message=_import_issue_message(issue),
                    source="manual",
                    source_path=str(excel_path),
                    row_index=row_index,
                )
            )
        status = _import_row_status(issues)
        status_counts[status] += 1
        preview.append(
            CurrentStudentsImportRowPreview(
                row_index=row_index,
                status=status,
                operation_type=values.get("operation_type") or None,
                student_no=values.get("student_no") or None,
                citizen_id=values.get("citizen_id") or None,
                full_name=_full_name(values.get("prefix", ""), values.get("first_name", ""), values.get("last_name", "")),
                issues=issues,
            )
        )

    return ValidateCurrentStudentsImportFormResponse(
        module=CURRENT_STUDENTS_MODULE,
        excel_path=str(excel_path),
        summary=CurrentStudentsImportSummary(
            rows_total=len(row_values),
            ready_rows=status_counts["ready"],
            needs_review_rows=status_counts["needs_review"],
            invalid_rows=status_counts["invalid"],
            duplicate_citizen_ids=len(duplicate_citizen_ids),
            warnings_total=len(warnings),
        ),
        preview=preview[:100],
        warnings=warnings,
    )


def _validate_current_students_json_import_form(json_path: Path) -> ValidateCurrentStudentsImportFormResponse:
    records = _read_dmc_form_json_import_records(json_path)
    duplicate_citizen_ids = _json_duplicate_citizen_ids(records)

    warnings: list[CurrentStudentsWarning] = []
    preview: list[CurrentStudentsImportRowPreview] = []
    status_counts: Counter[str] = Counter()
    for row_index, record in enumerate(records, start=1):
        issues = _validate_dmc_transfer_in_json_record(record, duplicate_citizen_ids)
        for issue in issues:
            warnings.append(
                CurrentStudentsWarning(
                    code=issue,
                    message=_import_issue_message(issue),
                    source="manual",
                    source_path=str(json_path),
                    row_index=row_index,
                )
            )
        status = _import_row_status(issues)
        status_counts[status] += 1
        preview.append(
            CurrentStudentsImportRowPreview(
                row_index=row_index,
                status=status,
                operation_type="transfer_in",
                student_no=_json_record_student_no(record),
                citizen_id=_json_record_citizen_id(record),
                full_name=_json_record_full_name(record),
                issues=issues,
            )
        )

    return ValidateCurrentStudentsImportFormResponse(
        module=CURRENT_STUDENTS_MODULE,
        excel_path=str(json_path),
        summary=CurrentStudentsImportSummary(
            rows_total=len(records),
            ready_rows=status_counts["ready"],
            needs_review_rows=status_counts["needs_review"],
            invalid_rows=status_counts["invalid"],
            duplicate_citizen_ids=len(duplicate_citizen_ids),
            warnings_total=len(warnings),
        ),
        preview=preview[:100],
        warnings=warnings,
    )


def load_dmc_transfer_in_import_records(json_path: Path) -> list[DmcTransferInImportRecord]:
    input_path = _validated_file(json_path, suffixes={".json"})
    records = _read_dmc_form_json_import_records(input_path)
    duplicate_citizen_ids = _json_duplicate_citizen_ids(records)
    import_records: list[DmcTransferInImportRecord] = []
    invalid_rows: list[str] = []
    for row_index, record in enumerate(records, start=1):
        issues = _validate_dmc_transfer_in_json_record(record, duplicate_citizen_ids)
        if issues:
            invalid_rows.append(f"{row_index}:{','.join(issues)}")
            continue
        student_no = _json_record_student_no(record)
        level_dtl_code = _json_record_level_code(record)
        classroom = _json_record_room(record)
        citizen_id = _clean_citizen_id(_json_record_citizen_id(record))
        if student_no is None or level_dtl_code is None or classroom is None or citizen_id is None:
            invalid_rows.append(f"{row_index}:CURRENT_STUDENTS_JSON_ROW_INCOMPLETE")
            continue
        import_records.append(
            DmcTransferInImportRecord(
                row_index=row_index,
                record_id=record.record_id,
                student_no=student_no,
                level_dtl_code=level_dtl_code,
                classroom=classroom,
                citizen_id=citizen_id,
                full_name=_json_record_full_name(record),
                dmc_form_values=record.dmc_form_values,
            )
        )

    if invalid_rows:
        raise DomainError(
            "CURRENT_STUDENTS_IMPORT_NOT_READY",
            "Current student JSON contains rows that are not ready for DMC transfer-in import.",
            details={"invalid_rows": invalid_rows[:20]},
        )
    return import_records


def read_thai_id_scan_csv(path: Path) -> tuple[list[ThaiIdScanRecord], list[CurrentStudentsWarning]]:
    text, _encoding = _decode_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect=dialect))
    if not rows:
        raise DomainError("CURRENT_STUDENTS_THAI_ID_CSV_EMPTY", "Thai ID scan CSV is empty.")

    warnings: list[CurrentStudentsWarning] = []
    records: list[ThaiIdScanRecord] = []
    for row_index, row in enumerate(rows[1:], start=2):
        if _row_is_empty(row):
            continue
        if len(row) < CSV_MIN_COLUMNS:
            warnings.append(
                CurrentStudentsWarning(
                    code="THAI_ID_CSV_SHORT_ROW",
                    message="Thai ID scan row has fewer columns than expected.",
                    source="thai_id_scan",
                    source_path=str(path),
                    row_index=row_index,
                )
            )
            continue
        raw_citizen_id = _clean_text(row[6])
        citizen_id = _clean_citizen_id(raw_citizen_id)
        citizen_id_valid = citizen_id is not None and is_valid_thai_citizen_id(citizen_id)
        if not citizen_id_valid:
            warnings.append(
                CurrentStudentsWarning(
                    code="THAI_ID_INVALID_OR_MASKED",
                    message="Citizen ID from Thai ID scan is missing, masked, or failed checksum validation.",
                    source="thai_id_scan",
                    source_path=str(path),
                    row_index=row_index,
                )
            )
        prefix = _clean_text(row[7])
        first_name = _clean_text(row[8])
        last_name = _clean_text(row[10])
        records.append(
            ThaiIdScanRecord(
                citizen_id=citizen_id,
                citizen_id_valid=citizen_id_valid,
                prefix=prefix,
                first_name=first_name,
                last_name=last_name,
                full_name=_full_name(prefix, first_name, last_name),
                name_key=normalized_name(first_name, last_name),
                first_name_en=_none_if_empty(row[12]),
                last_name_en=_none_if_empty(row[14]),
                sex=_none_if_empty(row[23]),
                birth_date=_none_if_empty(row[24]),
                house_no=_none_if_empty(row[15]),
                moo=_none_if_empty(row[16]),
                subdistrict=_none_if_empty(row[20]),
                district=_none_if_empty(row[21]),
                province=_none_if_empty(row[22]),
                source_path=str(path),
                row_index=row_index,
                raw_citizen_id=raw_citizen_id,
            )
        )
    return records, warnings


def read_student_roster(
    path: Path,
    *,
    school_year: int,
    grade_levels: Sequence[int] | None = None,
) -> tuple[list[RosterStudent], list[CurrentStudentsWarning]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    wanted_grades = set(grade_levels) if grade_levels is not None else None
    warnings: list[CurrentStudentsWarning] = []
    records: list[RosterStudent] = []
    for sheet in workbook.worksheets:
        sheet_name = str(sheet.title)
        header_records, header_warnings = _read_header_roster_sheet(
            sheet,
            path=path,
            school_year=school_year,
            wanted_grades=wanted_grades,
        )
        records.extend(header_records)
        warnings.extend(header_warnings)
        if header_records:
            continue

        match = re.fullmatch(r"(?P<grade>\d+)[.](?P<room>\d+)", sheet_name)
        if match is None:
            continue
        grade = int(match.group("grade"))
        room = int(match.group("room"))
        if wanted_grades is not None and grade not in wanted_grades:
            continue
        detected_year = _detect_roster_year(sheet)
        if detected_year != school_year:
            continue
        blank_streak = 0
        for row_index, row in enumerate(sheet.iter_rows(min_row=ROSTER_STUDENT_START_ROW, values_only=True), start=4):
            values = list(row[:4]) + [None] * max(0, 4 - len(row))
            seat_no, student_no, prefix, full_name = values[:4]
            if not any(_clean_text(value) for value in values):
                blank_streak += 1
                if blank_streak >= 10:
                    break
                continue
            blank_streak = 0
            if not _looks_numeric(student_no) or not _clean_text(full_name):
                continue
            first_name, last_name = _split_roster_name(_clean_text(full_name))
            records.append(
                RosterStudent(
                    student_no=str(int(float(str(student_no)))),
                    school_year=school_year,
                    grade=grade,
                    room=room,
                    seat_no=int(float(str(seat_no))) if _looks_numeric(seat_no) else None,
                    prefix=_clean_text(prefix),
                    first_name=first_name,
                    last_name=last_name,
                    full_name=_clean_text(full_name),
                    name_key=normalized_name(first_name, last_name),
                    source_path=str(path),
                    sheet_name=sheet_name,
                    row_index=row_index,
                )
            )
    if not records:
        warnings.append(
            CurrentStudentsWarning(
                code="ROSTER_NO_MATCHING_SHEETS",
                message="No roster sheets matched the requested school year and grade filters.",
                source="roster",
                source_path=str(path),
            )
        )
    return records, warnings


def _read_header_roster_sheet(
    sheet: Any,
    *,
    path: Path,
    school_year: int,
    wanted_grades: set[int] | None,
) -> tuple[list[RosterStudent], list[CurrentStudentsWarning]]:
    header_row_index: int | None = None
    header_columns: dict[str, int] = {}
    for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        columns: dict[str, int] = {}
        for column_index, value in enumerate(row):
            key = _roster_header_key(value)
            if key is not None:
                columns[key] = column_index
        required = {"student_no", "full_name", "citizen_id", "grade", "room"}
        if required.issubset(columns):
            header_row_index = row_index
            header_columns = columns
            break
    if header_row_index is None:
        return [], []

    records: list[RosterStudent] = []
    warnings: list[CurrentStudentsWarning] = []
    blank_streak = 0
    for row_index, row in enumerate(sheet.iter_rows(min_row=header_row_index + 1, values_only=True), start=header_row_index + 1):
        if _row_is_empty([str(value) if value is not None else "" for value in row]):
            blank_streak += 1
            if blank_streak >= 10:
                break
            continue
        blank_streak = 0
        student_no = _student_no_text(_cell_by_header(row, header_columns, "student_no"))
        full_name = _clean_text(_cell_by_header(row, header_columns, "full_name"))
        grade = _grade_level_number(_cell_by_header(row, header_columns, "grade"))
        room = _int_cell(_cell_by_header(row, header_columns, "room"))
        if student_no is None or not full_name or grade is None or room is None:
            continue
        if wanted_grades is not None and grade not in wanted_grades:
            continue

        raw_citizen_id = _clean_text(_cell_by_header(row, header_columns, "citizen_id"))
        citizen_id = _clean_citizen_id(raw_citizen_id)
        citizen_id_valid = citizen_id is not None and is_valid_thai_citizen_id(citizen_id)
        if raw_citizen_id and not citizen_id_valid:
            warnings.append(
                CurrentStudentsWarning(
                    code="ROSTER_CITIZEN_ID_INVALID",
                    message="Citizen ID from student roster is missing or failed checksum validation.",
                    source="roster",
                    source_path=str(path),
                    row_index=row_index,
                    sheet_name=str(sheet.title),
                )
            )

        prefix, first_name, last_name = _split_roster_full_name(full_name)
        records.append(
            RosterStudent(
                student_no=student_no,
                school_year=school_year,
                grade=grade,
                room=room,
                seat_no=_int_cell(_cell_by_header(row, header_columns, "seat_no")),
                citizen_id=citizen_id if citizen_id_valid else None,
                citizen_id_valid=citizen_id_valid,
                birth_date=_none_if_empty(_cell_by_header(row, header_columns, "birth_date")),
                prefix=prefix or "",
                first_name=first_name,
                last_name=last_name,
                full_name=_full_name(prefix or "", first_name, last_name) or full_name,
                name_key=normalized_name(first_name, last_name),
                source_path=str(path),
                sheet_name=str(sheet.title),
                row_index=row_index,
            )
        )
    return records, warnings


def read_ocr_records(
    path: Path,
    *,
    record_index_start: int = 1,
) -> tuple[list[OcrFormRecord], list[CurrentStudentsWarning]]:
    if path.suffix.lower() == ".json":
        return read_ocr_structured_json_records(path, record_index_start=record_index_start)
    record, warnings = read_ocr_markdown(path, record_index=record_index_start)
    return [record], warnings


def read_ocr_structured_json_records(
    path: Path,
    *,
    record_index_start: int = 1,
) -> tuple[list[OcrFormRecord], list[CurrentStudentsWarning]]:
    payload = _read_structured_ocr_payload(path)
    warnings: list[CurrentStudentsWarning] = []
    records: list[OcrFormRecord] = []
    for record_payload in payload.records:
        if _structured_record_kind(record_payload, payload.document_type) == "civil_registration":
            continue
        record, record_warnings = _structured_payload_to_ocr_form_record(
            path,
            record_payload,
            record_index=record_index_start + len(records),
        )
        records.append(record)
        warnings.extend(record_warnings)
    return records, warnings


def read_ocr_markdown(path: Path, *, record_index: int = 1) -> tuple[OcrFormRecord, list[CurrentStudentsWarning]]:
    if path.suffix.lower() == ".json":
        records, structured_warnings = read_ocr_structured_json_records(path, record_index_start=record_index)
        if not records:
            raise DomainError("OCR_STRUCTURED_JSON_NO_DMC_RECORDS", "Structured OCR JSON contains no DMC form records.")
        return records[0], structured_warnings

    text, _encoding = _decode_text(path)
    warnings: list[CurrentStudentsWarning] = []
    fields: dict[str, CurrentStudentField] = {}

    citizen_id = _clean_citizen_id(_extract_between(text, "เลขประจำตัวประชาชน*", ("เลขประจำตัวนักเรียน",)))
    citizen_id_valid = citizen_id is not None and is_valid_thai_citizen_id(citizen_id)
    if not citizen_id_valid:
        warnings.append(
            CurrentStudentsWarning(
                code="OCR_CITIZEN_ID_INVALID_OR_MISSING",
                message="OCR form citizen ID is missing or failed checksum validation.",
                source="ocr_form",
                source_path=str(path),
            )
        )
    student_no = _first_digits(_extract_between(text, "เลขประจำตัวนักเรียน", ("คำนำหน้านาม*",)))
    prefix = _none_if_empty(_extract_between(text, "คำนำหน้านาม*", ("เพศ*",)))
    sex = _none_if_empty(_extract_between(text, "เพศ*", ("ชื่อ*",)))
    first_name = _none_if_empty(_extract_between(text, "ชื่อ*", ("นามสกุล*",)))
    last_name = _none_if_empty(_extract_between(text, "นามสกุล*", ("ชื่อ (อังกฤษ)*",)))
    full_name = _full_name(prefix or "", first_name or "", last_name or "") if first_name or last_name else None
    name_key = normalized_name(first_name or "", last_name or "") if first_name or last_name else None

    _put_field(fields, "citizen_id", citizen_id, "ocr_form", "review")
    _put_field(fields, "student_no", student_no, "ocr_form", "review")
    _put_field(fields, "prefix", prefix, "ocr_form", "review")
    _put_field(fields, "sex", sex, "ocr_form", "review")
    _put_field(fields, "first_name", first_name, "ocr_form", "review")
    _put_field(fields, "last_name", last_name, "ocr_form", "review")
    _put_field(fields, "first_name_en", _extract_between(text, "ชื่อ (อังกฤษ)*", ("นามสกุล (อังกฤษ)*",)), "ocr_form", "review")
    _put_field(fields, "last_name_en", _extract_between(text, "นามสกุล (อังกฤษ)*", ("วัน/เดือน/ปีเกิด*",)), "ocr_form", "review")
    _put_field(fields, "birth_date", _extract_between(text, "วัน/เดือน/ปีเกิด*", ("จังหวัดที่เกิด*",)), "ocr_form", "review")
    _put_field(fields, "birth_province", _extract_between(text, "จังหวัดที่เกิด*", ("สถานพยาบาลที่เกิด*",)), "ocr_form", "review")
    _put_field(fields, "birth_hospital", _extract_between(text, "สถานพยาบาลที่เกิด*", ("กลุ่มเลือด*",)), "ocr_form", "review")
    _put_field(fields, "blood_type", _extract_between(text, "กลุ่มเลือด*", ("เชื้อชาติ*",)), "ocr_form", "review")
    _put_field(fields, "race", _extract_between(text, "เชื้อชาติ*", ("สัญชาติ*",)), "ocr_form", "review")
    _put_field(fields, "nationality", _extract_between(text, "สัญชาติ*", ("ศาสนา*",)), "ocr_form", "review")
    _put_field(fields, "religion", _extract_between(text, "ศาสนา*", ("๒. ที่อยู่อาศัย", "2. ที่อยู่อาศัย")), "ocr_form", "review")
    _put_address_fields(fields, text, prefix="registered_address", start="ตามทะเบียนบ้าน")
    _put_address_fields(fields, text, prefix="current_address", start="ที่อยู่ปัจจุบัน")
    _put_field(fields, "commute_method", _checked_option(_extract_between(text, "การเดินทางมาโรงเรียน*", ("ระยะทางจากบ้านมา",))), "ocr_form", "review")
    _put_field(fields, "distance_paved_road_km", _number_text(_extract_between(text, "ถนนลาดยาง (กม.)", ("กม. รวมระยะเวลา",))), "ocr_form", "review")
    _put_field(fields, "commute_minutes", _number_text(_extract_between(text, "รวมระยะเวลาการเดินทางมาโรงเรียน (นาที)*", ("นาที",))), "ocr_form", "review")
    _put_field(fields, "weight_kg", _number_text(_extract_between(text, "น้ำหนัก*", ("กิโลกรัม",))), "ocr_form", "review")
    _put_field(fields, "height_cm", _number_text(_extract_between(text, "ส่วนสูง*", ("เซนติเมตร",))), "ocr_form", "review")
    _put_field(fields, "parents_marital_status", _checked_option(_extract_between(text, "สถานภาพสมรส", ("ข้อมูลพี่น้อง",))), "ocr_form", "review")
    _put_field(fields, "older_brothers", _number_text(_extract_between(text, "จำนวนพี่ชาย*", ("คน จำนวนน้องชาย*",))), "ocr_form", "review")
    _put_field(fields, "younger_brothers", _number_text(_extract_between(text, "จำนวนน้องชาย*", ("คน จำนวนพี่สาว*",))), "ocr_form", "review")
    _put_field(fields, "older_sisters", _number_text(_extract_between(text, "จำนวนพี่สาว*", ("คน จำนวนน้องสาว*",))), "ocr_form", "review")
    _put_field(fields, "younger_sisters", _number_text(_extract_between(text, "จำนวนน้องสาว*", ("คน จำนวนพี่น้องที่ศึกษาอยู่", "ข้อมูลบิดา"))), "ocr_form", "review")
    _put_field(fields, "siblings_studying_count", _first_digits(_extract_between(text, "จำนวนพี่น้องที่ศึกษาอยู่", ("คน นักเรียนเป็นบุตรคนที่", "นักเรียนเป็นบุตรคนที่"))), "ocr_form", "review")
    _put_field(fields, "child_order", _first_digits(_extract_between(text, "นักเรียนเป็นบุตรคนที่", ("ข้อมูลบิดา",))), "ocr_form", "review")
    _put_guardian_fields(fields, text, person="father", start="ข้อมูลบิดา", end="ข้อมูลมารดา")
    _put_guardian_fields(fields, text, person="mother", start="ข้อมูลมารดา", end="ข้อมูลผู้ปกครอง")
    _put_guardian_fields(fields, text, person="guardian", start="ข้อมูลผู้ปกครอง", end="**หมายเหตุ")
    _put_field(fields, "guardian_relationship", _extract_between(text, "ความเกี่ยวข้องของผู้ปกครองกับนักเรียน*", ("**หมายเหตุ",)), "ocr_form", "review")

    return (
        OcrFormRecord(
            record_id=f"ocr-{record_index}",
            citizen_id=citizen_id,
            citizen_id_valid=citizen_id_valid,
            student_no=student_no,
            prefix=prefix,
            first_name=first_name,
            last_name=last_name,
            full_name=full_name,
            name_key=name_key,
            fields=fields,
            source_path=str(path),
        ),
        warnings,
    )


def read_civil_registration_markdown_records(
    path: Path,
    *,
    record_index_start: int = 1,
) -> tuple[list[CivilRegistrationRecord], list[CurrentStudentsWarning]]:
    if path.suffix.lower() == ".json":
        return read_civil_registration_structured_json_records(path, record_index_start=record_index_start)

    text, _encoding = _decode_text(path)
    page_texts = _split_ocr_markdown_pages(text)
    records: list[CivilRegistrationRecord] = []
    warnings: list[CurrentStudentsWarning] = []
    for page_text in page_texts:
        if not _looks_like_civil_registration_text(page_text):
            continue
        record, page_warnings = _parse_civil_registration_text(
            path,
            page_text,
            record_index=record_index_start + len(records),
        )
        records.append(record)
        warnings.extend(page_warnings)
    if records:
        return records, warnings

    record, fallback_warnings = _parse_civil_registration_text(path, text, record_index=record_index_start)
    return [record], fallback_warnings


def read_civil_registration_markdown(
    path: Path,
    *,
    record_index: int = 1,
) -> tuple[CivilRegistrationRecord, list[CurrentStudentsWarning]]:
    records, warnings = read_civil_registration_markdown_records(path, record_index_start=record_index)
    return records[0], warnings


def read_civil_registration_structured_json_records(
    path: Path,
    *,
    record_index_start: int = 1,
) -> tuple[list[CivilRegistrationRecord], list[CurrentStudentsWarning]]:
    payload = _read_structured_ocr_payload(path)
    warnings: list[CurrentStudentsWarning] = []
    records: list[CivilRegistrationRecord] = []
    for record_payload in payload.records:
        if _structured_record_kind(record_payload, payload.document_type) == "dmc_form":
            continue
        record, record_warnings = _structured_payload_to_civil_registration_record(
            path,
            record_payload,
            record_index=record_index_start + len(records),
        )
        records.append(record)
        warnings.extend(record_warnings)
    return records, warnings


def _read_structured_ocr_payload(path: Path) -> _StructuredOcrPayload:
    try:
        raw_payload = json.loads(path.read_text(encoding="utf-8"))
        return _StructuredOcrPayload.model_validate(raw_payload)
    except (OSError, json.JSONDecodeError, ValidationError) as exc:
        raise DomainError("OCR_STRUCTURED_JSON_INVALID", "Structured OCR JSON could not be read.") from exc


def _structured_record_kind(record_payload: _StructuredOcrRecordPayload, document_type: str | None) -> str:
    value = _clean_text(record_payload.record_type or document_type).lower().replace("-", "_").replace(" ", "_")
    if value in {"civil_registration", "house_registration", "thai_house_registration", "tabien_baan"}:
        return "civil_registration"
    if value in {"dmc_form", "dmc_student_form", "student_history_form"}:
        return "dmc_form"
    return "unknown"


def _structured_payload_to_ocr_form_record(
    path: Path,
    record_payload: _StructuredOcrRecordPayload,
    *,
    record_index: int,
) -> tuple[OcrFormRecord, list[CurrentStudentsWarning]]:
    structured_fields = _flatten_structured_fields(record_payload.fields)
    fields: dict[str, CurrentStudentField] = {}
    for field_name, value in structured_fields.items():
        if field_name in FORM_JSON_FIELD_NAMES:
            _put_field(fields, field_name, value, "ocr_form", "review")

    citizen_id = _clean_citizen_id(_structured_text(structured_fields.get("citizen_id")))
    citizen_id_valid = citizen_id is not None and is_valid_thai_citizen_id(citizen_id)
    warnings = _structured_identity_warnings(
        path,
        citizen_id=citizen_id,
        citizen_id_valid=citizen_id_valid,
        source="ocr_form",
    )
    student_no = _structured_text(structured_fields.get("student_no"))
    prefix = _structured_text(structured_fields.get("prefix"))
    first_name = _structured_text(structured_fields.get("first_name"))
    last_name = _structured_text(structured_fields.get("last_name"))
    full_name = _full_name(prefix or "", first_name or "", last_name or "") if first_name or last_name else None
    name_key = normalized_name(first_name or "", last_name or "") if first_name or last_name else None

    return (
        OcrFormRecord(
            record_id=f"ocr-{record_index}",
            citizen_id=citizen_id,
            citizen_id_valid=citizen_id_valid,
            student_no=student_no,
            prefix=prefix,
            first_name=first_name,
            last_name=last_name,
            full_name=full_name,
            name_key=name_key,
            fields=fields,
            source_path=str(path),
            row_index=record_payload.page_start,
        ),
        warnings,
    )


def _structured_payload_to_civil_registration_record(
    path: Path,
    record_payload: _StructuredOcrRecordPayload,
    *,
    record_index: int,
) -> tuple[CivilRegistrationRecord, list[CurrentStudentsWarning]]:
    structured_fields = _flatten_structured_fields(record_payload.fields)
    fields: dict[str, CurrentStudentField] = {}
    for field_name, value in structured_fields.items():
        if field_name in FORM_JSON_FIELD_NAMES:
            _put_field(fields, field_name, value, "civil_registration", "high")

    citizen_id = _clean_citizen_id(_structured_text(structured_fields.get("citizen_id")))
    citizen_id_valid = citizen_id is not None and is_valid_thai_citizen_id(citizen_id)
    warnings = _structured_identity_warnings(
        path,
        citizen_id=citizen_id,
        citizen_id_valid=citizen_id_valid,
        source="civil_registration",
    )
    prefix = _structured_text(structured_fields.get("prefix"))
    first_name = _structured_text(structured_fields.get("first_name"))
    last_name = _structured_text(structured_fields.get("last_name"))
    full_name = _full_name(prefix or "", first_name or "", last_name or "") if first_name or last_name else None
    name_key = normalized_name(first_name or "", last_name or "") if first_name or last_name else None

    return (
        CivilRegistrationRecord(
            record_id=f"civil-{record_index}",
            citizen_id=citizen_id,
            citizen_id_valid=citizen_id_valid,
            prefix=prefix,
            first_name=first_name,
            last_name=last_name,
            full_name=full_name,
            name_key=name_key,
            fields=fields,
            source_path=str(path),
        ),
        warnings,
    )


def _structured_identity_warnings(
    path: Path,
    *,
    citizen_id: str | None,
    citizen_id_valid: bool,
    source: SourceType,
) -> list[CurrentStudentsWarning]:
    if citizen_id_valid:
        return []
    code = "OCR_CITIZEN_ID_INVALID_OR_MISSING" if source == "ocr_form" else "CIVIL_REGISTRATION_CITIZEN_ID_INVALID_OR_MISSING"
    return [
        CurrentStudentsWarning(
            code=code,
            message="Structured OCR citizen ID is missing or failed checksum validation.",
            source=source,
            source_path=str(path),
        )
    ]


def _flatten_structured_fields(fields: dict[str, Any]) -> dict[str, str | int | float | bool | None]:
    flattened: dict[str, str | int | float | bool | None] = {}

    def visit(prefix: str, value: Any) -> None:
        if isinstance(value, dict):
            for child_key, child_value in value.items():
                child_name = str(child_key).strip()
                if child_name:
                    visit(f"{prefix}.{child_name}" if prefix else child_name, child_value)
            return
        flattened[prefix] = _structured_scalar(value)

    for key, value in fields.items():
        name = str(key).strip()
        if name:
            visit(name, value)
    return flattened


def _structured_scalar(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, list):
        parts = [_clean_text(item) for item in value if _clean_text(item)]
        return ", ".join(parts) if parts else None
    return _clean_text(value)


def _structured_text(value: Any) -> str | None:
    if value is None:
        return None
    return _none_if_empty(value)


def _parse_civil_registration_text(
    path: Path,
    text: str,
    *,
    record_index: int,
) -> tuple[CivilRegistrationRecord, list[CurrentStudentsWarning]]:
    compact_text = _clean_text(text)
    warnings: list[CurrentStudentsWarning] = []
    fields: dict[str, CurrentStudentField] = {}

    house_id = _clean_house_id(_regex_first(compact_text, r"เลขรหัสประจำบ้าน[:\s]+([0-9][0-9\-\s]{8,})"))
    address_match = re.search(
        r"รายการที่อยู่\s+(?P<house_no>\S+)\s+หมู่ที่\s+(?P<moo>\S+)\s+ตำบล(?P<subdistrict>.+?)\s+อำเภอ(?P<district>.+?)\s+จังหวัด(?P<province>.+?)(?:\s+ชื่อหมู่บ้าน|\s+ลงชื่อ|\s+ประเภทบ้าน|\s+วันเดือนปี|$)",
        compact_text,
    )
    if address_match is not None:
        _put_field(fields, "registered_address.house_no", address_match.group("house_no"), "civil_registration", "high")
        _put_field(fields, "registered_address.moo", address_match.group("moo"), "civil_registration", "high")
        _put_field(fields, "registered_address.subdistrict", address_match.group("subdistrict"), "civil_registration", "high")
        _put_field(fields, "registered_address.district", address_match.group("district"), "civil_registration", "high")
        _put_field(fields, "registered_address.province", address_match.group("province"), "civil_registration", "high")
    _put_field(fields, "registered_address.house_id", house_id, "civil_registration", "high")

    name_match = re.search(
        r"รายการบุคคลในบ้าน.*?(?:เพศ[:\s]*(?P<sex_before>ชาย|หญิง)\s+)?ชื่อ[:\s]+(?P<full_name>.+?)\s+สัญชาติ[:\s]+(?P<nationality>\S+)(?:\s+เพศ[:\s]*(?P<sex_after>\S+))?",
        compact_text,
    )
    prefix: str | None = None
    first_name: str | None = None
    last_name: str | None = None
    full_name: str | None = None
    name_key: str | None = None
    if name_match is not None:
        prefix, first_name, last_name = _split_civil_name(name_match.group("full_name"))
        full_name = _full_name(prefix or "", first_name or "", last_name or "") if first_name or last_name else None
        name_key = normalized_name(first_name or "", last_name or "") if first_name or last_name else None
        _put_field(fields, "prefix", prefix, "civil_registration", "high")
        _put_field(fields, "first_name", first_name, "civil_registration", "high")
        _put_field(fields, "last_name", last_name, "civil_registration", "high")
        _put_field(fields, "nationality", name_match.group("nationality"), "civil_registration", "high")
        _put_field(fields, "race", name_match.group("nationality"), "civil_registration", "high")
        _put_field(fields, "sex", name_match.group("sex_after") or name_match.group("sex_before"), "civil_registration", "high")

    citizen_id = _clean_citizen_id(
        _regex_first(compact_text, r"เลขประจำตัวประชาชน[:\s]+([0-9][0-9\-\s]{10,})(?=\s+สถานภาพ)")
    )
    citizen_id_valid = citizen_id is not None and is_valid_thai_citizen_id(citizen_id)
    if not citizen_id_valid:
        warnings.append(
            CurrentStudentsWarning(
                code="CIVIL_REGISTRATION_CITIZEN_ID_INVALID_OR_MISSING",
                message="Civil registration student citizen ID is missing or failed checksum validation.",
                source="civil_registration",
                source_path=str(path),
            )
        )
    _put_field(fields, "citizen_id", citizen_id, "civil_registration", "high")
    _put_field(fields, "birth_date", _regex_first(compact_text, r"เกิดเมื่อ\s+(.+?)(?:\s+มาจาก|\s+ลงชื่อ|$)"), "civil_registration", "high")
    _put_civil_parent_fields(fields, compact_text, person="mother", marker="มารดาผู้ให้กำเนิด")
    _put_civil_parent_fields(fields, compact_text, person="father", marker="บิดาผู้ให้กำเนิด")

    if citizen_id is None and name_key is None:
        warnings.append(
            CurrentStudentsWarning(
                code="CIVIL_REGISTRATION_NO_MATCH_KEY",
                message="Civil registration OCR did not contain a usable citizen ID or student name.",
                source="civil_registration",
                source_path=str(path),
            )
        )

    return (
        CivilRegistrationRecord(
            record_id=f"civil-{record_index}",
            citizen_id=citizen_id,
            citizen_id_valid=citizen_id_valid,
            prefix=prefix,
            first_name=first_name,
            last_name=last_name,
            full_name=full_name,
            name_key=name_key,
            fields=fields,
            source_path=str(path),
        ),
        warnings,
    )


def normalized_name(first_name: str | None, last_name: str | None = None) -> str:
    value = _clean_text(f"{first_name or ''}{last_name or ''}").replace(".", "")
    for prefix in TITLE_PREFIXES:
        compact_prefix = prefix.replace(".", "")
        if value.startswith(prefix):
            value = value[len(prefix) :].strip()
        if value.startswith(compact_prefix):
            value = value[len(compact_prefix) :].strip()
    return re.sub(r"\s+", "", value)


def is_valid_thai_citizen_id(citizen_id: str) -> bool:
    if not re.fullmatch(r"\d{13}", citizen_id):
        return False
    checksum = sum(int(citizen_id[index]) * (13 - index) for index in range(12))
    check_digit = (11 - (checksum % 11)) % 10
    return check_digit == int(citizen_id[-1])


def _build_canonical_records(
    *,
    roster: list[RosterStudent],
    scans: list[ThaiIdScanRecord],
    ocr_records: list[OcrFormRecord],
    civil_records: list[CivilRegistrationRecord],
    operation_type: OperationType,
    fuzzy_match_threshold: float,
    confirmed_ocr_matches: dict[str, str],
) -> tuple[list[CanonicalStudentRecord], int, int]:
    records_by_roster_key = {_roster_key(student): _base_record(student, operation_type) for student in roster}
    roster_by_name: dict[str, list[RosterStudent]] = defaultdict(list)
    roster_by_student_no: dict[str, RosterStudent] = {}
    for student in roster:
        roster_by_name[student.name_key].append(student)
        roster_by_student_no[student.student_no] = student

    scan_by_name: dict[str, list[ThaiIdScanRecord]] = defaultdict(list)
    for scan in scans:
        scan_by_name[scan.name_key].append(scan)

    duplicate_citizen_ids = _duplicates(scan.citizen_id for scan in scans if scan.citizen_id)
    duplicate_scan_names = _duplicates(scan.name_key for scan in scans if scan.name_key)
    consumed_scan_rows: set[int] = set()

    for student in roster:
        record = records_by_roster_key[_roster_key(student)]
        matching_scans = scan_by_name.get(student.name_key, [])
        if len(matching_scans) == 1:
            scan = matching_scans[0]
            consumed_scan_rows.add(scan.row_index)
            _attach_scan(record, scan)
            if scan.citizen_id in duplicate_citizen_ids or scan.name_key in duplicate_scan_names:
                _set_status(record, "duplicate", "duplicate_thai_id_scan")
            elif not scan.citizen_id_valid:
                _set_status(record, "invalid_id", "invalid_or_masked_citizen_id")
            else:
                record.match_status = "auto_matched"
                record.match_score = 1.0
                _remove_reason(record, "missing_thai_id_scan")
        elif len(matching_scans) > 1:
            for scan in matching_scans:
                consumed_scan_rows.add(scan.row_index)
                _attach_scan(record, scan)
            _set_status(record, "duplicate", "multiple_thai_id_scans_for_roster_name")

    records = list(records_by_roster_key.values())
    for scan in scans:
        if scan.row_index in consumed_scan_rows:
            continue
        suggestions = _best_suggestions(scan.name_key, roster, limit=3)
        best_score = suggestions[0].score if suggestions else None
        if scan.citizen_id_valid and best_score is not None and best_score >= fuzzy_match_threshold:
            suggested_student = roster_by_student_no.get(suggestions[0].student_no)
            if suggested_student is not None:
                suggested_record = records_by_roster_key[_roster_key(suggested_student)]
                if not _has_source(suggested_record, "thai_id_scan"):
                    consumed_scan_rows.add(scan.row_index)
                    _attach_scan(suggested_record, scan)
                    suggested_record.match_status = "needs_review"
                    suggested_record.match_score = best_score
                    suggested_record.suggestions = suggestions
                    _add_reason(suggested_record, "fuzzy_roster_match_candidate")
                    _remove_reason(suggested_record, "missing_thai_id_scan")
                    continue
        status: MatchStatus
        reason: str
        if not scan.citizen_id_valid:
            status = "invalid_id"
            reason = "invalid_or_masked_citizen_id"
        elif best_score is not None and best_score >= fuzzy_match_threshold:
            status = "needs_review"
            reason = "fuzzy_roster_match_candidate"
        else:
            status = "new_or_transfer_candidate"
            reason = "not_found_in_roster"
        records.append(_scan_only_record(scan, operation_type, status, reason, suggestions))

    records_by_citizen_id = {record.citizen_id: record for record in records if record.citizen_id}
    records_by_student_no = {record.student_no: record for record in records if record.student_no}
    records_by_name = {normalized_name(record.first_name, record.last_name): record for record in records if record.first_name or record.last_name}

    for civil_record in civil_records:
        target = None
        if civil_record.citizen_id:
            target = records_by_citizen_id.get(civil_record.citizen_id)
        if target is None and civil_record.name_key:
            target = records_by_name.get(civil_record.name_key)
        if target is None:
            target = _civil_registration_only_record(civil_record, operation_type)
            records.append(target)
        else:
            _attach_civil_registration(target, civil_record)
        if target.citizen_id:
            records_by_citizen_id[target.citizen_id] = target
        if target.student_no:
            records_by_student_no[target.student_no] = target
        if target.first_name or target.last_name:
            records_by_name[normalized_name(target.first_name, target.last_name)] = target

    ocr_attached = 0
    ocr_unmatched = 0
    for ocr_record in ocr_records:
        target = None
        suggestions = _best_ocr_suggestions(ocr_record, roster, limit=3)
        confirmed_student_no = confirmed_ocr_matches.get(_ocr_record_key(ocr_record)) or confirmed_ocr_matches.get(
            ocr_record.record_id
        )
        if confirmed_student_no:
            confirmed_student = roster_by_student_no.get(confirmed_student_no)
            if confirmed_student is not None:
                target = records_by_roster_key[_roster_key(confirmed_student)]
                target.match_score = _suggestion_score_for_student_no(suggestions, confirmed_student_no)
                target.suggestions = suggestions
                _add_reason(target, "user_confirmed_ocr_roster_match")
        if target is None and ocr_record.citizen_id:
            target = records_by_citizen_id.get(ocr_record.citizen_id)
        if target is None and ocr_record.student_no:
            target = records_by_student_no.get(ocr_record.student_no)
        if target is None and ocr_record.name_key:
            target = records_by_name.get(ocr_record.name_key)
        if target is None:
            best_score = suggestions[0].score if suggestions else None
            if best_score is not None and best_score >= OCR_FUZZY_MATCH_THRESHOLD:
                ocr_status: MatchStatus = "needs_review"
                reason = "ocr_fuzzy_roster_match_candidate"
            else:
                ocr_status = "new_or_transfer_candidate" if ocr_record.citizen_id_valid else "invalid_id"
                reason = "ocr_form_not_linked_to_roster_or_thai_id_scan"
            records.append(_ocr_only_record(ocr_record, operation_type, ocr_status, reason, suggestions))
            ocr_unmatched += 1
        else:
            _attach_ocr(target, ocr_record)
            if suggestions and not target.suggestions:
                target.suggestions = suggestions
            if target.citizen_id:
                records_by_citizen_id[target.citizen_id] = target
            ocr_attached += 1

    _promote_registered_address_to_current_when_matching(records)
    _fill_missing_mother_last_name_from_father(records)
    _fill_guardian_from_parent_relationship(records)
    return records, ocr_attached, ocr_unmatched


def _write_import_workbook(
    output_path: Path,
    *,
    records: list[CanonicalStudentRecord],
    default_school_year: int | None = None,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "current_students_import"
    labels = [label for _field_name, label, _required in IMPORT_FIELD_DEFINITIONS]
    field_names = [field_name for field_name, _label, _required in IMPORT_FIELD_DEFINITIONS]
    sheet.append(labels)
    sheet.append(field_names)
    sheet.freeze_panes = "A3"
    header_fill = PatternFill("solid", fgColor="E2E8F0")
    key_fill = PatternFill("solid", fgColor="F8FAFC")
    for column_index, (_field_name, _label, required) in enumerate(IMPORT_FIELD_DEFINITIONS, start=1):
        header_cell = sheet.cell(row=IMPORT_HEADER_ROW, column=column_index)
        key_cell = sheet.cell(row=IMPORT_FIELD_KEY_ROW, column=column_index)
        header_cell.font = Font(bold=True, color="991B1B" if required else "111827")
        header_cell.fill = header_fill
        key_cell.font = Font(color="64748B")
        key_cell.fill = key_fill
        sheet.column_dimensions[header_cell.column_letter].width = 18

    for record in records:
        sheet.append([_record_import_value(record, field_name, default_school_year) for field_name in field_names])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _record_import_value(
    record: CanonicalStudentRecord,
    field_name: str,
    default_school_year: int | None,
) -> str | int | float | bool | None:
    if field_name == "operation_type":
        return record.operation_type
    if field_name == "school_year":
        field = record.dmc_fields.get(field_name)
        return field.value if field is not None and field.value is not None else default_school_year
    if field_name == "student_no":
        return record.student_no
    if field_name == "citizen_id":
        return record.citizen_id
    if field_name == "grade":
        return record.grade
    if field_name == "room":
        return record.room
    if field_name == "seat_no":
        return record.seat_no
    if field_name == "prefix":
        return record.prefix
    if field_name == "first_name":
        return record.first_name
    if field_name == "last_name":
        return record.last_name
    field = record.dmc_fields.get(field_name)
    if field is None:
        return None
    return field.value


def _json_record(
    record: CanonicalStudentRecord,
    *,
    default_school_year: int | None,
    default_admission_date: str | None = None,
) -> DmcFormJsonRecord:
    fields = {
        field_name: _record_import_value(record, field_name, default_school_year)
        for field_name in FORM_JSON_FIELD_NAMES
    }
    return DmcFormJsonRecord(
        record_id=record.record_id,
        match_status=record.match_status,
        student_no=record.student_no,
        citizen_id=record.citizen_id,
        grade=record.grade,
        room=record.room,
        seat_no=record.seat_no,
        prefix=record.prefix,
        first_name=record.first_name,
        last_name=record.last_name,
        full_name=_export_full_name(record),
        match_score=record.match_score,
        review_reasons=record.review_reasons,
        suggestions=record.suggestions,
        sources=record.sources,
        fields=fields,
        dmc_form_values=_dmc_form_values(
            record,
            default_school_year=default_school_year,
            default_admission_date=default_admission_date,
        ),
        field_details={
            field_name: _json_field_detail(record, field_name, default_school_year=default_school_year)
            for field_name in FORM_JSON_FIELD_NAMES
        },
    )


def _dmc_form_values(
    record: CanonicalStudentRecord,
    *,
    default_school_year: int | None,
    default_admission_date: str | None = None,
) -> dict[str, DmcFormValue]:
    values: dict[str, DmcFormValue] = {}

    school_year = _dmc_field_text(record, "school_year", default_school_year)
    _dmc_put_value(values, "educationYear", school_year)
    _dmc_put_value(values, "admissionDate", _dmc_admission_date_text(default_admission_date))
    _dmc_put_value(values, "studentNo", _dmc_field_text(record, "student_no", default_school_year))
    _dmc_put_value(values, "levelDtlCode", _dmc_level_code(record, default_school_year=default_school_year))
    _dmc_put_value(values, "classroom", _dmc_field_text(record, "room", default_school_year))

    citizen_id = _clean_citizen_id(_dmc_field_text(record, "citizen_id", default_school_year))
    if citizen_id:
        _dmc_put_value(values, "cifNo", citizen_id)
        _dmc_put_value(values, "cifNoChk", citizen_id)
        _dmc_put_value(values, "cifType", "I")

    title_code = _dmc_title_code(_dmc_field_text(record, "prefix", default_school_year))
    _dmc_put_value(values, "titleCode", title_code)
    gender_code = _dmc_gender_code(_dmc_field_text(record, "sex", default_school_year))
    _dmc_put_value(values, "genderCode", gender_code or DMC_GENDER_BY_TITLE_CODE.get(title_code or ""))
    _dmc_put_value(values, "firstNameTh", _dmc_field_text(record, "first_name", default_school_year))
    _dmc_put_value(values, "lastNameTh", _dmc_field_text(record, "last_name", default_school_year))
    _dmc_put_value(values, "middleNameTh", "")
    _dmc_put_value(values, "firstNameEn", _dmc_field_text(record, "first_name_en", default_school_year))
    _dmc_put_value(values, "lastNameEn", _dmc_field_text(record, "last_name_en", default_school_year))
    _dmc_put_value(values, "middleNameEn", "")
    _dmc_put_value(values, "birthDate", _dmc_date_text(_dmc_field_text(record, "birth_date", default_school_year)))
    _dmc_put_value(
        values,
        "birthProvinceCode",
        _dmc_province_code(_dmc_field_text(record, "birth_province", default_school_year)),
    )
    _dmc_put_value(values, "bloodCode", _dmc_blood_code(_dmc_field_text(record, "blood_type", default_school_year)))
    _dmc_put_value(
        values,
        "nationCode",
        _dmc_lookup_code(_dmc_field_text(record, "nationality", default_school_year), DMC_NATION_CODES),
    )
    race = _dmc_field_text(record, "race", default_school_year)
    _dmc_put_value(
        values,
        "raceCode",
        DMC_DEFAULT_MISSING_RACE_CODE if race is None else _dmc_lookup_code(race, DMC_RACE_CODES),
    )
    _dmc_put_value(
        values,
        "religionCode",
        _dmc_lookup_code(_dmc_field_text(record, "religion", default_school_year), DMC_RELIGION_CODES),
    )
    _dmc_put_value(values, "studentTypeCode", "01")
    _dmc_put_value(values, "objectStatus", "Y")

    _dmc_put_address_values(
        values,
        record,
        source_prefix="registered_address",
        form_prefix="ps",
        default_school_year=default_school_year,
    )
    _dmc_put_address_values(
        values,
        record,
        source_prefix="current_address",
        form_prefix="",
        fallback_source_prefix="registered_address",
        default_school_year=default_school_year,
    )

    _dmc_put_value(
        values,
        "marriageStatusCode",
        _dmc_marriage_status_code(_dmc_field_text(record, "parents_marital_status", default_school_year)),
    )
    _dmc_put_value(
        values,
        "numOfOlderBrothers",
        _dmc_integer_text(
            record,
            "older_brothers",
            default_school_year,
            default=DMC_DEFAULT_MISSING_SIBLING_COUNT,
        ),
    )
    _dmc_put_value(
        values,
        "numOfYoungerBrothers",
        _dmc_integer_text(
            record,
            "younger_brothers",
            default_school_year,
            default=DMC_DEFAULT_MISSING_SIBLING_COUNT,
        ),
    )
    _dmc_put_value(
        values,
        "numOfOlderSisters",
        _dmc_integer_text(
            record,
            "older_sisters",
            default_school_year,
            default=DMC_DEFAULT_MISSING_SIBLING_COUNT,
        ),
    )
    _dmc_put_value(
        values,
        "numOfYoungerSisters",
        _dmc_integer_text(
            record,
            "younger_sisters",
            default_school_year,
            default=DMC_DEFAULT_MISSING_SIBLING_COUNT,
        ),
    )
    _dmc_put_value(
        values,
        "numOfStudyingSiblings",
        _dmc_integer_text(
            record,
            "siblings_studying_count",
            default_school_year,
            default=DMC_DEFAULT_MISSING_SIBLING_COUNT,
        ),
    )
    _dmc_put_value(values, "childIndex", _dmc_integer_text(record, "child_order", default_school_year))

    relation_code = _dmc_code_from_items(
        _dmc_field_text(record, "guardian_relationship", default_school_year),
        DMC_PARENT_RELATION_CODES,
    )
    _dmc_put_person_values(
        values,
        record,
        person="father",
        form_person="father",
        default_title_code="003",
        default_school_year=default_school_year,
    )
    _dmc_put_person_values(
        values,
        record,
        person="mother",
        form_person="mother",
        default_title_code="004",
        default_school_year=default_school_year,
    )
    _dmc_put_person_values(
        values,
        record,
        person="guardian",
        form_person="parent",
        default_title_code={"01": "003", "02": "004"}.get(relation_code or ""),
        default_school_year=default_school_year,
    )
    _dmc_put_value(values, "parentFamilyRelationCode", relation_code)

    commute_method = _dmc_field_text(record, "commute_method", default_school_year)
    _dmc_put_value(
        values,
        "journeyTypeCode",
        DMC_DEFAULT_MISSING_JOURNEY_TYPE_CODE
        if commute_method is None
        else _dmc_code_from_items(commute_method, DMC_JOURNEY_TYPE_CODES),
    )
    _dmc_put_value(
        values,
        "timeDt",
        _dmc_decimal_text(
            _dmc_field_text(record, "commute_minutes", default_school_year),
            default=DMC_DEFAULT_MISSING_COMMUTE_MINUTES,
        ),
    )
    _dmc_put_value(values, "waterDt", "0.0")
    _dmc_put_value(values, "rockDt", "0.0")
    _dmc_put_value(
        values,
        "rubberDt",
        _dmc_decimal_text(_dmc_field_text(record, "distance_paved_road_km", default_school_year), multiplier=1000),
    )
    _dmc_put_value(values, "weight", _dmc_decimal_text(_dmc_field_text(record, "weight_kg", default_school_year)))
    _dmc_put_value(values, "height", _dmc_decimal_text(_dmc_field_text(record, "height_cm", default_school_year)))

    return values


def _dmc_put_value(values: dict[str, DmcFormValue], name: str, value: DmcFormValue) -> None:
    if value is None:
        return
    if isinstance(value, str):
        text = _clean_text(value)
        if text in {"........................", "..................."}:
            return
    if isinstance(value, list) and not value:
        return
    values[name] = value


def _dmc_field_text(
    record: CanonicalStudentRecord,
    field_name: str,
    default_school_year: int | None,
) -> str | None:
    value = _record_import_value(record, field_name, default_school_year)
    if _is_missing_import_value(value):
        return None
    return _clean_text(value)


def _dmc_level_code(record: CanonicalStudentRecord, *, default_school_year: int | None) -> str | None:
    grade_value = _record_import_value(record, "grade", default_school_year)
    if isinstance(grade_value, int):
        return DMC_TRANSFER_IN_LEVEL_CODES.get(grade_value)
    grade = _first_digits(_clean_text(grade_value))
    return DMC_TRANSFER_IN_LEVEL_CODES.get(int(grade)) if grade else None


def _dmc_title_code(value: str | None) -> str | None:
    text = _clean_text(value)
    if not text:
        return None
    compact = _compact_thai_lookup_key(text)
    for title, code in DMC_TITLE_CODE_ITEMS:
        if compact == _compact_thai_lookup_key(title):
            return code
    return None


def _dmc_split_prefixed_first_name(value: str | None, default_title_code: str | None) -> tuple[str | None, str | None]:
    text = _clean_text(value)
    if not text:
        return default_title_code, None
    compact = _compact_thai_lookup_key(text)
    for title, code in sorted(DMC_TITLE_CODE_ITEMS, key=lambda item: len(_compact_thai_lookup_key(item[0])), reverse=True):
        title_compact = _compact_thai_lookup_key(title)
        if not compact.startswith(title_compact):
            continue
        if text.startswith(title):
            return code, _none_if_empty(text[len(title) :])
        title_without_dot = title.replace(".", "")
        if text.startswith(title_without_dot):
            return code, _none_if_empty(text[len(title_without_dot) :])
    return default_title_code, text


def _dmc_gender_code(value: str | None) -> str | None:
    text = _clean_text(value)
    if not text:
        return None
    for label, code in DMC_GENDER_CODES.items():
        if label in text:
            return code
    return None


def _dmc_lookup_code(value: str | None, mapping: dict[str, str]) -> str | None:
    text = _dmc_area_lookup_key(value)
    if not text:
        return None
    return mapping.get(text)


def _dmc_code_from_items(value: str | None, items: Sequence[tuple[str, str]]) -> str | None:
    text = _clean_text(value)
    if not text:
        return None
    lowered = text.lower()
    for marker, code in items:
        if marker.lower() in lowered:
            return code
    return None


def _dmc_marriage_status_code(value: str | None) -> str | None:
    text = _clean_text(value)
    if not text:
        return DMC_DEFAULT_MISSING_MARRIAGE_STATUS_CODE
    return _dmc_code_from_items(text, DMC_MARRIAGE_STATUS_CODES)


def _dmc_blood_code(value: str | None) -> str | None:
    text = _clean_text(value).replace(" ", "")
    if not text:
        return None
    normalized = text.upper().replace("RH", "Rh")
    return DMC_BLOOD_CODES.get(normalized) or DMC_BLOOD_CODES.get(text)


def _dmc_province_code(value: str | None) -> str | None:
    text = _dmc_area_lookup_key(value)
    if not text:
        return None
    digits = "".join(re.findall(r"\d", text))
    if len(digits) == 8:
        return digits
    if len(digits) == 2:
        return f"{digits}000000"
    admin_codes = _thai_admin_code_index()
    return admin_codes.province_codes.get(_admin_area_lookup_key(text)) or DMC_PROVINCE_CODES.get(text)


def _dmc_district_code(value: str | None, province_code: str | None) -> str | None:
    text = _dmc_area_lookup_key(value)
    if not text:
        return None
    digits = "".join(re.findall(r"\d", text))
    if len(digits) == 8:
        return digits
    admin_codes = _thai_admin_code_index()
    if province_code is not None:
        code = admin_codes.district_codes.get((province_code, _admin_area_lookup_key(text)))
        if code is not None:
            return code
    if province_code == "81000000":
        return DMC_KRABI_DISTRICT_CODES.get(text)
    return None


def _dmc_subdistrict_code(value: str | None, province_code: str | None, district_code: str | None) -> str | None:
    text = _dmc_area_lookup_key(value)
    if not text:
        return None
    digits = "".join(re.findall(r"\d", text))
    if len(digits) == 8:
        return digits
    admin_codes = _thai_admin_code_index()
    lookup_key = _admin_area_lookup_key(text)
    if province_code is not None and district_code is not None:
        code = admin_codes.subdistrict_codes.get((province_code, district_code, lookup_key))
        if code is not None:
            return code
    if province_code is not None:
        code = admin_codes.unique_subdistrict_codes_by_province.get((province_code, lookup_key))
        if code is not None:
            return code
    if province_code == "81000000":
        return DMC_KRABI_SUBDISTRICT_CODES.get(text)
    return None


def _dmc_area_lookup_key(value: str | None) -> str:
    text = _clean_text(value)
    text = re.sub(r"^(จังหวัด|จ\.|อำเภอ|อ\.|เขต|ตำบล|ต\.|แขวง)\s*", "", text)
    return text.strip()


def _compact_thai_lookup_key(value: str) -> str:
    return re.sub(r"[\s.]+", "", _clean_text(value))


def _admin_area_lookup_key(value: str | None) -> str:
    return _compact_thai_lookup_key(_dmc_area_lookup_key(value)).lower()


def _thai_admin_code_index() -> ThaiAdminCodeIndex:
    global _THAI_ADMIN_CODE_INDEX
    if _THAI_ADMIN_CODE_INDEX is not None:
        return _THAI_ADMIN_CODE_INDEX

    province_codes: dict[str, str] = {}
    district_codes: dict[tuple[str, str], str] = {}
    subdistrict_codes: dict[tuple[str, str, str], str] = {}
    unique_subdistrict_codes_by_province: dict[tuple[str, str], str] = {}
    ambiguous_subdistricts: set[tuple[str, str]] = set()
    data = json.loads(THAI_ADMIN_CODES_PATH.read_text(encoding="utf-8"))
    records = data.get("records", []) if isinstance(data, dict) else []
    for row in records:
        if not isinstance(row, dict):
            continue
        province = _clean_text(row.get("province"))
        district = _clean_text(row.get("district"))
        subdistrict = _clean_text(row.get("subdistrict"))
        province_code = _clean_text(row.get("province_code"))
        district_code = _clean_text(row.get("district_code"))
        subdistrict_code = _clean_text(row.get("subdistrict_code"))
        if not (
            province
            and district
            and subdistrict
            and re.fullmatch(r"\d{8}", province_code)
            and re.fullmatch(r"\d{8}", district_code)
            and re.fullmatch(r"\d{8}", subdistrict_code)
        ):
            continue

        province_key = _admin_area_lookup_key(province)
        district_key = _admin_area_lookup_key(district)
        subdistrict_key = _admin_area_lookup_key(subdistrict)
        province_codes[province_key] = province_code
        district_codes[(province_code, district_key)] = district_code
        subdistrict_codes[(province_code, district_code, subdistrict_key)] = subdistrict_code

        unique_key = (province_code, subdistrict_key)
        if unique_key in ambiguous_subdistricts:
            continue
        existing_subdistrict_code = unique_subdistrict_codes_by_province.get(unique_key)
        if existing_subdistrict_code is None:
            unique_subdistrict_codes_by_province[unique_key] = subdistrict_code
        elif existing_subdistrict_code != subdistrict_code:
            unique_subdistrict_codes_by_province.pop(unique_key, None)
            ambiguous_subdistricts.add(unique_key)

    _THAI_ADMIN_CODE_INDEX = ThaiAdminCodeIndex(
        province_codes=province_codes,
        district_codes=district_codes,
        subdistrict_codes=subdistrict_codes,
        unique_subdistrict_codes_by_province=unique_subdistrict_codes_by_province,
    )
    return _THAI_ADMIN_CODE_INDEX


def _dmc_date_text(value: str | None) -> str | None:
    text = _clean_text(value)
    if not text:
        return None
    slash_match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
    if slash_match:
        day, month, year = slash_match.groups()
        return f"{int(day):02d}/{int(month):02d}/{year}"
    thai_match = re.search(r"(\d{1,2})\s+([ก-๙.]+)\s+(\d{4})", text)
    if thai_match is None:
        return None
    day, month_text, year = thai_match.groups()
    month = THAI_MONTH_CODES.get(month_text)
    if month is None:
        return None
    return f"{int(day):02d}/{month}/{year}"


def _dmc_admission_date_text(value: str | None) -> str | None:
    text = _clean_text(value)
    if not text:
        return None
    iso_match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if iso_match:
        year, month, day = (int(part) for part in iso_match.groups())
        if year < 2400:
            year += 543
        return f"{day:02d}/{month:02d}/{year}"
    return _dmc_date_text(text)


def _dmc_integer_text(
    record: CanonicalStudentRecord,
    field_name: str,
    default_school_year: int | None,
    *,
    default: str | None = None,
) -> str | None:
    value = _dmc_field_text(record, field_name, default_school_year)
    return _first_digits(value) or default


def _dmc_decimal_text(value: str | None, *, multiplier: float = 1.0, default: str | None = None) -> str | None:
    text = _clean_text(value)
    if not text:
        return default
    numbers = [float(match) for match in re.findall(r"\d+(?:\.\d+)?", text)]
    if not numbers:
        return default
    return f"{numbers[0] * multiplier:.1f}"


def _dmc_salary_text(value: str | None) -> str:
    text = _clean_text(value)
    numbers = [float(match.replace(",", "")) for match in re.findall(r"\d[\d,]*(?:\.\d+)?", text)]
    if not numbers:
        return "0.0"
    salary = sum(numbers[:2]) / min(len(numbers), 2)
    return f"{salary:.1f}"


def _dmc_card_type(value: str | None, *, has_citizen_id: bool) -> str:
    code = _dmc_code_from_items(value, DMC_CARD_TYPE_CODES)
    if code is not None:
        return code
    return "I" if has_citizen_id else "O"


def _dmc_put_address_values(
    values: dict[str, DmcFormValue],
    record: CanonicalStudentRecord,
    *,
    source_prefix: Literal["registered_address", "current_address"],
    form_prefix: Literal["ps", ""],
    default_school_year: int | None,
    fallback_source_prefix: Literal["registered_address", "current_address"] | None = None,
) -> None:
    def text(suffix: str) -> str | None:
        field_value = _dmc_field_text(record, f"{source_prefix}.{suffix}", default_school_year)
        if field_value is None and fallback_source_prefix is not None:
            return _dmc_field_text(record, f"{fallback_source_prefix}.{suffix}", default_school_year)
        return field_value

    house_id = _dmc_house_id(text("house_id"))
    house_no = text("house_no")
    moo = _first_digits(text("moo"))
    road = text("road")
    province_code = _dmc_province_code(text("province"))
    district_code = _dmc_district_code(text("district"), province_code)
    subdistrict_code = _dmc_subdistrict_code(text("subdistrict"), province_code, district_code)
    if subdistrict_code is not None:
        inferred_district_code = f"{subdistrict_code[:4]}0000"
        if district_code is None or district_code[:4] != subdistrict_code[:4]:
            district_code = inferred_district_code
    if province_code is None and district_code is not None:
        province_code = f"{district_code[:2]}000000"
    postal_code = _first_digits(text("postal_code"))
    has_address = any([house_id, house_no, moo, road, province_code, district_code, subdistrict_code, postal_code])
    if not has_address:
        return

    if form_prefix == "ps":
        names = {
            "house_id": "psHomeIdNo",
            "house_no": "psHomeNo",
            "moo": "psMoo",
            "road": "psStreet",
            "province": "psProvinceCode",
            "district": "psAmphurCode",
            "subdistrict": "psTumbolCode",
            "postal_code": "psPostalCode",
            "tel": "psTelNo",
        }
    else:
        names = {
            "house_id": "homeIdNo",
            "house_no": "homeNo",
            "moo": "moo",
            "road": "street",
            "province": "provinceCode",
            "district": "amphurCode",
            "subdistrict": "tumbolCode",
            "postal_code": "postalCode",
            "tel": "telNo",
        }

    _dmc_put_value(values, names["house_id"], house_id)
    _dmc_put_value(values, names["house_no"], house_no)
    _dmc_put_value(values, names["moo"], moo or "0")
    _dmc_put_value(values, names["road"], road or "-")
    _dmc_put_value(values, names["province"], province_code)
    _dmc_put_value(values, names["district"], district_code)
    _dmc_put_value(values, names["subdistrict"], subdistrict_code)
    _dmc_put_value(values, names["postal_code"], postal_code)
    _dmc_put_value(values, names["tel"], "-")


def _dmc_house_id(value: str | None) -> str | None:
    digits = "".join(re.findall(r"\d", _clean_text(value)))
    return digits if digits else None


def _dmc_put_person_values(
    values: dict[str, DmcFormValue],
    record: CanonicalStudentRecord,
    *,
    person: Literal["father", "mother", "guardian"],
    form_person: Literal["father", "mother", "parent"],
    default_title_code: str | None,
    default_school_year: int | None,
) -> None:
    citizen_id = _clean_citizen_id(_dmc_field_text(record, f"{person}.citizen_id", default_school_year))
    title_code, first_name = _dmc_split_prefixed_first_name(
        _dmc_field_text(record, f"{person}.first_name", default_school_year),
        default_title_code,
    )
    last_name = _dmc_field_text(record, f"{person}.last_name", default_school_year)
    card_type = _dmc_field_text(record, f"{person}.card_type", default_school_year)
    blood_type = _dmc_blood_code(_dmc_field_text(record, f"{person}.blood_type", default_school_year))
    occupation_text = _dmc_field_text(record, f"{person}.occupation", default_school_year)
    occupation = _dmc_code_from_items(occupation_text, DMC_OCCUPATION_CODES)
    income = _dmc_field_text(record, f"{person}.income_text", default_school_year)
    phone = _dmc_field_text(record, f"{person}.phone", default_school_year)
    has_person_without_occupation = any([citizen_id, first_name, last_name, card_type, blood_type, income, phone])
    if occupation is None:
        occupation = DMC_OTHER_OCCUPATION_CODE if _clean_text(occupation_text) else None
    if occupation is None and has_person_without_occupation:
        occupation = DMC_DEFAULT_MISSING_OCCUPATION_CODE
    has_person = any([has_person_without_occupation, occupation])
    if not has_person:
        return

    prefix = form_person
    _dmc_put_value(values, f"{prefix}CifNo", citizen_id or "-")
    _dmc_put_value(values, f"{prefix}CifType", _dmc_card_type(card_type, has_citizen_id=citizen_id is not None))
    _dmc_put_value(values, f"{prefix}TitleCode", title_code)
    _dmc_put_value(values, f"{prefix}OccupationCode", occupation)
    _dmc_put_value(values, f"{prefix}FirstNameTh", first_name)
    _dmc_put_value(values, f"{prefix}LastNameTh", last_name)
    _dmc_put_value(values, f"{prefix}MiddleNameTh", "")
    _dmc_put_value(values, f"{prefix}BloodCode", blood_type)
    _dmc_put_value(values, f"{prefix}Salary", _dmc_salary_text(income))
    _dmc_put_value(values, f"{prefix}TelNo", phone or "-")


def _json_field_detail(
    record: CanonicalStudentRecord,
    field_name: str,
    *,
    default_school_year: int | None,
) -> CurrentStudentField:
    field = record.dmc_fields.get(field_name)
    if field is not None:
        return field

    value = _record_import_value(record, field_name, default_school_year)
    source = _selected_import_source(record, field_name) or "manual"
    return CurrentStudentField(
        value=value,
        source=source,
        confidence=_json_field_confidence(field_name, value=value, source=source),
        raw_value=None,
    )


def _json_field_confidence(
    field_name: str,
    *,
    value: str | int | float | bool | None,
    source: SourceType,
) -> FieldConfidence:
    if _is_missing_import_value(value):
        return "missing"
    if source == "thai_id_scan" and field_name == "citizen_id":
        return "authoritative"
    if source == "ocr_form":
        return "review"
    return "high"


def _export_missing_blockers(
    records: list[CanonicalStudentRecord],
    *,
    default_school_year: int | None,
) -> list[CurrentStudentsFieldConflict]:
    conflicts: list[CurrentStudentsFieldConflict] = []
    for record in records:
        for field_name in DMC_BLOCKING_MISSING_FIELDS:
            selected_value = _record_import_value(record, field_name, default_school_year)
            if not _is_missing_import_value(selected_value):
                continue
            conflicts.append(
                CurrentStudentsFieldConflict(
                    record_id=record.record_id,
                    full_name=_export_full_name(record),
                    student_no=record.student_no,
                    citizen_id=record.citizen_id,
                    field_name=field_name,
                    field_label=IMPORT_FIELD_LABELS.get(field_name, field_name),
                    selected_value=None,
                    selected_source=None,
                    selected_basis=MISSING_BLOCKER_BASIS,
                    reason="missing_after_all_sources",
                    source_values=[],
                )
            )
    return conflicts


def _is_missing_import_value(value: str | int | float | bool | None) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return _clean_text(value) in EMPTY_MARKERS
    return False


def _export_conflicts(
    records: list[CanonicalStudentRecord],
    *,
    default_school_year: int | None,
) -> list[CurrentStudentsFieldConflict]:
    conflicts: list[CurrentStudentsFieldConflict] = []
    for record in records:
        for conflict in record.conflicts:
            selected_value, selected_source = _selected_import_field(record, conflict.field_name, default_school_year)
            conflicts.append(
                conflict.model_copy(
                    update={
                        "full_name": _export_full_name(record),
                        "student_no": record.student_no,
                        "citizen_id": record.citizen_id,
                        "selected_value": selected_value,
                        "selected_source": selected_source,
                        "selected_basis": _selected_basis(conflict.field_name, selected_source),
                    }
                )
            )
    return conflicts


def _add_field_conflict(
    record: CanonicalStudentRecord,
    *,
    field_name: str,
    reason: str,
    source_fields: Sequence[CurrentStudentField],
) -> None:
    source_values = _conflict_values(record, source_fields)
    if len({_value_key(value.value) for value in source_values}) < 2:
        return

    selected_value, selected_source = _selected_import_field(record, field_name, default_school_year=None)
    conflict = CurrentStudentsFieldConflict(
        record_id=record.record_id,
        full_name=_export_full_name(record),
        student_no=record.student_no,
        citizen_id=record.citizen_id,
        field_name=field_name,
        field_label=IMPORT_FIELD_LABELS.get(field_name, field_name),
        selected_value=selected_value,
        selected_source=selected_source,
        selected_basis=_selected_basis(field_name, selected_source),
        reason=reason,
        source_values=source_values,
    )

    existing = next((item for item in record.conflicts if item.field_name == field_name), None)
    if existing is None:
        record.conflicts.append(conflict)
        return

    existing.reason = reason
    existing.selected_value = conflict.selected_value
    existing.selected_source = conflict.selected_source
    existing.selected_basis = conflict.selected_basis
    existing.source_values = _merge_conflict_values(existing.source_values, conflict.source_values)


def _conflict_values(
    record: CanonicalStudentRecord,
    source_fields: Sequence[CurrentStudentField],
) -> list[CurrentStudentsConflictValue]:
    values: list[CurrentStudentsConflictValue] = []
    for field in source_fields:
        reference = _source_reference_for(record, field.source)
        values.append(
            CurrentStudentsConflictValue(
                source=field.source,
                value=field.value,
                confidence=field.confidence,
                raw_value=field.raw_value,
                source_path=reference.source_path if reference is not None else None,
                row_index=reference.row_index if reference is not None else None,
                sheet_name=reference.sheet_name if reference is not None else None,
            )
        )
    return _merge_conflict_values([], values)


def _merge_conflict_values(
    existing: Sequence[CurrentStudentsConflictValue],
    incoming: Sequence[CurrentStudentsConflictValue],
) -> list[CurrentStudentsConflictValue]:
    merged: list[CurrentStudentsConflictValue] = []
    seen: set[tuple[SourceType, str, str | None]] = set()
    for value in [*existing, *incoming]:
        key = (value.source, _value_key(value.value), value.source_path)
        if key in seen:
            continue
        seen.add(key)
        merged.append(value)
    return merged


def _selected_import_field(
    record: CanonicalStudentRecord,
    field_name: str,
    default_school_year: int | None,
) -> tuple[str | int | float | bool | None, SourceType | None]:
    return _record_import_value(record, field_name, default_school_year), _selected_import_source(record, field_name)


def _selected_import_source(record: CanonicalStudentRecord, field_name: str) -> SourceType | None:
    if field_name == "operation_type":
        return "manual"
    if field_name == "school_year":
        field = record.dmc_fields.get(field_name)
        return field.source if field is not None else "manual"
    if field_name in {"student_no", "grade", "room", "seat_no"}:
        if _has_source(record, "roster"):
            return "roster"
        field = record.dmc_fields.get(field_name)
        return field.source if field is not None else _single_record_source(record)
    if field_name == "citizen_id":
        field = record.dmc_fields.get(field_name)
        if field is not None:
            return field.source
        if _has_source(record, "thai_id_scan"):
            return "thai_id_scan"
        if _has_source(record, "ocr_form"):
            return "ocr_form"
        return _single_record_source(record)
    if field_name in {"prefix", "first_name", "last_name"}:
        if _has_source(record, "roster"):
            return "roster"
        if _has_source(record, "thai_id_scan"):
            return "thai_id_scan"
        if _has_source(record, "ocr_form"):
            return "ocr_form"
        return _single_record_source(record)

    field = record.dmc_fields.get(field_name)
    return field.source if field is not None else _single_record_source(record)


def _selected_basis(field_name: str, source: SourceType | None) -> str:
    if source == "roster":
        if field_name in {"student_no", "school_year", "grade", "room", "seat_no"}:
            return "ยึดบัญชีรายชื่อ เพราะเป็นแหล่งเลขประจำตัว ชั้น ห้อง และเลขที่ของโรงเรียนในปีการศึกษานี้"
        if field_name in {"prefix", "first_name", "last_name"}:
            return "ยึดบัญชีรายชื่อ เพราะเป็นรายชื่อหลักของโรงเรียนและใช้รูปแบบคำนำหน้าเต็มสำหรับนำเข้า DMC"
        return "ยึดบัญชีรายชื่อ เพราะเป็นแหล่งข้อมูลหลักของนักเรียนในโรงเรียน"
    if source == "thai_id_scan":
        if field_name == "citizen_id":
            return "ยึด CSV เครื่องสแกนบัตร เพราะเลขบัตรมาจากบัตรประชาชนและผ่านการตรวจ checksum"
        return "ยึด CSV เครื่องสแกนบัตร เพราะเป็นข้อมูลจากบัตรประชาชนและมีความน่าเชื่อถือสูงกว่า OCR"
    if source == "ocr_form":
        return "ยึด OCR แบบฟอร์ม เพราะเป็นข้อมูลที่นักเรียนกรอกในแบบฟอร์ม DMC และไม่มีแหล่งที่น่าเชื่อถือกว่ามาทับ"
    if source == "civil_registration":
        return "ยึด OCR สำเนาทะเบียนบ้าน เพราะตัวเลขทะเบียนบ้าน เลขบัตรผู้ปกครอง และชื่อบิดามารดามาจากเอกสารทะเบียนราษฎร"
    if source == "derived":
        if field_name.startswith("guardian."):
            return "เติมข้อมูลผู้ปกครองจากข้อมูลบิดาหรือมารดา เพราะความเกี่ยวข้องผู้ปกครองระบุว่าเป็นบิดาหรือมารดา"
        if field_name == "mother.last_name":
            return "เติมนามสกุลมารดาจากนามสกุลบิดา เพราะตรวจจากทุกไฟล์แล้วไม่พบนามสกุลมารดาโดยตรง"
        return "เติมค่าจากข้อมูลที่ระบบสรุปได้จากไฟล์ที่เลือก"
    if source == "manual":
        return "ยึดค่าที่ผู้ใช้เลือกในหน้าจอสร้างไฟล์"
    return "ยังระบุแหล่งหลักไม่ได้ ต้องตรวจข้อมูลด้วยผู้ใช้"


def _source_reference_for(record: CanonicalStudentRecord, source: SourceType) -> SourceReference | None:
    return next((reference for reference in record.sources if reference.source == source), None)


def _single_record_source(record: CanonicalStudentRecord) -> SourceType | None:
    if len(record.sources) == 1:
        return record.sources[0].source
    return None


def _export_full_name(record: CanonicalStudentRecord) -> str | None:
    full_name = _full_name(record.prefix or "", record.first_name or "", record.last_name or "")
    return full_name or record.full_name


def _value_key(value: str | int | float | bool | None) -> str:
    return "" if value is None else str(value).strip()


def _default_import_output_path(school_year: int) -> Path:
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    return reports_dir() / f"current-students-import-{school_year}-{timestamp}.xlsx"


def _default_form_json_output_path(school_year: int) -> Path:
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    return reports_dir() / f"dmc-form-data-{school_year}-{timestamp}.json"


def _read_import_field_columns(sheet: Any) -> dict[str, int]:
    for row_index in (IMPORT_FIELD_KEY_ROW, IMPORT_HEADER_ROW):
        columns: dict[str, int] = {}
        row: Sequence[Any] = next(sheet.iter_rows(min_row=row_index, max_row=row_index, values_only=True), [])
        for column_index, value in enumerate(row, start=1):
            field_name = _clean_text(value)
            if field_name in IMPORT_FIELD_NAMES:
                columns[field_name] = column_index
        if set(REQUIRED_IMPORT_FIELDS).issubset(columns):
            return columns
    return {}


def _read_dmc_form_json_import_records(json_path: Path) -> list[DmcFormJsonRecord]:
    text, _encoding = _decode_text(json_path)
    try:
        payload = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DomainError("CURRENT_STUDENTS_JSON_INVALID", "Current student JSON file is not valid JSON.") from exc

    if not isinstance(payload, dict):
        raise DomainError(
            "CURRENT_STUDENTS_JSON_UNSUPPORTED_SCHEMA",
            "Current student JSON file must be a dmc_form_json.v1 object.",
        )

    try:
        import_payload = _DmcFormJsonImportPayload.model_validate(payload)
    except ValidationError as exc:
        raise DomainError(
            "CURRENT_STUDENTS_JSON_UNSUPPORTED_SCHEMA",
            "Current student JSON file must use schema_version dmc_form_json.v1 and include records.",
        ) from exc
    return import_payload.records


def _json_duplicate_citizen_ids(records: Sequence[DmcFormJsonRecord]) -> set[str]:
    citizen_ids = [
        cleaned
        for record in records
        if (cleaned := _clean_citizen_id(_json_record_citizen_id(record))) is not None
    ]
    citizen_counts = Counter(citizen_ids)
    return {citizen_id for citizen_id, count in citizen_counts.items() if count > 1}


def _json_record_text(record: DmcFormJsonRecord, field_name: str) -> str | None:
    value = record.fields.get(field_name)
    return _none_if_empty(value)


def _json_record_int(record: DmcFormJsonRecord, field_name: str, top_level_value: int | None) -> int | None:
    value: Any = top_level_value
    if value is None:
        value = record.fields.get(field_name)
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(str(value))
    except ValueError:
        return None
    if not number.is_integer():
        return None
    return int(number)


def _json_record_student_no(record: DmcFormJsonRecord) -> str | None:
    return _none_if_empty(record.student_no) or _json_record_text(record, "student_no")


def _json_record_citizen_id(record: DmcFormJsonRecord) -> str | None:
    return _none_if_empty(record.citizen_id) or _json_record_text(record, "citizen_id")


def _json_record_grade(record: DmcFormJsonRecord) -> int | None:
    return _json_record_int(record, "grade", record.grade)


def _json_record_room(record: DmcFormJsonRecord) -> str | None:
    if record.room is not None:
        return str(record.room)
    return _json_record_text(record, "room") or _json_record_text(record, "classroom")


def _json_record_level_code(record: DmcFormJsonRecord) -> str | None:
    raw_level_code = _json_record_text(record, "levelDtlCode") or _json_record_text(record, "level_dtl_code")
    if raw_level_code:
        digits = "".join(re.findall(r"\d", raw_level_code))
        if not digits:
            return None
        level_code = digits.zfill(2)
        return level_code if level_code in DMC_TRANSFER_IN_VALID_LEVEL_CODES else None
    grade = _json_record_grade(record)
    return DMC_TRANSFER_IN_LEVEL_CODES.get(grade) if grade is not None else None


def _json_record_full_name(record: DmcFormJsonRecord) -> str:
    full_name = _none_if_empty(record.full_name)
    if full_name:
        return full_name
    prefix = _none_if_empty(record.prefix) or _json_record_text(record, "prefix") or ""
    first_name = _none_if_empty(record.first_name) or _json_record_text(record, "first_name") or ""
    last_name = _none_if_empty(record.last_name) or _json_record_text(record, "last_name") or ""
    return _full_name(prefix, first_name, last_name)


def _validate_dmc_transfer_in_json_record(
    record: DmcFormJsonRecord,
    duplicate_citizen_ids: set[str],
) -> list[str]:
    issues: list[str] = []
    student_no = _json_record_student_no(record)
    if not student_no:
        issues.append("MISSING_STUDENT_NO")
    elif not re.fullmatch(r"\d{1,10}", student_no):
        issues.append("INVALID_STUDENT_NO")

    has_raw_level_code = bool(_json_record_text(record, "levelDtlCode") or _json_record_text(record, "level_dtl_code"))
    if _json_record_grade(record) is None and not has_raw_level_code:
        issues.append("MISSING_GRADE")
    elif _json_record_level_code(record) is None:
        issues.append("DMC_TRANSFER_LEVEL_UNSUPPORTED")

    room = _json_record_room(record)
    if not room:
        issues.append("MISSING_ROOM")
    elif not re.fullmatch(r"\d{1,2}", room):
        issues.append("INVALID_ROOM")

    citizen_id = _json_record_citizen_id(record)
    cleaned_citizen_id = _clean_citizen_id(citizen_id)
    if not citizen_id:
        issues.append("MISSING_CITIZEN_ID")
    elif cleaned_citizen_id is None or not is_valid_thai_citizen_id(cleaned_citizen_id):
        issues.append("INVALID_CITIZEN_ID")
    elif cleaned_citizen_id in duplicate_citizen_ids:
        issues.append("DUPLICATE_CITIZEN_ID")
    return issues


def _validate_import_row(values: dict[str, str], duplicate_citizen_ids: set[str]) -> list[str]:
    issues: list[str] = []
    for field_name in REQUIRED_IMPORT_FIELDS:
        if not values.get(field_name):
            issues.append(f"MISSING_{field_name.upper().replace('.', '_')}")
    operation_type = values.get("operation_type")
    if operation_type and operation_type not in {"current", "transfer_in", "add_new"}:
        issues.append("INVALID_OPERATION_TYPE")
    citizen_id = values.get("citizen_id")
    if citizen_id:
        cleaned = _clean_citizen_id(citizen_id)
        if cleaned is None or not is_valid_thai_citizen_id(cleaned):
            issues.append("INVALID_CITIZEN_ID")
        elif cleaned in duplicate_citizen_ids:
            issues.append("DUPLICATE_CITIZEN_ID")
    return issues


def _import_row_status(issues: list[str]) -> Literal["ready", "needs_review", "invalid"]:
    if not issues:
        return "ready"
    if any(
        issue.startswith("MISSING_")
        or issue.startswith("INVALID_")
        or issue in {"DMC_TRANSFER_LEVEL_UNSUPPORTED"}
        for issue in issues
    ):
        return "invalid"
    return "needs_review"


def _import_issue_message(issue: str) -> str:
    messages = {
        "INVALID_OPERATION_TYPE": "Operation type must be current, transfer_in, or add_new.",
        "INVALID_CITIZEN_ID": "Citizen ID is invalid or failed checksum validation.",
        "INVALID_STUDENT_NO": "Student number must contain 1-10 digits for the DMC transfer-in form.",
        "INVALID_ROOM": "Classroom must contain 1-2 digits for the DMC transfer-in form.",
        "DUPLICATE_CITIZEN_ID": "Citizen ID is duplicated in this import form.",
        "DMC_TRANSFER_LEVEL_UNSUPPORTED": "Grade could not be mapped to the DMC transfer-in levelDtlCode field.",
    }
    if issue.startswith("MISSING_"):
        return "Required field is missing."
    return messages.get(issue, issue)


def _base_record(student: RosterStudent, operation_type: OperationType) -> CanonicalStudentRecord:
    fields = {
        "student_no": _field(student.student_no, "roster", "authoritative"),
        "school_year": _field(student.school_year, "roster", "authoritative"),
        "grade": _field(student.grade, "roster", "authoritative"),
        "room": _field(student.room, "roster", "authoritative"),
        "seat_no": _field(student.seat_no, "roster", "authoritative"),
        "prefix": _field(student.prefix, "roster", "high"),
        "first_name": _field(student.first_name, "roster", "high"),
        "last_name": _field(student.last_name, "roster", "high"),
        "full_name": _field(student.full_name, "roster", "high"),
    }
    if student.citizen_id:
        fields["citizen_id"] = _field(student.citizen_id, "roster", "authoritative")
    if student.birth_date:
        fields["birth_date"] = _field(student.birth_date, "roster", "high")
    has_roster_citizen_id = student.citizen_id is not None and student.citizen_id_valid
    return CanonicalStudentRecord(
        record_id=_roster_key(student),
        operation_type=operation_type,
        match_status="auto_matched" if has_roster_citizen_id else "needs_review",
        student_no=student.student_no,
        citizen_id=student.citizen_id if has_roster_citizen_id else None,
        grade=student.grade,
        room=student.room,
        seat_no=student.seat_no,
        prefix=student.prefix,
        first_name=student.first_name,
        last_name=student.last_name,
        full_name=student.full_name,
        review_reasons=[] if has_roster_citizen_id else ["missing_thai_id_scan"],
        dmc_fields=fields,
        sources=[
            SourceReference(
                source="roster",
                source_path=student.source_path,
                sheet_name=student.sheet_name,
                row_index=student.row_index,
            )
        ],
    )


def _attach_scan(record: CanonicalStudentRecord, scan: ThaiIdScanRecord) -> None:
    record.sources.append(
        SourceReference(source="thai_id_scan", source_path=scan.source_path, row_index=scan.row_index)
    )
    record.citizen_id = record.citizen_id or scan.citizen_id
    _merge_authoritative_field(record, "citizen_id", scan.citizen_id, "thai_id_scan")
    _merge_authoritative_field(record, "prefix", scan.prefix, "thai_id_scan")
    _merge_authoritative_field(record, "first_name", scan.first_name, "thai_id_scan")
    _merge_authoritative_field(record, "last_name", scan.last_name, "thai_id_scan")
    _put_record_field(record, "first_name_en", scan.first_name_en, "thai_id_scan", "authoritative")
    _put_record_field(record, "last_name_en", scan.last_name_en, "thai_id_scan", "authoritative")
    _put_record_field(record, "sex", scan.sex, "thai_id_scan", "authoritative")
    _put_record_field(record, "birth_date", scan.birth_date, "thai_id_scan", "authoritative")
    _put_record_field(record, "registered_address.house_no", scan.house_no, "thai_id_scan", "authoritative")
    _put_record_field(record, "registered_address.moo", scan.moo, "thai_id_scan", "authoritative")
    _put_record_field(record, "registered_address.subdistrict", scan.subdistrict, "thai_id_scan", "authoritative")
    _put_record_field(record, "registered_address.district", scan.district, "thai_id_scan", "authoritative")
    _put_record_field(record, "registered_address.province", scan.province, "thai_id_scan", "authoritative")


def _attach_ocr(record: CanonicalStudentRecord, ocr_record: OcrFormRecord) -> None:
    record.sources.append(SourceReference(source="ocr_form", source_path=ocr_record.source_path))
    if record.citizen_id is None and ocr_record.citizen_id and ocr_record.citizen_id_valid:
        record.citizen_id = ocr_record.citizen_id
    if ocr_record.citizen_id and record.citizen_id and ocr_record.citizen_id != record.citizen_id:
        _add_reason(record, "ocr_conflicts_with_citizen_id")
        if record.match_status == "auto_matched":
            record.match_status = "needs_review"
    for field_name, field_value in ocr_record.fields.items():
        if field_value.value is None:
            continue
        existing = record.dmc_fields.get(field_name)
        if existing is not None and existing.value not in {None, "", field_value.value}:
            _add_field_conflict(
                record,
                field_name=field_name,
                reason=f"ocr_conflicts_with_{field_name}",
                source_fields=[existing, field_value],
            )
            if field_name in {"citizen_id", "student_no", "first_name", "last_name"}:
                _add_reason(record, f"ocr_conflicts_with_{field_name}")
                if record.match_status == "auto_matched":
                    record.match_status = "needs_review"
            continue
        record.dmc_fields.setdefault(field_name, field_value)


def _attach_civil_registration(record: CanonicalStudentRecord, civil_record: CivilRegistrationRecord) -> None:
    record.sources.append(SourceReference(source="civil_registration", source_path=civil_record.source_path))
    if civil_record.citizen_id:
        if record.citizen_id and record.citizen_id != civil_record.citizen_id:
            _add_reason(record, "civil_registration_conflicts_with_citizen_id")
            if record.match_status == "auto_matched":
                record.match_status = "needs_review"
            _add_field_conflict(
                record,
                field_name="citizen_id",
                reason="civil_registration_conflicts_with_citizen_id",
                source_fields=[_field(record.citizen_id, _selected_import_source(record, "citizen_id") or "manual", "high"), civil_record.fields["citizen_id"]],
            )
        else:
            record.citizen_id = civil_record.citizen_id
    for field_name, field_value in civil_record.fields.items():
        if field_value.value is None:
            continue
        if field_name in {"prefix", "first_name", "last_name"}:
            if not getattr(record, field_name):
                setattr(record, field_name, field_value.value)
                record.dmc_fields[field_name] = field_value
            continue
        _merge_civil_registration_field(record, field_name, field_value)


def _merge_civil_registration_field(
    record: CanonicalStudentRecord,
    field_name: str,
    incoming: CurrentStudentField,
) -> None:
    existing = record.dmc_fields.get(field_name)
    if existing is None or existing.value in {None, ""}:
        record.dmc_fields[field_name] = incoming
        return
    if existing.value == incoming.value:
        return
    _add_field_conflict(
        record,
        field_name=field_name,
        reason=f"civil_registration_conflicts_with_{field_name}",
        source_fields=[existing, incoming],
    )
    if existing.source == "ocr_form":
        record.dmc_fields[field_name] = incoming
        return
    if field_name in {
        "registered_address.house_id",
        "father.citizen_id",
        "father.first_name",
        "father.last_name",
        "mother.citizen_id",
        "mother.first_name",
        "mother.last_name",
    } and existing.source not in {"roster", "thai_id_scan"}:
        record.dmc_fields[field_name] = incoming


def _promote_registered_address_to_current_when_matching(records: list[CanonicalStudentRecord]) -> None:
    for record in records:
        if not _registered_address_matches_current_address(record):
            continue
        for suffix in ADDRESS_FIELD_SUFFIXES:
            registered_field = record.dmc_fields.get(_address_field_name("registered_address", suffix))
            if registered_field is None or _is_missing_import_value(registered_field.value):
                continue
            record.dmc_fields[_address_field_name("current_address", suffix)] = CurrentStudentField(
                value=registered_field.value,
                source=registered_field.source,
                confidence=registered_field.confidence,
                raw_value=registered_field.raw_value,
            )


def _registered_address_matches_current_address(record: CanonicalStudentRecord) -> bool:
    comparable = 0
    matched = 0
    matched_house_id = False
    for suffix in ADDRESS_FIELD_SUFFIXES:
        registered_value = _address_compare_value(
            suffix,
            record.dmc_fields.get(_address_field_name("registered_address", suffix)),
        )
        current_value = _address_compare_value(
            suffix,
            record.dmc_fields.get(_address_field_name("current_address", suffix)),
        )
        if registered_value is None or current_value is None:
            continue
        comparable += 1
        if _address_values_match(suffix, registered_value, current_value):
            matched += 1
            if suffix == "house_id":
                matched_house_id = True

    if matched_house_id:
        return True
    if comparable < ADDRESS_COPY_MIN_COMPARABLE_FIELDS:
        return False
    return matched / comparable >= ADDRESS_COPY_MIN_MATCH_RATIO


def _address_field_name(prefix: Literal["registered_address", "current_address"], suffix: str) -> str:
    return f"{prefix}.{suffix}"


def _address_compare_value(suffix: str, field: CurrentStudentField | None) -> str | None:
    if field is None:
        return None
    value = field.value
    if _is_missing_import_value(value):
        return None
    text = _clean_text(value).lower()
    text = re.sub(
        r"^(เลขที่|บ้านเลขที่|หมู่ที่|หมู่|ม\.|ถนน|ถ\.|จังหวัด|จ\.|อำเภอ|อ\.|เขต|ตำบล|ต\.|แขวง)\s*",
        "",
        text,
    )
    if suffix in {"house_id", "postal_code"}:
        digits = "".join(re.findall(r"\d", text))
        return digits or None
    return re.sub(r"[^\w/]", "", text, flags=re.UNICODE) or None


def _address_values_match(suffix: str, registered_value: str, current_value: str) -> bool:
    if registered_value == current_value:
        return True
    if suffix in {"house_id", "postal_code", "house_no", "moo"}:
        return False
    return SequenceMatcher(None, registered_value, current_value).ratio() >= 0.86


def _fill_missing_mother_last_name_from_father(records: list[CanonicalStudentRecord]) -> None:
    for record in records:
        mother_last_name = record.dmc_fields.get("mother.last_name")
        if mother_last_name is not None and not _is_missing_import_value(mother_last_name.value):
            continue

        father_last_name = record.dmc_fields.get("father.last_name")
        if father_last_name is None or _is_missing_import_value(father_last_name.value):
            continue

        record.dmc_fields["mother.last_name"] = CurrentStudentField(
            value=father_last_name.value,
            source="derived",
            confidence="review",
            raw_value="derived_from_father.last_name",
        )


def _fill_guardian_from_parent_relationship(records: list[CanonicalStudentRecord]) -> None:
    for record in records:
        relationship_field = record.dmc_fields.get("guardian_relationship")
        relationship_value = relationship_field.value if relationship_field is not None else None
        parent = _guardian_parent_from_relationship(relationship_value)
        if parent is None:
            continue

        for suffix in GUARDIAN_PARENT_FIELD_SUFFIXES:
            parent_field = record.dmc_fields.get(f"{parent}.{suffix}")
            if parent_field is None or _is_missing_import_value(parent_field.value):
                continue
            record.dmc_fields[f"guardian.{suffix}"] = CurrentStudentField(
                value=parent_field.value,
                source="derived",
                confidence=parent_field.confidence,
                raw_value=f"derived_from_{parent}.{suffix}:{parent_field.source}",
            )


def _guardian_parent_from_relationship(value: str | int | float | bool | None) -> Literal["father", "mother"] | None:
    text = _clean_text(value)
    if not text:
        return None
    has_father = GUARDIAN_FATHER_MARKER in text
    has_mother = GUARDIAN_MOTHER_MARKER in text
    if has_father == has_mother:
        return None
    return "father" if has_father else "mother"


def _merge_authoritative_field(
    record: CanonicalStudentRecord,
    field_name: str,
    value: str | int | float | bool | None,
    source: SourceType,
) -> None:
    if value is None or value == "":
        return
    existing = record.dmc_fields.get(field_name)
    incoming = _field(value, source, "authoritative")
    if existing is not None and existing.value not in {None, "", value}:
        _add_reason(record, f"{source}_conflicts_with_{field_name}")
        if record.match_status == "auto_matched":
            record.match_status = "needs_review"
        record.dmc_fields[field_name] = incoming
        _add_field_conflict(
            record,
            field_name=field_name,
            reason=f"{source}_conflicts_with_{field_name}",
            source_fields=[existing, incoming],
        )
        return
    record.dmc_fields[field_name] = incoming


def _put_record_field(
    record: CanonicalStudentRecord,
    field_name: str,
    value: str | int | float | bool | None,
    source: SourceType,
    confidence: FieldConfidence,
) -> None:
    if value is None or value == "":
        return
    record.dmc_fields.setdefault(field_name, _field(value, source, confidence))


def _scan_only_record(
    scan: ThaiIdScanRecord,
    operation_type: OperationType,
    status: MatchStatus,
    reason: str,
    suggestions: list[MatchSuggestion],
) -> CanonicalStudentRecord:
    record = CanonicalStudentRecord(
        record_id=f"scan:{Path(scan.source_path).name}:{scan.row_index}",
        operation_type=operation_type,
        match_status=status,
        student_no=None,
        citizen_id=scan.citizen_id,
        grade=None,
        room=None,
        seat_no=None,
        prefix=scan.prefix,
        first_name=scan.first_name,
        last_name=scan.last_name,
        full_name=_full_name(scan.prefix, scan.first_name, scan.last_name),
        match_score=suggestions[0].score if suggestions else None,
        review_reasons=[reason],
        suggestions=suggestions,
        sources=[SourceReference(source="thai_id_scan", source_path=scan.source_path, row_index=scan.row_index)],
    )
    _attach_scan(record, scan)
    return record


def _ocr_only_record(
    ocr_record: OcrFormRecord,
    operation_type: OperationType,
    status: MatchStatus,
    reason: str,
    suggestions: list[MatchSuggestion],
) -> CanonicalStudentRecord:
    reasons = [reason]
    if not ocr_record.citizen_id_valid:
        reasons.append("invalid_or_missing_ocr_citizen_id")
    return CanonicalStudentRecord(
        record_id=f"ocr:{Path(ocr_record.source_path).name}:{ocr_record.record_id}",
        operation_type=operation_type,
        match_status=status,
        student_no=ocr_record.student_no,
        citizen_id=ocr_record.citizen_id,
        grade=None,
        room=None,
        seat_no=None,
        prefix=ocr_record.prefix,
        first_name=ocr_record.first_name,
        last_name=ocr_record.last_name,
        full_name=ocr_record.full_name,
        match_score=suggestions[0].score if suggestions else None,
        review_reasons=reasons,
        suggestions=suggestions,
        dmc_fields=ocr_record.fields,
        sources=[SourceReference(source="ocr_form", source_path=ocr_record.source_path)],
    )


def _civil_registration_only_record(
    civil_record: CivilRegistrationRecord,
    operation_type: OperationType,
) -> CanonicalStudentRecord:
    reasons = ["civil_registration_not_linked_to_roster_or_thai_id_scan"]
    status: MatchStatus = "new_or_transfer_candidate" if civil_record.citizen_id_valid else "invalid_id"
    if not civil_record.citizen_id_valid:
        reasons.append("invalid_or_missing_civil_registration_citizen_id")
    return CanonicalStudentRecord(
        record_id=f"civil:{Path(civil_record.source_path).name}:{civil_record.record_id}",
        operation_type=operation_type,
        match_status=status,
        student_no=None,
        citizen_id=civil_record.citizen_id,
        grade=None,
        room=None,
        seat_no=None,
        prefix=civil_record.prefix,
        first_name=civil_record.first_name,
        last_name=civil_record.last_name,
        full_name=civil_record.full_name,
        review_reasons=reasons,
        dmc_fields=civil_record.fields,
        sources=[SourceReference(source="civil_registration", source_path=civil_record.source_path)],
    )


def _summary(
    *,
    roster: list[RosterStudent],
    scans: list[ThaiIdScanRecord],
    ocr_records: list[OcrFormRecord],
    civil_records: list[CivilRegistrationRecord],
    records: list[CanonicalStudentRecord],
    ocr_attached: int,
    ocr_unmatched: int,
    warnings: list[CurrentStudentsWarning],
) -> CurrentStudentsSummary:
    status_counts = Counter(record.match_status for record in records)
    duplicate_scan_records = sum(count for _key, count in Counter(scan.citizen_id for scan in scans if scan.citizen_id).items() if count > 1)
    return CurrentStudentsSummary(
        roster_records=len(roster),
        thai_id_scan_records=len(scans),
        ocr_form_records=len(ocr_records),
        civil_registration_records=len(civil_records),
        records_total=len(records),
        auto_matched=status_counts["auto_matched"],
        needs_review=status_counts["needs_review"],
        duplicate_records=status_counts["duplicate"],
        duplicate_scan_records=duplicate_scan_records,
        invalid_id_records=status_counts["invalid_id"],
        new_or_transfer_candidates=status_counts["new_or_transfer_candidate"],
        roster_without_thai_id=sum(1 for record in records if "missing_thai_id_scan" in record.review_reasons),
        ocr_attached_records=ocr_attached,
        ocr_unmatched_records=ocr_unmatched,
        review_queue_records=sum(1 for record in records if record.match_status != "auto_matched"),
        warnings_total=len(warnings),
    )


def _export_summary(
    source_summary: CurrentStudentsSummary,
    records: list[CanonicalStudentRecord],
    *,
    warnings: list[CurrentStudentsWarning],
) -> CurrentStudentsSummary:
    status_counts = Counter(record.match_status for record in records)
    return source_summary.model_copy(
        update={
            "records_total": len(records),
            "auto_matched": status_counts["auto_matched"],
            "needs_review": status_counts["needs_review"],
            "duplicate_records": status_counts["duplicate"],
            "invalid_id_records": status_counts["invalid_id"],
            "new_or_transfer_candidates": status_counts["new_or_transfer_candidate"],
            "roster_without_thai_id": sum(1 for record in records if "missing_thai_id_scan" in record.review_reasons),
            "review_queue_records": sum(1 for record in records if record.match_status != "auto_matched"),
            "warnings_total": len(warnings),
        }
    )


def _has_source(record: CanonicalStudentRecord, source: SourceType) -> bool:
    return any(reference.source == source for reference in record.sources)


def _is_unconfirmed_ocr_record(record: CanonicalStudentRecord) -> bool:
    return (
        _has_source(record, "ocr_form")
        and not _has_source(record, "roster")
        and "user_confirmed_ocr_roster_match" not in record.review_reasons
    )


def _confirmed_match_map(confirmations: Sequence[DmcFormMatchConfirmation]) -> dict[str, str]:
    return {confirmation.record_id: confirmation.student_no for confirmation in confirmations}


def _ocr_record_key(ocr_record: OcrFormRecord) -> str:
    return f"ocr:{Path(ocr_record.source_path).name}:{ocr_record.record_id}"


def _record_sort_key(record: CanonicalStudentRecord) -> tuple[int, int, int, str]:
    grade = record.grade if record.grade is not None else 999
    room = record.room if record.room is not None else 999
    seat = record.seat_no if record.seat_no is not None else 999
    return (grade, room, seat, record.record_id)


def _best_ocr_suggestions(
    ocr_record: OcrFormRecord,
    roster: list[RosterStudent],
    *,
    limit: int,
) -> list[MatchSuggestion]:
    if not ocr_record.name_key:
        return []
    ocr_birth_date = _field_text_value(ocr_record.fields.get("birth_date"))
    scored: list[MatchSuggestion] = []
    for student in roster:
        name_score = SequenceMatcher(None, ocr_record.name_key, student.name_key).ratio()
        if name_score <= 0:
            continue
        birth_date_matches = _date_texts_match(ocr_birth_date, student.birth_date)
        score = min(1.0, name_score + 0.08) if birth_date_matches and name_score >= 0.75 else name_score
        scored.append(
            MatchSuggestion(
                student_no=student.student_no,
                full_name=_full_name(student.prefix, student.first_name, student.last_name),
                grade=student.grade,
                room=student.room,
                score=score,
                source_path=student.source_path,
                sheet_name=student.sheet_name,
                row_index=student.row_index,
            )
        )
    return sorted(scored, key=lambda item: item.score, reverse=True)[:limit]


def _suggestion_score_for_student_no(suggestions: Sequence[MatchSuggestion], student_no: str) -> float | None:
    suggestion = next((item for item in suggestions if item.student_no == student_no), None)
    return suggestion.score if suggestion is not None else None


def _field_text_value(field: CurrentStudentField | None) -> str | None:
    if field is None:
        return None
    return _none_if_empty(field.value)


def _date_texts_match(left: str | None, right: str | None) -> bool:
    if not left or not right:
        return False
    return re.sub(r"\s+", "", _clean_text(left)) == re.sub(r"\s+", "", _clean_text(right))


def _best_suggestions(name_key: str, roster: list[RosterStudent], *, limit: int) -> list[MatchSuggestion]:
    scored: list[MatchSuggestion] = []
    for student in roster:
        score = SequenceMatcher(None, name_key, student.name_key).ratio()
        if score <= 0:
            continue
        scored.append(
            MatchSuggestion(
                student_no=student.student_no,
                full_name=_full_name(student.prefix, student.first_name, student.last_name),
                grade=student.grade,
                room=student.room,
                score=score,
                source_path=student.source_path,
                sheet_name=student.sheet_name,
                row_index=student.row_index,
            )
        )
    return sorted(scored, key=lambda item: item.score, reverse=True)[:limit]


def _duplicates(values: Iterable[str | None]) -> set[str]:
    counts = Counter(value for value in values if value)
    return {value for value, count in counts.items() if count > 1}


def _set_status(record: CanonicalStudentRecord, status: MatchStatus, reason: str) -> None:
    priority = {
        "auto_matched": 0,
        "new_or_transfer_candidate": 1,
        "needs_review": 2,
        "invalid_id": 3,
        "duplicate": 4,
    }
    if priority[status] >= priority[record.match_status]:
        record.match_status = status
    _add_reason(record, reason)
    _remove_reason(record, "missing_thai_id_scan")


def _add_reason(record: CanonicalStudentRecord, reason: str) -> None:
    if reason not in record.review_reasons:
        record.review_reasons.append(reason)


def _remove_reason(record: CanonicalStudentRecord, reason: str) -> None:
    record.review_reasons = [item for item in record.review_reasons if item != reason]


def _field(
    value: str | int | float | bool | None,
    source: SourceType,
    confidence: FieldConfidence,
    raw_value: str | None = None,
) -> CurrentStudentField:
    return CurrentStudentField(value=value, source=source, confidence=confidence, raw_value=raw_value)


def _put_field(
    fields: dict[str, CurrentStudentField],
    field_name: str,
    value: str | int | float | bool | None,
    source: SourceType,
    confidence: FieldConfidence,
) -> None:
    clean_value = _normalized_field_value(value)
    if clean_value is None:
        return
    fields[field_name] = _field(clean_value, source, confidence, raw_value=str(value) if value is not None else None)


def _put_address_fields(fields: dict[str, CurrentStudentField], text: str, *, prefix: str, start: str) -> None:
    segment = _extract_between(text, start, ("ที่อยู่ปัจจุบัน", "๓. รายละเอียดนักเรียน", "3. รายละเอียดนักเรียน"))
    if not segment:
        return
    _put_field(fields, f"{prefix}.house_id", _extract_between(segment, "รหัสประจำบ้าน*", ("บ้านเลขที่*",)), "ocr_form", "review")
    _put_field(fields, f"{prefix}.house_no", _extract_between(segment, "บ้านเลขที่*", ("หมู่ที่",)), "ocr_form", "review")
    _put_field(fields, f"{prefix}.moo", _extract_between(segment, "หมู่ที่ (ถ้าไม่มีใส่ -)", ("ถนน",)), "ocr_form", "review")
    _put_field(fields, f"{prefix}.road", _extract_between(segment, "ถนน (ถ้าไม่มีใส่ -)", ("จังหวัด*",)), "ocr_form", "review")
    _put_field(fields, f"{prefix}.province", _extract_between(segment, "จังหวัด*", ("อำเภอ*",)), "ocr_form", "review")
    _put_field(fields, f"{prefix}.district", _extract_between(segment, "อำเภอ*", ("ตำบล*",)), "ocr_form", "review")
    _put_field(fields, f"{prefix}.subdistrict", _extract_between(segment, "ตำบล*", ("รหัสไปรษณีย์*",)), "ocr_form", "review")
    _put_field(fields, f"{prefix}.postal_code", _first_digits(_extract_between(segment, "รหัสไปรษณีย์*", ("หมายเลขโทรศัพท์บ้าน",))), "ocr_form", "review")


def _put_guardian_fields(
    fields: dict[str, CurrentStudentField],
    text: str,
    *,
    person: Literal["father", "mother", "guardian"],
    start: str,
    end: str,
) -> None:
    segment = _extract_between(text, start, (end,))
    if not segment:
        return
    _put_field(fields, f"{person}.citizen_id", _clean_citizen_id(_extract_between(segment, "เลขบัตรประจำตัวประชาชน", ("ชนิดบัตร*",))), "ocr_form", "review")
    _put_field(fields, f"{person}.card_type", _checked_option(_extract_between(segment, "ชนิดบัตร*", ("ชื่อ",))), "ocr_form", "review")
    _put_field(fields, f"{person}.first_name", _extract_between(segment, "ชื่อ" + _person_label(person) + "*", ("นามสกุล*",)), "ocr_form", "review")
    _put_field(fields, f"{person}.last_name", _extract_between(segment, "นามสกุล*", ("กลุ่มเลือด",)), "ocr_form", "review")
    _put_field(fields, f"{person}.blood_type", _extract_between(segment, "กลุ่มเลือด" + _person_label(person) + "*", ("อาชีพ*",)), "ocr_form", "review")
    _put_field(fields, f"{person}.occupation", _extract_between(segment, "อาชีพ*", ("รายได้ต่อเดือน",)), "ocr_form", "review")
    _put_field(fields, f"{person}.income_text", _extract_between(segment, "รายได้ต่อเดือน(บาท)*", ("หมายเลขโทรศัพท์",)), "ocr_form", "review")
    _put_field(fields, f"{person}.phone", _phone_text(segment), "ocr_form", "review")


def _put_civil_parent_fields(
    fields: dict[str, CurrentStudentField],
    text: str,
    *,
    person: Literal["father", "mother"],
    marker: str,
) -> None:
    pattern = rf"{re.escape(marker)}[:\s]+ชื่อ[:\s]+(.+?)\s+(?:เลขประจำตัวประชาชน[:\s]+)?([0-9][0-9\-\s]{{10,}})(?=\s+(?:สัญชาติ|สถานภาพ|มาจาก|เกิดเมื่อ|บิดาผู้ให้กำเนิด|มารดาผู้ให้กำเนิด)|$)"
    match = re.search(pattern, text)
    if match is None:
        return
    _prefix, first_name, last_name = _split_civil_name(match.group(1))
    _put_field(fields, f"{person}.citizen_id", _clean_citizen_id(match.group(2)), "civil_registration", "high")
    _put_field(fields, f"{person}.first_name", first_name, "civil_registration", "high")
    _put_field(fields, f"{person}.last_name", last_name, "civil_registration", "high")


def _split_civil_name(value: str | None) -> tuple[str | None, str | None, str | None]:
    text = _clean_text(value)
    if not text:
        return None, None, None
    prefix: str | None = None
    for candidate in TITLE_PREFIXES:
        compact_candidate = candidate.replace(".", "")
        if text.startswith(candidate):
            prefix = candidate
            text = text[len(candidate) :].strip()
            break
        if text.startswith(compact_candidate):
            prefix = candidate
            text = text[len(compact_candidate) :].strip()
            break
    parts = [part for part in re.split(r"\s+", text) if part]
    if not parts:
        return prefix, None, None
    if len(parts) == 1:
        return prefix, parts[0], None
    return prefix, parts[0], " ".join(parts[1:])


def _regex_first(text: str, pattern: str) -> str | None:
    match = re.search(pattern, text)
    return _none_if_empty(match.group(1)) if match is not None else None


def _marker_variants(marker: str) -> tuple[str, ...]:
    clean_marker = _clean_text(marker)
    stripped_marker = _clean_text(clean_marker.replace("*", ""))
    candidates = [
        clean_marker,
        stripped_marker,
        stripped_marker.replace("/", ""),
        stripped_marker.replace(" ", ""),
        stripped_marker.replace("/", "").replace(" ", ""),
    ]
    unique: list[str] = []
    for candidate in candidates:
        if candidate and candidate not in unique:
            unique.append(candidate)
    return tuple(unique)


def _find_marker(text: str, marker: str, start: int = 0) -> tuple[int, int] | None:
    for candidate in _marker_variants(marker):
        index = text.find(candidate, start)
        if index >= 0:
            return index, len(candidate)
    return None


def _clean_house_id(value: str | None) -> str | None:
    text = _clean_text(value)
    if not text:
        return None
    return re.sub(r"\s+", "", text)


def _person_label(person: Literal["father", "mother", "guardian"]) -> str:
    return {"father": "บิดา", "mother": "มารดา", "guardian": "ผู้ปกครอง"}[person]


def _split_ocr_markdown_pages(text: str) -> list[str]:
    parts = re.split(r"<!--\s*Page\s+\d+\s+confidence:\s*[^>]*-->", text, flags=re.IGNORECASE)
    pages = [part.strip() for part in parts[1:] if _clean_text(part)]
    return pages or [text]


def _looks_like_civil_registration_text(text: str) -> bool:
    compact_text = _clean_text(text)
    return any(
        marker in compact_text
        for marker in (
            "รายการเกี่ยวกับบ้าน",
            "รายการบุคคลในบ้าน",
            "เลขรหัสประจำบ้าน",
        )
    )


def _decode_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp874", "tis-620"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8-replace"


def _validated_file(path: Path, *, suffixes: set[str]) -> Path:
    if not path.exists():
        raise DomainError("CURRENT_STUDENTS_INPUT_NOT_FOUND", "Input file was not found.")
    if not path.is_file():
        raise DomainError("CURRENT_STUDENTS_INPUT_NOT_FILE", "Selected input path is not a file.")
    if path.suffix.lower() not in suffixes:
        raise DomainError("CURRENT_STUDENTS_UNSUPPORTED_INPUT_TYPE", "Selected file type is not supported.")
    return path


def _detect_roster_year(sheet: Any) -> int | None:
    first_row: Sequence[Any] = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
    first_cell = _clean_text(first_row[0] if first_row else "")
    years = re.findall(r"\d{4}", first_cell)
    return int(years[-1]) if years else None


def _roster_header_key(value: Any) -> str | None:
    text = re.sub(r"\s+", "", _clean_text(value).lower())
    text = text.replace("-", "")
    aliases = {
        "ลำดับ": "seat_no",
        "เลขที่": "seat_no",
        "รหัสนักเรียน": "student_no",
        "เลขประจำตัวนักเรียน": "student_no",
        "เลขนักเรียน": "student_no",
        "ชื่อ–นามสกุล": "full_name",
        "ชื่อนามสกุล": "full_name",
        "ชื่อ": "full_name",
        "วันเกิด": "birth_date",
        "วันเดือนปีเกิด": "birth_date",
        "เลขบัตรประชาชน": "citizen_id",
        "เลขประจำตัวประชาชน": "citizen_id",
        "เลขบัตรประจำตัวประชาชน": "citizen_id",
        "ชั้น": "grade",
        "ระดับชั้น": "grade",
        "ห้อง": "room",
    }
    return aliases.get(text)


def _cell_by_header(row: Sequence[Any], columns: dict[str, int], name: str) -> Any:
    index = columns.get(name)
    if index is None or index >= len(row):
        return None
    return row[index]


def _int_cell(value: Any) -> int | None:
    if _looks_numeric(value):
        return int(float(str(value)))
    digits = _first_digits(_clean_text(value))
    return int(digits) if digits is not None else None


def _student_no_text(value: Any) -> str | None:
    if _looks_numeric(value):
        return str(int(float(str(value))))
    digits = _first_digits(_clean_text(value))
    return digits


def _grade_level_number(value: Any) -> int | None:
    text = _clean_text(value)
    if not text:
        return None
    match = re.search(r"(?:ม\.?|มัธยมศึกษาปีที่)\s*(\d+)", text)
    if match is not None:
        return int(match.group(1))
    return _int_cell(value)


def _split_roster_name(full_name: str) -> tuple[str, str]:
    parts = [part for part in re.split(r"\s+", _clean_text(full_name)) if part]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def _split_roster_full_name(full_name: str) -> tuple[str | None, str, str]:
    prefix, first_name, last_name = _split_civil_name(full_name)
    if first_name:
        return prefix, first_name, last_name or ""
    first_name, last_name = _split_roster_name(full_name)
    return None, first_name, last_name


def _clean_citizen_id(value: str | None) -> str | None:
    digits = "".join(re.findall(r"\d", _clean_text(value)))
    if len(digits) != 13:
        return None
    return digits


def _first_digits(value: str | None) -> str | None:
    match = re.search(r"\d+", _clean_text(value))
    return match.group(0) if match else None


def _number_text(value: str | None) -> str | None:
    text = _clean_text(value)
    if text in EMPTY_MARKERS:
        return None
    match = re.search(r"\d+(?:[.,]\d+)?", text)
    if match is None:
        return None
    return match.group(0)


def _phone_text(value: str) -> str | None:
    matches = re.findall(r"(?<!\d)\d{2,3}[- ]?\d{3}[- ]?\d{4}(?!\d)", value)
    return matches[-1] if matches else None


def _checked_option(value: str | None) -> str | None:
    text = _clean_text(value)
    match = re.search(r"\[x\]\s*([^\[\n]+?)(?=\s*\[|$)", text, flags=re.IGNORECASE)
    if match is None:
        return None
    return _clean_text(match.group(1))


def _extract_between(text: str, start: str, end_markers: Sequence[str]) -> str | None:
    start_match = _find_marker(text, start)
    if start_match is None:
        return None
    start_index, start_length = start_match
    value_start = start_index + start_length
    value_end = len(text)
    for marker in end_markers:
        marker_match = _find_marker(text, marker, value_start)
        if marker_match is not None:
            marker_index, _marker_length = marker_match
            value_end = min(value_end, marker_index)
    return _none_if_empty(text[value_start:value_end])


def _none_if_empty(value: Any) -> str | None:
    text = _clean_text(value)
    return None if text in EMPTY_MARKERS else text


def _normalized_field_value(value: str | int | float | bool | None) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, str):
        return _none_if_empty(value)
    return value


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"^(?:\.{2,}|…+)+\s*", "", text)
    text = re.sub(r"\s*(?:\.{2,}|…+)+$", "", text)
    return text.strip()


def _row_is_empty(row: Sequence[str]) -> bool:
    return not any(_clean_text(value) for value in row)


def _looks_numeric(value: Any) -> bool:
    if value is None:
        return False
    try:
        float(str(value))
    except ValueError:
        return False
    return True


def _full_name(prefix: str, first_name: str, last_name: str) -> str:
    return _clean_text(" ".join(part for part in (prefix, first_name, last_name) if part))


def _roster_key(student: RosterStudent) -> str:
    return f"roster:{student.sheet_name}:{student.row_index}:{student.student_no}"
