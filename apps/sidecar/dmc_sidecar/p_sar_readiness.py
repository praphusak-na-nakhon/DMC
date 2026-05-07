from __future__ import annotations

import hashlib
import json
import re
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Sequence, cast
from xml.etree import ElementTree
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile
from xml.sax.saxutils import escape

from openpyxl import load_workbook  # type: ignore[import-untyped]

from .config import bundled_resources_root, repo_root, reports_dir
from .db import utc_now
from .errors import DomainError
from .schemas import (
    AddPsarEvidenceResponse,
    EvidenceMappingStatus,
    GeneratePsarReportResponse,
    PsarEvidenceMapping,
    PsarMissingEvidenceRecommendation,
    PsarReadinessRequirement,
    PsarReadinessResponse,
    PsarReadinessSection,
    PsarRequirement,
    PsarUploadedFile,
    ReadinessPriority,
    ReadinessRequirementStatus,
)


DEFAULT_CONFIDENCE_THRESHOLD = 0.75
DEFAULT_WARNING_THRESHOLD = 0.8
SUPPORTED_EVIDENCE_SUFFIXES = {".pdf", ".docx", ".xlsx", ".xlsm", ".xls", ".png", ".jpg", ".jpeg"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}
PDF_PAGE_PATTERN = re.compile(rb"/Type\s*/Page\b")
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_NS = "http://www.w3.org/XML/1998/namespace"
WORD_NAMESPACES = {"w": WORD_NS}
DOCX_XML_NAMESPACES: tuple[tuple[str, str], ...] = (
    ("mc", "http://schemas.openxmlformats.org/markup-compatibility/2006"),
    ("o", "urn:schemas-microsoft-com:office:office"),
    ("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships"),
    ("m", "http://schemas.openxmlformats.org/officeDocument/2006/math"),
    ("v", "urn:schemas-microsoft-com:vml"),
    ("wp", "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"),
    ("w10", "urn:schemas-microsoft-com:office:word"),
    ("w", WORD_NS),
    ("wne", "http://schemas.microsoft.com/office/word/2006/wordml"),
    ("sl", "http://schemas.openxmlformats.org/schemaLibrary/2006/main"),
    ("a", "http://schemas.openxmlformats.org/drawingml/2006/main"),
    ("pic", "http://schemas.openxmlformats.org/drawingml/2006/picture"),
    ("c", "http://schemas.openxmlformats.org/drawingml/2006/chart"),
    ("lc", "http://schemas.openxmlformats.org/drawingml/2006/lockedCanvas"),
    ("dgm", "http://schemas.openxmlformats.org/drawingml/2006/diagram"),
    ("wps", "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"),
    ("wpg", "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"),
    ("w14", "http://schemas.microsoft.com/office/word/2010/wordml"),
    ("w15", "http://schemas.microsoft.com/office/word/2012/wordml"),
    ("w16", "http://schemas.microsoft.com/office/word/2018/wordml"),
    ("w16cex", "http://schemas.microsoft.com/office/word/2018/wordml/cex"),
    ("w16cid", "http://schemas.microsoft.com/office/word/2016/wordml/cid"),
    ("cr", "http://schemas.microsoft.com/office/comments/2020/reactions"),
)


THAI_EVIDENCE_LABELS: dict[str, str] = {
    "school/profile information": "ข้อมูลโรงเรียน/ข้อมูลพื้นฐาน",
    "teacher/personnel information": "ข้อมูลครู/บุคลากร",
    "student count": "จำนวนนักเรียน",
    "context/background information": "บริบทและข้อมูลพื้นฐาน",
    "teaching assignment": "ข้อมูลการมอบหมายงานสอน",
    "class schedule": "ตารางสอน",
    "advisor duty": "ข้อมูลครูที่ปรึกษา",
    "training certificates": "เกียรติบัตร/วุฒิบัตรการอบรม",
    "awards or recognitions": "รางวัลหรือประกาศเกียรติคุณ",
    "academic achievement results": "ผลสัมฤทธิ์ทางการเรียน",
    "grade distribution by semester": "ผลการเรียนแยกตามภาคเรียน",
    "reading/thinking/analyzing/writing assessment": "ผลการอ่าน คิด วิเคราะห์ และเขียน",
    "desirable characteristics assessment": "ผลการประเมินคุณลักษณะอันพึงประสงค์",
    "learner competency assessment": "ผลการประเมินสมรรถนะสำคัญของผู้เรียน",
    "learner development activities": "กิจกรรมพัฒนาผู้เรียน",
    "supporting evidence for learning outcomes": "หลักฐานสนับสนุนผลลัพธ์การเรียนรู้",
    "school/action plan": "แผนปฏิบัติการของโรงเรียน",
    "quality development plan": "แผนพัฒนาคุณภาพ",
    "meeting records": "บันทึกการประชุม",
    "committee orders": "คำสั่งแต่งตั้ง/คำสั่งปฏิบัติหน้าที่",
    "internal quality assurance process": "กระบวนการประกันคุณภาพภายใน",
    "monitoring and evaluation evidence": "หลักฐานการนิเทศ ติดตาม และประเมินผล",
    "lesson plans": "แผนการจัดการเรียนรู้",
    "course plan": "แผนรายวิชา",
    "learning management records": "บันทึกการจัดการเรียนรู้",
    "post-teaching reflection": "บันทึกหลังสอน/การสะท้อนผลหลังสอน",
    "classroom observation/supervision records": "บันทึกการนิเทศหรือสังเกตชั้นเรียน",
    "student assessment evidence": "หลักฐานการวัดและประเมินผลผู้เรียน",
    "assessment tools": "เครื่องมือวัดและประเมินผล",
    "teaching innovation or learning media evidence": "สื่อ นวัตกรรม หรือเทคโนโลยีการสอน",
    "strengths": "จุดเด่น",
    "weaknesses or areas for improvement": "จุดที่ควรพัฒนา",
    "next improvement plan": "แผนพัฒนาครั้งต่อไป",
    "responsible persons": "ผู้รับผิดชอบ",
    "timeline or action plan": "ระยะเวลา/แผนปฏิบัติการ",
}


SECTION_TITLES: dict[str, tuple[int, str]] = {
    "basic_information": (1, "Basic Information"),
    "standard_1_learner_quality": (2, "Standard 1: Learner Quality"),
    "standard_2_management": (3, "Standard 2: Management and Administration"),
    "standard_3_teaching": (4, "Standard 3: Teaching and Learning Process"),
    "summary_improvement": (5, "Summary and Improvement Plan"),
}


P_SAR_REQUIREMENT_MATRIX: tuple[PsarRequirement, ...] = (
    PsarRequirement(
        requirement_id="basic.profile",
        section_id="basic_information",
        section_title=SECTION_TITLES["basic_information"][1],
        section_order=SECTION_TITLES["basic_information"][0],
        title="School, teacher, and student profile",
        category="profile",
        required_evidence=[
            "school/profile information",
            "teacher/personnel information",
            "student count",
            "context/background information",
        ],
        optional_evidence=["education history", "leave record", "teacher appointment"],
        weight=8,
        minimum_required_items=3,
        description="Baseline school, teacher, personnel, class, student, and context fields from the P-SAR cover and Part 1.",
    ),
    PsarRequirement(
        requirement_id="basic.teaching_assignment",
        section_id="basic_information",
        section_title=SECTION_TITLES["basic_information"][1],
        section_order=SECTION_TITLES["basic_information"][0],
        title="Teaching assignment and advisory duties",
        category="profile",
        required_evidence=["teaching assignment", "class schedule", "advisor duty"],
        optional_evidence=["special duty order", "weekly teaching hours"],
        weight=5,
        minimum_required_items=2,
        description="Courses, class levels, teaching load, advisor assignment, and additional duties for the reporting year.",
    ),
    PsarRequirement(
        requirement_id="basic.professional_development",
        section_id="basic_information",
        section_title=SECTION_TITLES["basic_information"][1],
        section_order=SECTION_TITLES["basic_information"][0],
        title="Professional development and recognition",
        category="development",
        required_evidence=["training certificates", "awards or recognitions"],
        optional_evidence=["speaker invitation", "committee invitation"],
        weight=4,
        minimum_required_items=1,
        description="Training, certificates, awards, and invited professional service used in Part 2 of the P-SAR form.",
    ),
    PsarRequirement(
        requirement_id="learner.academic_results",
        section_id="standard_1_learner_quality",
        section_title=SECTION_TITLES["standard_1_learner_quality"][1],
        section_order=SECTION_TITLES["standard_1_learner_quality"][0],
        title="Academic achievement results",
        category="learner_quality",
        required_evidence=["academic achievement results", "grade distribution by semester"],
        optional_evidence=["target comparison", "subject achievement summary"],
        weight=10,
        minimum_required_items=1,
        description="Student achievement tables and target comparisons for each semester and subject area.",
    ),
    PsarRequirement(
        requirement_id="learner.reading_thinking_writing",
        section_id="standard_1_learner_quality",
        section_title=SECTION_TITLES["standard_1_learner_quality"][1],
        section_order=SECTION_TITLES["standard_1_learner_quality"][0],
        title="Reading, thinking, analyzing, and writing assessment",
        category="learner_quality",
        required_evidence=["reading/thinking/analyzing/writing assessment"],
        optional_evidence=["assessment summary by semester"],
        weight=7,
        minimum_required_items=1,
        description="Assessment results for reading, analytical thinking, writing, and communication.",
    ),
    PsarRequirement(
        requirement_id="learner.desirable_characteristics",
        section_id="standard_1_learner_quality",
        section_title=SECTION_TITLES["standard_1_learner_quality"][1],
        section_order=SECTION_TITLES["standard_1_learner_quality"][0],
        title="Desirable characteristics assessment",
        category="learner_quality",
        required_evidence=["desirable characteristics assessment"],
        optional_evidence=["8 desirable characteristics order"],
        weight=7,
        minimum_required_items=1,
        description="Learner desirable-characteristics evidence and summary results.",
    ),
    PsarRequirement(
        requirement_id="learner.competencies",
        section_id="standard_1_learner_quality",
        section_title=SECTION_TITLES["standard_1_learner_quality"][1],
        section_order=SECTION_TITLES["standard_1_learner_quality"][0],
        title="Learner competency assessment",
        category="learner_quality",
        required_evidence=["learner competency assessment"],
        optional_evidence=["competency summary by semester"],
        weight=6,
        minimum_required_items=1,
        description="Core learner competency results from the curriculum assessment section.",
    ),
    PsarRequirement(
        requirement_id="learner.development_activities",
        section_id="standard_1_learner_quality",
        section_title=SECTION_TITLES["standard_1_learner_quality"][1],
        section_order=SECTION_TITLES["standard_1_learner_quality"][0],
        title="Learner development activities and outcomes",
        category="learner_quality",
        required_evidence=["learner development activities", "supporting evidence for learning outcomes"],
        optional_evidence=["activity photos", "student work samples"],
        weight=6,
        minimum_required_items=1,
        description="Activity records, photos, learner work, and evidence supporting learning outcomes.",
    ),
    PsarRequirement(
        requirement_id="management.action_plan",
        section_id="standard_2_management",
        section_title=SECTION_TITLES["standard_2_management"][1],
        section_order=SECTION_TITLES["standard_2_management"][0],
        title="School action plan and quality development plan",
        category="management",
        required_evidence=["school/action plan", "quality development plan"],
        optional_evidence=["vision and mission document"],
        weight=8,
        minimum_required_items=1,
        description="Plans and goals that support school quality development and the P-SAR management standard.",
    ),
    PsarRequirement(
        requirement_id="management.meeting_records",
        section_id="standard_2_management",
        section_title=SECTION_TITLES["standard_2_management"][1],
        section_order=SECTION_TITLES["standard_2_management"][0],
        title="Meeting records and committee orders",
        category="management",
        required_evidence=["meeting records", "committee orders"],
        optional_evidence=["working group assignment"],
        weight=6,
        minimum_required_items=1,
        description="Meeting minutes, school orders, and committee assignments related to management and administration.",
    ),
    PsarRequirement(
        requirement_id="management.quality_assurance",
        section_id="standard_2_management",
        section_title=SECTION_TITLES["standard_2_management"][1],
        section_order=SECTION_TITLES["standard_2_management"][0],
        title="Internal quality assurance and monitoring",
        category="management",
        required_evidence=["internal quality assurance process", "monitoring and evaluation evidence"],
        optional_evidence=["IQA award evidence", "control self assessment"],
        weight=8,
        minimum_required_items=1,
        description="Internal QA process, monitoring, evaluation, and follow-up evidence.",
    ),
    PsarRequirement(
        requirement_id="teaching.lesson_plans",
        section_id="standard_3_teaching",
        section_title=SECTION_TITLES["standard_3_teaching"][1],
        section_order=SECTION_TITLES["standard_3_teaching"][0],
        title="Lesson plans and learning design",
        category="teaching_process",
        required_evidence=["lesson plans", "course plan"],
        optional_evidence=["learning unit design", "active learning plan"],
        weight=9,
        minimum_required_items=1,
        description="Learning plans, course planning, and assessment design from the teaching-process standard.",
    ),
    PsarRequirement(
        requirement_id="teaching.learning_records",
        section_id="standard_3_teaching",
        section_title=SECTION_TITLES["standard_3_teaching"][1],
        section_order=SECTION_TITLES["standard_3_teaching"][0],
        title="Learning management records and reflection",
        category="teaching_process",
        required_evidence=["learning management records", "post-teaching reflection"],
        optional_evidence=["classroom activity record"],
        weight=7,
        minimum_required_items=1,
        description="Teaching records and reflections showing implementation of learning activities.",
    ),
    PsarRequirement(
        requirement_id="teaching.supervision",
        section_id="standard_3_teaching",
        section_title=SECTION_TITLES["standard_3_teaching"][1],
        section_order=SECTION_TITLES["standard_3_teaching"][0],
        title="Classroom observation and supervision",
        category="teaching_process",
        required_evidence=["classroom observation/supervision records"],
        optional_evidence=["supervision feedback", "PLC feedback"],
        weight=7,
        minimum_required_items=1,
        description="Observation, supervision, and feedback evidence for classroom practice.",
    ),
    PsarRequirement(
        requirement_id="teaching.assessment_evidence",
        section_id="standard_3_teaching",
        section_title=SECTION_TITLES["standard_3_teaching"][1],
        section_order=SECTION_TITLES["standard_3_teaching"][0],
        title="Student assessment evidence",
        category="teaching_process",
        required_evidence=["student assessment evidence", "assessment tools"],
        optional_evidence=["rubric", "test paper", "student score record"],
        weight=7,
        minimum_required_items=1,
        description="Assessment tools, records, rubrics, and results used to evaluate learners systematically.",
    ),
    PsarRequirement(
        requirement_id="teaching.media_innovation",
        section_id="standard_3_teaching",
        section_title=SECTION_TITLES["standard_3_teaching"][1],
        section_order=SECTION_TITLES["standard_3_teaching"][0],
        title="Teaching media and innovation",
        category="teaching_process",
        required_evidence=["teaching innovation or learning media evidence"],
        optional_evidence=["technology tools", "local wisdom learning source"],
        weight=5,
        minimum_required_items=1,
        description="Media, technology, innovation, and learning-resource evidence used in teaching.",
    ),
    PsarRequirement(
        requirement_id="improvement.strengths_weaknesses",
        section_id="summary_improvement",
        section_title=SECTION_TITLES["summary_improvement"][1],
        section_order=SECTION_TITLES["summary_improvement"][0],
        title="Strengths and areas for improvement",
        category="improvement",
        required_evidence=["strengths", "weaknesses or areas for improvement"],
        optional_evidence=["self-assessment summary"],
        weight=7,
        minimum_required_items=2,
        description="Summary of strengths, weaknesses, and improvement needs from Part 3.",
    ),
    PsarRequirement(
        requirement_id="improvement.next_plan",
        section_id="summary_improvement",
        section_title=SECTION_TITLES["summary_improvement"][1],
        section_order=SECTION_TITLES["summary_improvement"][0],
        title="Next improvement plan",
        category="improvement",
        required_evidence=["next improvement plan", "responsible persons", "timeline or action plan"],
        optional_evidence=["support request"],
        weight=8,
        minimum_required_items=2,
        description="Future improvement plan, responsible people, timeline, and requested support.",
    ),
)


@dataclass(frozen=True)
class KeywordRule:
    requirement_id: str
    evidence_type: str
    keywords: tuple[str, ...]
    base_confidence: float


@dataclass(frozen=True)
class PsarTemplateFacts:
    teacher_name: str | None = None
    subject_group: str | None = None


KEYWORD_RULES: tuple[KeywordRule, ...] = (
    KeywordRule("basic.profile", "school/profile information", ("profile", "school", "teacher", "personnel", "student count", "ข้อมูลพื้นฐาน", "ข้อมูลทั่วไป", "ประวัติ"), 0.86),
    KeywordRule("basic.teaching_assignment", "teaching assignment", ("teaching assignment", "schedule", "ภาระการสอน", "ตารางสอน", "ครูที่ปรึกษา", "advisor"), 0.86),
    KeywordRule("basic.professional_development", "training certificates", ("certificate", "training", "อบรม", "เกียรติบัตร", "วุฒิบัตร", "รางวัล"), 0.84),
    KeywordRule("learner.academic_results", "academic achievement results", ("achievement", "grade", "ผลสัมฤทธิ์", "ระดับผลการเรียน", "คะแนน", "grade distribution"), 0.88),
    KeywordRule("learner.reading_thinking_writing", "reading/thinking/analyzing/writing assessment", ("reading", "thinking", "writing", "อ่าน", "คิดวิเคราะห์", "เขียน"), 0.88),
    KeywordRule("learner.desirable_characteristics", "desirable characteristics assessment", ("desirable", "characteristics", "คุณลักษณะ", "8 ประการ"), 0.88),
    KeywordRule("learner.competencies", "learner competency assessment", ("competency", "competencies", "สมรรถนะ"), 0.86),
    KeywordRule("learner.development_activities", "learner development activities", ("learner development", "activity", "กิจกรรมพัฒนาผู้เรียน", "ผลงานนักเรียน", "student work"), 0.82),
    KeywordRule("management.action_plan", "school/action plan", ("action plan", "development plan", "แผนปฏิบัติการ", "แผนพัฒนา", "vision", "mission"), 0.86),
    KeywordRule("management.meeting_records", "meeting records", ("meeting", "minutes", "committee", "คำสั่ง", "ประชุม", "แต่งตั้ง"), 0.84),
    KeywordRule("management.quality_assurance", "internal quality assurance process", ("quality assurance", "IQA", "ประกันคุณภาพ", "monitoring", "evaluation", "นิเทศ", "ติดตาม"), 0.86),
    KeywordRule("teaching.lesson_plans", "lesson plans", ("lesson plan", "แผนการจัดการเรียนรู้", "course plan", "learning design", "หน่วยการเรียนรู้"), 0.88),
    KeywordRule("teaching.learning_records", "learning management records", ("teaching record", "reflection", "บันทึกหลังสอน", "บันทึกการสอน", "learning management"), 0.84),
    KeywordRule("teaching.supervision", "classroom observation/supervision records", ("observation", "supervision", "classroom", "นิเทศ", "สังเกตการสอน", "PLC"), 0.84),
    KeywordRule("teaching.assessment_evidence", "student assessment evidence", ("assessment", "rubric", "test", "แบบทดสอบ", "วัดผล", "ประเมินผล", "score"), 0.84),
    KeywordRule("teaching.media_innovation", "teaching innovation or learning media evidence", ("media", "innovation", "สื่อ", "นวัตกรรม", "technology", "canva", "kahoot"), 0.82),
    KeywordRule("improvement.strengths_weaknesses", "strengths", ("strength", "weakness", "improvement area", "จุดเด่น", "จุดควรพัฒนา", "แนวทางการพัฒนา"), 0.82),
    KeywordRule("improvement.next_plan", "next improvement plan", ("next plan", "timeline", "responsible", "action plan", "แผนพัฒนา", "ผู้รับผิดชอบ", "ระยะเวลา", "ความต้องการช่วยเหลือ"), 0.82),
)


def compute_psar_readiness(
    *,
    project_id: str,
    mappings: Sequence[PsarEvidenceMapping],
    uploaded_files: Sequence[PsarUploadedFile] = (),
    requirements: Sequence[PsarRequirement] = P_SAR_REQUIREMENT_MATRIX,
    confidence_threshold: float = DEFAULT_CONFIDENCE_THRESHOLD,
    warning_threshold: float = DEFAULT_WARNING_THRESHOLD,
) -> PsarReadinessResponse:
    mapped_by_requirement: dict[str, list[PsarEvidenceMapping]] = {}
    for mapping in mappings:
        if mapping.status == "rejected":
            continue
        mapped_by_requirement.setdefault(mapping.requirement_id, []).append(mapping)

    readiness_items: list[tuple[PsarRequirement, PsarReadinessRequirement]] = []
    complete_count = partial_count = missing_count = needs_review_count = 0
    weighted_score = 0.0
    total_weight = sum(requirement.weight for requirement in requirements)

    for requirement in requirements:
        item = _score_requirement(
            requirement,
            mapped_by_requirement.get(requirement.requirement_id, []),
            confidence_threshold=confidence_threshold,
        )
        readiness_items.append((requirement, item))
        weighted_score += item.completion_score * requirement.weight
        if item.status == "complete":
            complete_count += 1
        elif item.status == "partial":
            partial_count += 1
        elif item.status == "needs_review":
            needs_review_count += 1
        else:
            missing_count += 1

    sections = _group_sections(readiness_items)
    recommendations = _build_recommendations(readiness_items)
    mapped_files = sorted(
        [mapping for mapping in mappings if mapping.status != "rejected"],
        key=lambda item: (item.requirement_id, item.source_file_name, item.evidence_type),
    )

    return PsarReadinessResponse(
        project_id=project_id,
        overall_completion_score=round(weighted_score / total_weight, 4) if total_weight else 0,
        warning_threshold=warning_threshold,
        confidence_threshold=confidence_threshold,
        total_requirements=len(requirements),
        complete_count=complete_count,
        partial_count=partial_count,
        missing_count=missing_count,
        needs_review_count=needs_review_count,
        sections=sections,
        missing_evidence_recommendations=recommendations,
        mapped_files=mapped_files,
        uploaded_files=list(uploaded_files),
    )


def build_ai_evidence_mapping_prompt(
    *,
    document_filename: str,
    extracted_text_summary: str,
    metadata: dict[str, Any],
    requirements: Sequence[PsarRequirement] = P_SAR_REQUIREMENT_MATRIX,
) -> str:
    schema_hint = {
        "mappings": [
            {
                "requirement_id": "basic.profile",
                "evidence_type": "school/profile information",
                "confidence_score": 0.0,
                "extracted_summary": "",
                "reason": "",
                "page_number": None,
            }
        ],
        "unmapped_reason": None,
    }
    return "\n".join(
        [
            "Classify this uploaded evidence against the fixed P-SAR requirement matrix.",
            "Return structured JSON only. Do not include markdown or explanations.",
            "Do not fabricate evidence. If support is unclear, return no mapping or a low-confidence mapping.",
            "The requirement matrix is authoritative; do not invent new P-SAR requirements.",
            f"Document filename: {document_filename}",
            f"Extracted text summary: {extracted_text_summary}",
            f"Metadata: {json.dumps(metadata, ensure_ascii=False, sort_keys=True)}",
            f"Available requirements: {json.dumps([item.model_dump() for item in requirements], ensure_ascii=False)}",
            f"Response shape: {json.dumps(schema_hint, ensure_ascii=False)}",
        ]
    )


class PsarReadinessService:
    def get_readiness(self, project_id: str = "default") -> PsarReadinessResponse:
        state = _load_project_state(project_id)
        files = [PsarUploadedFile.model_validate(item) for item in state.get("uploaded_files", [])]
        mappings = [PsarEvidenceMapping.model_validate(item) for item in state.get("mappings", [])]
        return compute_psar_readiness(project_id=project_id, mappings=mappings, uploaded_files=files)

    def add_evidence(self, *, project_id: str = "default", file_path: str) -> AddPsarEvidenceResponse:
        path = Path(file_path)
        _validate_evidence_path(path)
        state = _load_project_state(project_id)
        now = utc_now()
        uploaded_file = _uploaded_file_from_path(path, now=now)
        previous_files = [
            PsarUploadedFile.model_validate(item)
            for item in state.get("uploaded_files", [])
            if str(item.get("file_id")) != uploaded_file.file_id
        ]
        previous_mappings = [
            PsarEvidenceMapping.model_validate(item)
            for item in state.get("mappings", [])
            if str(item.get("source_file_id")) != uploaded_file.file_id
        ]
        mappings = _map_uploaded_file(uploaded_file, now=now)
        files = [*previous_files, uploaded_file]
        all_mappings = [*previous_mappings, *mappings]
        _save_project_state(project_id, uploaded_files=files, mappings=all_mappings)
        readiness = compute_psar_readiness(project_id=project_id, mappings=all_mappings, uploaded_files=files)
        return AddPsarEvidenceResponse(
            project_id=project_id,
            file=uploaded_file,
            mappings=mappings,
            readiness=readiness,
        )

    def generate_report(self, project_id: str = "default") -> GeneratePsarReportResponse:
        readiness = self.get_readiness(project_id)
        generated_at = utc_now()
        output_path = _project_dir(project_id) / f"p-sar-report-{_timestamp_slug(generated_at)}.docx"
        _write_psar_report_docx(output_path, readiness=readiness, generated_at=generated_at)
        return GeneratePsarReportResponse(
            project_id=project_id,
            report_path=str(output_path),
            readiness=readiness,
            generated_at=generated_at,
        )


def _score_requirement(
    requirement: PsarRequirement,
    mappings: Sequence[PsarEvidenceMapping],
    *,
    confidence_threshold: float,
) -> PsarReadinessRequirement:
    found = sorted(mappings, key=lambda item: (-item.confidence_score, item.source_file_name, item.evidence_type))
    accepted_types = {
        _normalize_evidence_type(mapping.evidence_type)
        for mapping in found
        if mapping.status == "accepted" and mapping.confidence_score >= confidence_threshold
    }
    high_review_types = {
        _normalize_evidence_type(mapping.evidence_type)
        for mapping in found
        if mapping.status in {"suggested", "needs_review"} and mapping.confidence_score >= confidence_threshold
    }
    low_types = {
        _normalize_evidence_type(mapping.evidence_type)
        for mapping in found
        if mapping.confidence_score < confidence_threshold or mapping.status == "needs_review"
    }
    accepted_count = len(accepted_types)
    reviewable_count = len(high_review_types - accepted_types)
    low_count = len(low_types - accepted_types - high_review_types)
    minimum = requirement.minimum_required_items
    completion_score = min((accepted_count + (0.5 * reviewable_count) + (0.25 * low_count)) / minimum, 1.0)
    missing_evidence = _missing_evidence(requirement, accepted_types)

    if accepted_count >= minimum:
        status: ReadinessRequirementStatus = "complete"
        completion_score = 1.0
        missing_evidence = []
    elif found and (low_types or high_review_types):
        status = "needs_review"
    elif found:
        status = "partial"
    else:
        status = "missing"
        completion_score = 0.0

    return PsarReadinessRequirement(
        requirement_id=requirement.requirement_id,
        requirement_title=requirement.title,
        category=requirement.category,
        description=requirement.description,
        required_evidence=requirement.required_evidence,
        optional_evidence=requirement.optional_evidence,
        weight=requirement.weight,
        minimum_required_items=minimum,
        status=status,
        completion_score=round(completion_score, 4),
        missing_evidence=missing_evidence,
        found_evidence=list(found),
        recommendation=_recommendation_text(status, requirement, missing_evidence),
        priority=_priority_for(status, requirement.weight),
    )


def _missing_evidence(requirement: PsarRequirement, accepted_types: set[str]) -> list[str]:
    missing = [
        evidence
        for evidence in requirement.required_evidence
        if _normalize_evidence_type(evidence) not in accepted_types
    ]
    if len(requirement.required_evidence) - len(missing) >= requirement.minimum_required_items:
        return []
    return missing


def _group_sections(
    readiness_items: Sequence[tuple[PsarRequirement, PsarReadinessRequirement]],
) -> list[PsarReadinessSection]:
    grouped: dict[str, list[tuple[PsarRequirement, PsarReadinessRequirement]]] = {}
    for requirement, item in readiness_items:
        grouped.setdefault(requirement.section_id, []).append((requirement, item))

    sections: list[PsarReadinessSection] = []
    for section_id, items in sorted(grouped.items(), key=lambda entry: min(item[0].section_order for item in entry[1])):
        total_weight = sum(requirement.weight for requirement, _item in items)
        score = (
            sum(item.completion_score * requirement.weight for requirement, item in items) / total_weight
            if total_weight
            else 0.0
        )
        requirements = [item for _requirement, item in items]
        sections.append(
            PsarReadinessSection(
                section_id=section_id,
                section_title=items[0][0].section_title,
                completion_score=round(score, 4),
                complete_count=sum(1 for item in requirements if item.status == "complete"),
                partial_count=sum(1 for item in requirements if item.status == "partial"),
                missing_count=sum(1 for item in requirements if item.status == "missing"),
                needs_review_count=sum(1 for item in requirements if item.status == "needs_review"),
                requirements=requirements,
            )
        )
    return sections


def _build_recommendations(
    readiness_items: Sequence[tuple[PsarRequirement, PsarReadinessRequirement]],
) -> list[PsarMissingEvidenceRecommendation]:
    recommendations: list[PsarMissingEvidenceRecommendation] = []
    for requirement, item in readiness_items:
        if item.status == "complete":
            continue
        missing = item.missing_evidence or requirement.required_evidence[: requirement.minimum_required_items]
        for evidence in missing:
            recommendations.append(
                PsarMissingEvidenceRecommendation(
                    evidence_name=evidence,
                    requirement_id=requirement.requirement_id,
                    requirement_title=requirement.title,
                    section_id=requirement.section_id,
                    section_title=requirement.section_title,
                    why_needed=item.recommendation,
                    priority=item.priority,
                    suggested_file_types=["PDF", "DOCX", "XLSX", "image"],
                )
            )
    priority_rank = {"high": 0, "medium": 1, "low": 2}
    return sorted(
        recommendations,
        key=lambda item: (priority_rank[item.priority], item.section_id, item.requirement_id, item.evidence_name),
    )


def _recommendation_text(
    status: ReadinessRequirementStatus,
    requirement: PsarRequirement,
    missing_evidence: Sequence[str],
) -> str:
    if status == "complete":
        return "Evidence found meets the current readiness rule for this requirement."
    if status == "needs_review":
        return "Evidence found needs human review or a clearer upload before this requirement should be treated as complete."
    if status == "partial":
        missing_text = ", ".join(missing_evidence[:3]) if missing_evidence else "additional required evidence"
        return f"Not enough evidence yet. Upload or confirm {missing_text}."
    first_missing = missing_evidence[0] if missing_evidence else requirement.required_evidence[0]
    return f"Recommended upload: {first_missing}."


def _priority_for(status: ReadinessRequirementStatus, weight: float) -> ReadinessPriority:
    if status in {"missing", "needs_review"} and weight >= 7:
        return "high"
    if status in {"missing", "needs_review", "partial"}:
        return "medium"
    return "low"


def _uploaded_file_from_path(path: Path, *, now: str) -> PsarUploadedFile:
    digest = hashlib.sha256(f"{path.resolve()}:{path.stat().st_mtime_ns}:{path.stat().st_size}".encode("utf-8")).hexdigest()
    return PsarUploadedFile(
        file_id=digest[:24],
        file_name=path.name,
        file_path=str(path),
        file_type=path.suffix.lower().lstrip("."),
        extracted_summary=_extract_summary(path),
        created_at=now,
        updated_at=now,
    )


def _map_uploaded_file(uploaded_file: PsarUploadedFile, *, now: str) -> list[PsarEvidenceMapping]:
    haystack = f"{uploaded_file.file_name}\n{uploaded_file.extracted_summary}".casefold()
    suffix = f".{uploaded_file.file_type.casefold()}"
    mappings: list[PsarEvidenceMapping] = []
    seen: set[tuple[str, str]] = set()
    for rule in KEYWORD_RULES:
        matched_keywords = [keyword for keyword in rule.keywords if keyword.casefold() in haystack]
        if not matched_keywords:
            continue
        confidence = min(rule.base_confidence + (0.02 * (len(matched_keywords) - 1)), 0.95)
        if suffix in IMAGE_SUFFIXES:
            confidence = min(confidence, 0.68)
        status: EvidenceMappingStatus = (
            "accepted"
            if confidence >= 0.82
            else "suggested"
            if confidence >= DEFAULT_CONFIDENCE_THRESHOLD
            else "needs_review"
        )
        key = (rule.requirement_id, rule.evidence_type)
        if key in seen:
            continue
        seen.add(key)
        mappings.append(
            PsarEvidenceMapping(
                requirement_id=rule.requirement_id,
                evidence_type=rule.evidence_type,
                source_file_id=uploaded_file.file_id,
                source_file_name=uploaded_file.file_name,
                extracted_summary=uploaded_file.extracted_summary,
                confidence_score=round(confidence, 2),
                page_number=None,
                location="filename and extracted metadata",
                status=status,
                created_at=now,
                updated_at=now,
            )
        )
    return mappings


def _extract_summary(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _extract_docx_summary(path)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return _extract_workbook_summary(path)
    if suffix == ".pdf":
        return _extract_pdf_summary(path)
    if suffix in IMAGE_SUFFIXES:
        return f"Image evidence file: {path.name}. Content requires human review unless an AI extractor is added."
    return f"Evidence file: {path.name}"


def _extract_docx_summary(path: Path) -> str:
    try:
        with ZipFile(path) as archive:
            document_xml = archive.read("word/document.xml")
    except (BadZipFile, KeyError, OSError):
        return f"DOCX evidence file: {path.name}. Text summary unavailable."
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ElementTree.fromstring(document_xml)
    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:p", namespace):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace)).strip()
        if text:
            paragraphs.append(text)
        if len(paragraphs) >= 20:
            break
    joined = " ".join(paragraphs)
    return _truncate(joined or f"DOCX evidence file: {path.name}")


def _extract_workbook_summary(path: Path) -> str:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return f"Workbook evidence file: {path.name}. Text summary unavailable."
    try:
        parts = [f"sheets: {', '.join(workbook.sheetnames[:8])}"]
        for sheet in workbook.worksheets[:3]:
            values: list[str] = []
            for row in sheet.iter_rows(max_row=8, max_col=8, values_only=True):
                for value in row:
                    if value is not None:
                        values.append(str(value))
                    if len(values) >= 16:
                        break
                if len(values) >= 16:
                    break
            if values:
                parts.append(f"{sheet.title}: {' | '.join(values)}")
        return _truncate(" ".join(parts))
    finally:
        workbook.close()


def _extract_pdf_summary(path: Path) -> str:
    try:
        data = path.read_bytes()
    except OSError:
        return f"PDF evidence file: {path.name}. Text summary unavailable."
    page_count = len(PDF_PAGE_PATTERN.findall(data))
    if page_count <= 0:
        page_count = 1 if data.startswith(b"%PDF") else 0
    return f"PDF evidence file: {path.name}. Estimated pages: {page_count}."


def _truncate(value: str, limit: int = 900) -> str:
    normalized = " ".join(value.split())
    if len(normalized) <= limit:
        return normalized
    return f"{normalized[:limit].rstrip()}..."


def _validate_evidence_path(path: Path) -> None:
    if not path.exists():
        raise DomainError("PSAR_EVIDENCE_FILE_NOT_FOUND", "Evidence file was not found.")
    if not path.is_file():
        raise DomainError("PSAR_EVIDENCE_PATH_NOT_FILE", "Selected evidence path is not a file.")
    if path.suffix.lower() not in SUPPORTED_EVIDENCE_SUFFIXES:
        raise DomainError("PSAR_EVIDENCE_UNSUPPORTED_TYPE", "Evidence must be PDF, DOCX, XLSX, or image.")


def _project_dir(project_id: str) -> Path:
    normalized = project_id.strip() or "default"
    if not re.fullmatch(r"[A-Za-z0-9_.-]{1,80}", normalized):
        raise DomainError("PSAR_PROJECT_ID_INVALID", "P-SAR project id contains unsupported characters.")
    path = reports_dir() / "psar" / normalized
    path.mkdir(parents=True, exist_ok=True)
    return path


def _project_state_path(project_id: str) -> Path:
    return _project_dir(project_id) / "readiness.json"


def _load_project_state(project_id: str) -> dict[str, Any]:
    path = _project_state_path(project_id)
    if not path.exists():
        return {"uploaded_files": [], "mappings": []}
    return cast(dict[str, Any], json.loads(path.read_text(encoding="utf-8")))


def _save_project_state(
    project_id: str,
    *,
    uploaded_files: Sequence[PsarUploadedFile],
    mappings: Sequence[PsarEvidenceMapping],
) -> None:
    path = _project_state_path(project_id)
    payload = {
        "uploaded_files": [item.model_dump() for item in uploaded_files],
        "mappings": [item.model_dump() for item in mappings],
    }
    temp_path: Path | None = None
    try:
        with NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as handle:
            temp_path = Path(handle.name)
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        temp_path.replace(path)
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _write_psar_report_docx(
    path: Path,
    *,
    readiness: PsarReadinessResponse,
    generated_at: str,
) -> None:
    template_path = _resolved_psar_template_path()
    with ZipFile(template_path, "r") as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        patched_document = _patch_psar_template_document(
            source.read("word/document.xml"),
            readiness=readiness,
            generated_at=generated_at,
        )
        for item in source.infolist():
            data = patched_document if item.filename == "word/document.xml" else source.read(item.filename)
            target.writestr(item, data)


def _resolved_psar_template_path() -> Path:
    bundled_root = bundled_resources_root()
    candidates = []
    if bundled_root is not None:
        candidates.append(bundled_root / "templates" / "P-SAR-form.docx")
    candidates.extend(
        [
            repo_root() / "P-SAR-form.docx",
            Path(__file__).resolve().parent / "templates" / "P-SAR-form.docx",
        ]
    )
    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return candidate
    raise DomainError("PSAR_TEMPLATE_NOT_FOUND", "Official P-SAR form template was not found.")


def _patch_psar_template_document(
    document_xml: bytes,
    *,
    readiness: PsarReadinessResponse,
    generated_at: str,
) -> bytes:
    xml = document_xml.decode("utf-8")
    facts = _extract_template_facts(readiness)

    if facts.teacher_name:
        xml = _replace_text_node_pattern_once(xml, r"…+\.", facts.teacher_name)
        xml = _replace_text_node_once(xml, "(        )", f"({facts.teacher_name})")
        xml = _replace_text_node_once(xml, "()", f"({facts.teacher_name})")
        xml = _replace_text_node_text_once(
            xml,
            "ตามที่ข้าพเจ้า   ได้รับมอบหมาย",
            f"ตามที่ข้าพเจ้า {facts.teacher_name} ได้รับมอบหมาย",
        )
    if facts.subject_group:
        xml = _replace_text_node_pattern_once(
            xml,
            r"กลุ่มสาระการเรียนรู้\.+",
            f"กลุ่มสาระการเรียนรู้{facts.subject_group}",
        )

    for marker, fill_value in _part_3_fill_values(readiness).items():
        if fill_value:
            xml = _append_to_marker_text_node_once(xml, marker, fill_value)

    return xml.encode("utf-8")


def _extract_template_facts(readiness: PsarReadinessResponse) -> PsarTemplateFacts:
    text = "\n".join(
        mapping.extracted_summary
        for mapping in readiness.mapped_files
        if mapping.status != "rejected" and mapping.extracted_summary
    )
    teacher_name = _first_regex_group(
        text,
        (
            r"(?:ผู้รายงาน|ข้าพเจ้า)\s*((?:นาย|นางสาว|นาง|ดร\.|ว่าที่ร้อยตรี)\s*[\u0e00-\u0e7f]+(?:\s+[\u0e00-\u0e7f]+){1,3})",
        ),
    )
    subject_group = _first_regex_group(
        text,
        (
            r"กลุ่มสาระการเรียนรู้\s*([A-Za-z0-9\u0e00-\u0e7f /().-]{2,80})",
        ),
    )
    clean_subject_group = _clean_template_fact(subject_group)
    if clean_subject_group and any(term in clean_subject_group for term in ("โรงเรียน", "สำนักงาน", "กระทรวง")):
        clean_subject_group = None
    return PsarTemplateFacts(
        teacher_name=_clean_template_fact(teacher_name),
        subject_group=clean_subject_group,
    )


def _part_3_fill_values(readiness: PsarReadinessResponse) -> dict[str, str]:
    learner = _recommendation_lines_for_section(readiness, "standard_1_learner_quality", limit=5)
    management = _recommendation_lines_for_section(readiness, "standard_2_management", limit=3)
    teaching = _recommendation_lines_for_section(readiness, "standard_3_teaching", limit=5)
    summary = _recommendation_lines_for_section(readiness, "summary_improvement", limit=3)

    return {
        "1.1.1": _line_or_empty(learner, 0),
        "1.1.2": _line_or_empty(learner, 1),
        "1.1.3": _line_or_empty(learner, 2),
        "1.2.1": _line_or_empty(management, 0),
        "1.3.1": _line_or_empty(teaching, 0),
        "1.3.2": _line_or_empty(teaching, 1),
        "2.1.1": _line_or_empty(learner, 3),
        "2.1.2": _line_or_empty(summary, 0),
        "2.2.1": _line_or_empty(management, 1),
        "2.2.2": _line_or_empty(management, 2),
        "3.1": _line_or_empty(teaching, 2),
        "3.2": _line_or_empty(teaching, 3),
        "3.3": _line_or_empty(teaching, 4),
    }


def _recommendation_lines_for_section(
    readiness: PsarReadinessResponse,
    section_id: str,
    *,
    limit: int,
) -> list[str]:
    lines: list[str] = []
    for item in readiness.missing_evidence_recommendations:
        if item.section_id != section_id:
            continue
        evidence_name = _thai_evidence_name(item.evidence_name)
        lines.append(f"จัดเตรียม/แนบ{evidence_name} เพื่อให้หลักฐานในหัวข้อนี้ครบถ้วน")
        if len(lines) >= limit:
            return lines
    return lines


def _thai_evidence_name(evidence_name: str) -> str:
    return THAI_EVIDENCE_LABELS.get(evidence_name, evidence_name)


def _line_or_empty(lines: Sequence[str], index: int) -> str:
    return lines[index] if index < len(lines) else ""


def _first_regex_group(text: str, patterns: Sequence[str]) -> str | None:
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    return None


def _clean_template_fact(value: str | None) -> str | None:
    if value is None:
        return None
    cleaned = re.sub(r"[._]+", " ", value)
    cleaned = " ".join(cleaned.split())
    return cleaned or None


def _replace_text_node_once(xml: str, old_text: str, new_text: str) -> str:
    return _replace_text_node_pattern_once(xml, re.escape(_xml_text(old_text)), new_text)


def _replace_text_node_text_once(xml: str, old_text: str, new_text: str) -> str:
    pattern = re.compile(r"(<w:t\b[^>]*>)(?P<content>.*?)(</w:t>)")
    old_xml = _xml_text(old_text)
    new_xml = _xml_text(new_text)

    def replace(match: re.Match[str]) -> str:
        content = match.group("content")
        if old_xml not in content:
            return match.group(0)
        return f"{match.group(1)}{content.replace(old_xml, new_xml, 1)}{match.group(3)}"

    return pattern.sub(replace, xml, count=1)


def _replace_text_node_pattern_once(xml: str, old_text_pattern: str, new_text: str) -> str:
    pattern = re.compile(rf"(<w:t\b[^>]*>){old_text_pattern}(</w:t>)")
    return pattern.sub(lambda match: f"{match.group(1)}{_xml_text(new_text)}{match.group(2)}", xml, count=1)


def _append_to_marker_text_node_once(xml: str, marker: str, value: str) -> str:
    marker_pattern = re.escape(_xml_text(marker))
    pattern = re.compile(rf"(<w:t\b[^>]*>)(?P<content>\s*{marker_pattern}\s*)(</w:t>)")
    def replace(match: re.Match[str]) -> str:
        content = match.group("content")
        separator = "" if content.endswith((" ", "\t", "\n", "\r")) else " "
        return f"{match.group(1)}{content}{separator}{_xml_text(value)}{match.group(3)}"

    return pattern.sub(replace, xml, count=1)


def _xml_text(value: str) -> str:
    return escape(_clean_xml_text(value))


def _paragraph(text: str, *, style: str) -> str:
    style_xml = "" if style == "Normal" else f'<w:pStyle w:val="{escape(style)}"/>'
    return (
        "<w:p>"
        f"<w:pPr>{style_xml}</w:pPr>"
        "<w:r>"
        "<w:t xml:space=\"preserve\">"
        f"{escape(_clean_xml_text(text))}"
        "</w:t>"
        "</w:r>"
        "</w:p>"
    )


def _docx_document(body_xml: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body>"
        f"{body_xml}"
        '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="720" w:footer="720" w:gutter="0"/></w:sectPr>'
        "</w:body>"
        "</w:document>"
    )


def _docx_content_types() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
        "</Types>"
    )


def _docx_relationships() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )


def _docx_document_relationships() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        "</Relationships>"
    )


def _docx_styles() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:rPr><w:rFonts w:ascii="Aptos" w:eastAsia="TH Sarabun New"/><w:sz w:val="24"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="240"/></w:pPr><w:rPr><w:b/><w:rFonts w:ascii="Aptos Display" w:eastAsia="TH Sarabun New"/><w:sz w:val="40"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Subtitle"><w:name w:val="Subtitle"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="240"/></w:pPr><w:rPr><w:i/><w:color w:val="666666"/><w:sz w:val="24"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="Heading 1"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:before="360" w:after="160"/></w:pPr><w:rPr><w:b/><w:sz w:val="32"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="Heading 2"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:before="240" w:after="120"/></w:pPr><w:rPr><w:b/><w:sz w:val="28"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading3"><w:name w:val="Heading 3"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:before="160" w:after="80"/></w:pPr><w:rPr><w:b/><w:sz w:val="24"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Warning"><w:name w:val="Warning"/><w:basedOn w:val="Normal"/><w:rPr><w:b/><w:color w:val="9A3412"/></w:rPr></w:style>'
        "</w:styles>"
    )


def _format_percent(value: float) -> str:
    return f"{round(max(0.0, min(1.0, value)) * 100)}%"


def _timestamp_slug(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z]+", "-", value).strip("-") or "generated"


def _clean_xml_text(value: str) -> str:
    return "".join(
        character
        for character in value
        if character in {"\t", "\n", "\r"} or ord(character) >= 32
    )


def _normalize_evidence_type(value: str) -> str:
    return re.sub(r"[^a-z0-9ก-๙]+", " ", value.casefold()).strip()
