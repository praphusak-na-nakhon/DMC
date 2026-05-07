from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from dmc_sidecar import config as sidecar_config
from dmc_sidecar.student_basic_info import _read_students, export_student_basic_info_form


DMC_HEADERS = [
    "รหัสโรงเรียน",
    "ชื่อโรงเรียน",
    "เลขประจำตัวนักเรียน",
    "ชั้น",
    "ห้อง",
    "เลขประจำตัวนักเรียน",
    "เพศ",
    "คำนำหน้าชื่อ",
    "ชื่อ",
    "นามสกุล",
    "วันเกิด",
    "อายุ(ปี)",
    "น้ำหนัก",
    "ส่วนสูง",
    "กลุ่มเลือด",
    "ศาสนา",
    "เชื้อชาติ",
    "สัญชาติ",
    "บ้านเลขที่",
    "หมู่",
    "ถนน/ซอย",
    "ตำบล",
    "อำเภอ",
    "จังหวัด",
    "ชื่อผู้ปกครอง",
    "นามสกุลผู้ปกครอง",
    "อาชีพของผู้ปกครอง",
    "ความเกี่ยวข้องของผู้ปกครองกับนักเรียน",
    "ชื่อบิดา",
    "นามสกุลบิดา",
    "อาชีพของบิดา",
    "ชื่อมารดา",
    "นามสกุลมารดา",
    "อาชีพของมารดา",
    "ความด้อยโอกาส",
    "ยังไม่สามารถจำหน่ายได้ (3.1.8)",
]


def _write_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "แบบฟอร์มนักเรียน"
    sheet.merge_cells("A1:I1")
    sheet.merge_cells("A2:E2")
    sheet.merge_cells("F2:G2")
    sheet.merge_cells("H2:I2")
    sheet.merge_cells("A3:I3")
    sheet["A1"] = "แบบฟอร์มบันทึกข้อมูลพื้นฐานนักเรียน"
    sheet["A2"] = "รายชื่อนักเรียนโรงเรียน..............................................................................."
    sheet["F2"] = "ปีการศึกษา........................."
    sheet["H2"] = "เทอม................"
    sheet["A3"] = "ชั้น................................................../................"
    sheet.append(
        [
            "ลำดับ",
            "ชื่อ-สกุล",
            "เลขบัตรประชาชน",
            "วัน/เดือน/ปีเกิด",
            "ที่อยู่",
            "ชื่อ-สกุลบิดา (ถ้ามี)",
            "ชื่อ-สกุลมารดา(ถ้ามี)",
            "น้ำหนัก",
            "ส่วนสูง",
        ]
    )
    for _ in range(25):
        sheet.append([None] * 9)
    workbook.save(path)


def _write_dmc_export(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["วันและเวลาที่สร้างรายงาน 2026-05-02 04:10:12.244315"])
    sheet.append(DMC_HEADERS)
    for index in range(1, 28):
        _append_student(sheet, index=index, level="ม.1", room=1)
    _append_student(sheet, index=28, level="ม.1", room=2)
    workbook.save(path)


def _append_student(sheet: object, *, index: int, level: str, room: int) -> None:
    sheet.append(  # type: ignore[attr-defined]
        [
            "81012017",
            "เหนือคลองประชาบำรุง",
            f"181990000{index:04d}",
            level,
            room,
            f"19{index:03d}",
            "ช",
            "เด็กชาย",
            f"ชื่อ{index}",
            f"สกุล{index}",
            "13/05/2555",
            13,
            40 + index,
            140 + index,
            "O",
            "พุทธ",
            "ไทย",
            "ไทย",
            f"{index}",
            "5",
            "-",
            "เหนือคลอง",
            "เหนือคลอง",
            "กระบี่",
            "ผู้ปกครอง",
            "ทดสอบ",
            "รับจ้าง",
            "บิดา",
            "พ่อ",
            f"สกุล{index}",
            "รับจ้าง",
            "แม่",
            f"สกุล{index}",
            "ค้าขาย",
            "-",
            "-",
        ]
    )


def test_export_student_basic_info_splits_by_class_and_extends_template_rows(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setattr(sidecar_config, "default_data_dir", lambda: tmp_path)
    source_path = tmp_path / "2568-3-student.xlsx"
    template_path = tmp_path / "student_basic_info_form.xlsx"
    _write_dmc_export(source_path)
    _write_template(template_path)

    result = export_student_basic_info_form(excel_path=source_path, template_path=template_path)

    assert result.module == "studentBasicInfo"
    assert result.school_name == "เหนือคลองประชาบำรุง"
    assert result.school_year == "2568"
    assert result.term == "2"
    assert result.rows_total == 28
    assert result.students_exported == 28
    assert result.classes_exported == 2
    assert [item.sheet_name for item in result.classes] == ["ม.1-1", "ม.1-2"]

    workbook = load_workbook(result.output_path, data_only=True)
    assert workbook.sheetnames == ["ม.1-1", "ม.1-2"]

    first_class = workbook["ม.1-1"]
    assert first_class["A1"].value == "แบบฟอร์มบันทึกข้อมูลพื้นฐานนักเรียน(งานพยาบาล)"
    assert first_class["A2"].value == "รายชื่อนักเรียนโรงเรียนเหนือคลองประชาบำรุง"
    assert first_class["F2"].value == "ปีการศึกษา 2568"
    assert first_class["H2"].value == "เทอม 2"
    assert first_class["A3"].value == "ชั้น ม.1/1"
    assert first_class["A5"].value == 1
    assert first_class["B5"].value == "เด็กชาย ชื่อ1 สกุล1"
    assert first_class["C5"].value == "1819900000001"
    assert "บ้านเลขที่ 1" in first_class["E5"].value
    assert first_class["F5"].value == "พ่อ สกุล1"
    assert first_class["G5"].value == "แม่ สกุล1"
    assert first_class["A31"].value == 27
    assert first_class["B31"].value == "เด็กชาย ชื่อ27 สกุล27"
    assert first_class["A1"].font.name == "TH Sarabun New"
    assert first_class["A1"].font.sz == 16
    assert first_class["B5"].font.name == "TH Sarabun New"
    assert first_class["B5"].font.sz == 16
    assert first_class["A1"].alignment.wrap_text is not True
    assert first_class["E5"].alignment.wrap_text is not True

    second_class = workbook["ม.1-2"]
    assert second_class["A3"].value == "ชั้น ม.1/2"
    assert second_class["A5"].value == 1
    assert second_class["B5"].value == "เด็กชาย ชื่อ28 สกุล28"


def test_read_students_streams_large_dmc_export(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setattr(sidecar_config, "default_data_dir", lambda: tmp_path)
    source_path = tmp_path / "2568-3-student-large.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["generated 2026-05-02"])
    sheet.append(DMC_HEADERS)
    for index in range(1, 601):
        _append_student(sheet, index=index, level="ม.1", room=1)
    workbook.save(source_path)

    students, rows_total = _read_students(source_path)

    assert rows_total == 600
    assert len(students) == 600
    assert students[0].citizen_id == "1819900000001"
    assert students[-1].citizen_id == "1819900000600"
