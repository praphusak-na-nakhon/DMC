from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]

from dmc_sidecar import config
from dmc_sidecar.checkpoint import JobCheckpoint
from dmc_sidecar.job_completion_summary import build_completion_summary, write_completion_summary_workbook
from dmc_sidecar.job_store import JobStore


def test_build_completion_summary_counts_success_and_failure() -> None:
    summary = build_completion_summary(
        [
            {
                "row_index": 3,
                "record_id": "record-1",
                "student_no": "1001",
                "full_name": "Student One",
                "status": "success",
                "note": "dry_run",
                "applied": False,
            },
            {
                "page": 1,
                "portal_row_index": 2,
                "portal_name": "Student Two",
                "portal_student_no": "1002",
                "note": "no_match",
                "applied": False,
            },
            {
                "page": 1,
                "portal_row_index": 3,
                "portal_name": "Student Three",
                "portal_student_no": "1003",
                "note": "filled",
                "applied": True,
            },
            {
                "row_index": 4,
                "full_name": "Student Four",
                "note": "submitted",
                "applied": False,
            },
        ]
    )

    assert summary["total"] == 4
    assert summary["succeeded"] == 3
    assert summary["failed"] == 1
    assert [item["full_name"] for item in summary["success_items"]] == ["Student One", "Student Three", "Student Four"]
    assert summary["failure_items"][0]["full_name"] == "Student Two"


def test_write_completion_summary_workbook(tmp_path: Path) -> None:
    path = write_completion_summary_workbook(
        tmp_path / "summary.xlsx",
        job_id="job-1",
        module="currentStudents",
        source_file="source.xlsx",
        results=[
            {
                "row_index": 1,
                "record_id": "record-1",
                "student_no": "1001",
                "full_name": "Student One",
                "status": "success",
                "note": "dry_run",
                "applied": False,
            }
        ],
    )

    workbook = load_workbook(path)

    assert workbook.sheetnames == ["Summary", "Success", "Failed", "All Records"]
    assert workbook["Summary"]["B5"].value == 1
    assert workbook["Success"]["E2"].value == "Student One"


def test_job_store_exposes_completion_summary_and_excel_path(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(config, "default_data_dir", lambda: tmp_path)
    store = JobStore()
    job_id = "job-summary"
    store.create_pending_job(job_id, "currentStudents", "source.xlsx")
    checkpoint = JobCheckpoint.initial(level_label="DMC transfer-in", base_url="https://example.test")
    checkpoint.processed = 2
    checkpoint.succeeded = 1
    checkpoint.failed = 1
    store.mark_running(job_id, total_records=2, checkpoint=checkpoint, started_at="2026-05-17T00:00:00Z")
    store.append_results(
        job_id,
        [
            {
                "page": 1,
                "portal_row_index": 1,
                "row_index": 1,
                "record_id": "record-1",
                "student_no": "1001",
                "full_name": "Student One",
                "status": "success",
                "note": "dry_run",
                "applied": False,
            },
            {
                "page": 2,
                "portal_row_index": 2,
                "row_index": 2,
                "record_id": "record-2",
                "student_no": "1002",
                "full_name": "Student Two",
                "status": "review",
                "note": "dmc_validation_error",
                "applied": False,
            },
        ],
    )

    store.mark_done(job_id, checkpoint=checkpoint, finished_at="2026-05-17T00:01:00Z")
    status = store.get_status(job_id)

    assert status is not None
    assert status["completion_summary"]["succeeded"] == 1
    assert status["completion_summary"]["failed"] == 1
    assert status["completion_summary"]["success_items"][0]["full_name"] == "Student One"
    assert status["completion_summary"]["failure_items"][0]["full_name"] == "Student Two"
    assert status["summary_report_path"] is not None
    assert Path(status["summary_report_path"]).exists()
