from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from dmc_sidecar.current_students import (
    CanonicalStudentRecord,
    CurrentStudentField,
    DmcFormMatchConfirmation,
    DmcTransferInImportRecord,
    ExportDmcFormJsonRequest,
    ExportCurrentStudentsBlankFormRequest,
    ExportCurrentStudentsImportExcelRequest,
    PreviewDmcFormJsonRequest,
    ReconcileCurrentStudentsRequest,
    ValidateCurrentStudentsImportFormRequest,
    export_dmc_form_json,
    export_current_students_blank_form,
    export_current_students_import_excel,
    load_dmc_transfer_in_import_records,
    preview_dmc_form_json,
    read_civil_registration_markdown_records,
    read_ocr_markdown,
    read_ocr_records,
    read_student_roster,
    reconcile_current_students,
    validate_current_students_import_form,
    _dmc_form_values,
)
from dmc_sidecar.checkpoint import JobCheckpoint
from dmc_sidecar.modules.current_students import CurrentStudentsModule, _today_buddhist_date_text
from dmc_sidecar.rpc import RpcServer


def _ocr_field(value: str | None) -> CurrentStudentField:
    return CurrentStudentField(value=value, source="ocr_form", confidence="review", raw_value=value)


def _write_roster(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1.1"
    sheet["A1"] = "รายชื่อนักเรียนชั้นมัธยมศึกษาปีที่ 1/1 ปีการศึกษา 2569"
    sheet["A2"] = "โรงเรียนเหนือคลองประชาบำรุง"
    sheet.append(["เลขที่", "เลขประจำตัว", None, "ชื่อ - สกุล"])
    sheet.append([1, 19984, "เด็กหญิง", "กัญญาณัฐ  กุลดี"])
    sheet.append([2, 19985, "เด็กชาย", "ภูริณัฐ  หาเหม"])
    sheet.append([3, 19986, "เด็กหญิง", "กนกวรรณ  นวลปั้น"])
    sheet.append([4, 19987, "เด็กชาย", "อนุวัฒน์  เดชอุดม"])
    workbook.save(path)


def _thai_id_row(
    *,
    citizen_id: str,
    prefix: str,
    first_name: str,
    last_name: str,
    row_no: int,
) -> list[str]:
    row = [""] * 45
    row[0] = "2026-03-19"
    row[1] = "08:39:05"
    row[2] = f"{row_no:06d}"
    row[3] = "ลงบันทึก"
    row[6] = f'="{citizen_id}"' if citizen_id.isdigit() else citizen_id
    row[7] = prefix
    row[8] = first_name
    row[10] = last_name
    row[15] = "28"
    row[16] = "4"
    row[20] = "คลองขนาน"
    row[21] = "เหนือคลอง"
    row[22] = "กระบี่"
    row[23] = "ชาย" if prefix == "ด.ช." else "หญิง"
    row[24] = "13 กุมภาพันธ์ 2557"
    return row


def _write_thai_id_csv(path: Path) -> None:
    header = [f"col-{index}" for index in range(45)]
    rows = [
        _thai_id_row(citizen_id="1819900905157", prefix="ด.ญ.", first_name="กัญญาณัฐ", last_name="กุลดี", row_no=1),
        _thai_id_row(citizen_id="1819900915055", prefix="ด.ช.", first_name="ภุริณัฐ", last_name="หาเหม", row_no=2),
        _thai_id_row(citizen_id='="18199..968591"', prefix="ด.ญ.", first_name="กนกวรรณ", last_name="นวลปั้น", row_no=3),
        _thai_id_row(citizen_id="1810800164408", prefix="ด.ช.", first_name="อนุวัฒน์", last_name="เดชอุดม", row_no=4),
        _thai_id_row(citizen_id="1810800164408", prefix="ด.ช.", first_name="อนุวัฒน์", last_name="เดชอุดม", row_no=5),
        _thai_id_row(citizen_id="1819900943555", prefix="ด.ช.", first_name="ปารมี", last_name="เชาว์ช่าง", row_no=6),
    ]
    lines = [",".join(header), *(",".join(row) for row in rows)]
    path.write_bytes(("\n".join(lines) + "\n").encode("cp874"))


def test_current_students_dry_run_opens_history_and_fills_it(monkeypatch: pytest.MonkeyPatch) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace(url="https://portal.test/studentin/add_cif")
    calls: list[str] = []
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20018",
        level_dtl_code="12",
        classroom="3",
        citizen_id="1810800164491",
        full_name="Test Student",
        dmc_form_values={"firstNameTh": "Test", "lastNameTh": "Student"},
    )

    def fake_open_transfer_form(**_kwargs: object) -> SimpleNamespace:
        calls.append("open_transfer_form")
        return page

    def fake_fill_transfer_form(_page: SimpleNamespace, _record: DmcTransferInImportRecord) -> None:
        calls.append("fill_transfer_form")

    def fake_submit_transfer_form(
        _page: SimpleNamespace,
        _record: DmcTransferInImportRecord,
        *,
        dry_run: bool,
    ) -> dict[str, object]:
        calls.append("submit_transfer_form")
        assert dry_run is True
        page.url = "https://portal.test/studentin/add"
        return {
            "applied": False,
            "note": "history_filled_for_review",
            "status": "review",
            "message": "History form filled.",
        }

    monkeypatch.setattr(module, "_open_transfer_form", fake_open_transfer_form)
    monkeypatch.setattr(module, "_fill_transfer_form", fake_fill_transfer_form)
    monkeypatch.setattr(module, "_submit_transfer_form", fake_submit_transfer_form)

    result = module._process_record(
        page=page,  # type: ignore[arg-type]
        record=record,
        dry_run=True,
        context=object(),  # type: ignore[arg-type]
        checkpoint=JobCheckpoint.initial(level_label="DMC transfer-in", base_url="https://portal.test"),
        record_number=1,
    )

    assert calls == ["open_transfer_form", "fill_transfer_form", "submit_transfer_form"]
    assert result["status"] == "success"
    assert result["note"] == "dry_run"
    assert "student history form" in result["message"]
    assert result["page_url"] == "https://portal.test/studentin/add"


def test_current_students_real_import_submits_history_form(monkeypatch: pytest.MonkeyPatch) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace(url="https://portal.test/studentin/add_cif")
    calls: list[str] = []
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20018",
        level_dtl_code="12",
        classroom="3",
        citizen_id="1810800164491",
        full_name="Test Student",
        dmc_form_values={"firstNameTh": "Test", "lastNameTh": "Student"},
    )

    def fake_open_transfer_form(**_kwargs: object) -> SimpleNamespace:
        calls.append("open_transfer_form")
        return page

    def fake_fill_transfer_form(_page: SimpleNamespace, _record: DmcTransferInImportRecord) -> None:
        calls.append("fill_transfer_form")

    def fake_submit_transfer_form(
        _page: SimpleNamespace,
        _record: DmcTransferInImportRecord,
        *,
        dry_run: bool,
    ) -> dict[str, object]:
        calls.append("submit_transfer_form")
        assert dry_run is False
        page.url = "https://portal.test/studentin/"
        return {
            "applied": True,
            "note": "submitted",
            "status": "success",
            "message": "DMC transfer-in history form was saved.",
        }

    monkeypatch.setattr(module, "_open_transfer_form", fake_open_transfer_form)
    monkeypatch.setattr(module, "_fill_transfer_form", fake_fill_transfer_form)
    monkeypatch.setattr(module, "_submit_transfer_form", fake_submit_transfer_form)

    result = module._process_record(
        page=page,  # type: ignore[arg-type]
        record=record,
        dry_run=False,
        context=object(),  # type: ignore[arg-type]
        checkpoint=JobCheckpoint.initial(level_label="DMC transfer-in", base_url="https://portal.test"),
        record_number=1,
    )

    assert calls == ["open_transfer_form", "fill_transfer_form", "submit_transfer_form"]
    assert result["status"] == "success"
    assert result["note"] == "submitted"
    assert result["applied"] is True
    assert result["page_url"] == "https://portal.test/studentin/"


def test_current_students_submit_transfer_form_saves_history_when_not_dry_run(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace(url="https://portal.test/studentin/add")
    calls: list[str] = []
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20018",
        level_dtl_code="12",
        classroom="3",
        citizen_id="1810800164491",
        full_name="Test Student",
        dmc_form_values={"firstNameTh": "Test", "lastNameTh": "Student"},
    )

    monkeypatch.setattr(module, "_is_history_form", lambda _page: True)

    def fake_fill_history(_page: SimpleNamespace, _record: DmcTransferInImportRecord) -> None:
        calls.append("fill_history")

    def fake_submit_history(_page: SimpleNamespace) -> dict[str, object]:
        calls.append("submit_history")
        return {
            "applied": True,
            "note": "submitted",
            "status": "success",
            "message": "DMC transfer-in history form was saved.",
        }

    monkeypatch.setattr(module, "_fill_student_history_form", fake_fill_history)
    monkeypatch.setattr(module, "_submit_student_history_form", fake_submit_history)

    result = module._submit_transfer_form(page, record, dry_run=False)  # type: ignore[arg-type]

    assert calls == ["fill_history", "submit_history"]
    assert result["applied"] is True
    assert result["note"] == "submitted"
    assert result["status"] == "success"


def test_current_students_submit_transfer_form_dry_run_still_stops_before_history_save(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace(url="https://portal.test/studentin/add")
    calls: list[str] = []
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20018",
        level_dtl_code="12",
        classroom="3",
        citizen_id="1810800164491",
        full_name="Test Student",
        dmc_form_values={"firstNameTh": "Test", "lastNameTh": "Student"},
    )

    monkeypatch.setattr(module, "_is_history_form", lambda _page: True)

    def fake_fill_history(_page: SimpleNamespace, _record: DmcTransferInImportRecord) -> None:
        calls.append("fill_history")

    def fake_submit_history(_page: SimpleNamespace) -> dict[str, object]:
        calls.append("submit_history")
        return {"applied": True, "note": "submitted", "status": "success", "message": "saved"}

    monkeypatch.setattr(module, "_fill_student_history_form", fake_fill_history)
    monkeypatch.setattr(module, "_submit_student_history_form", fake_submit_history)

    result = module._submit_transfer_form(page, record, dry_run=True)  # type: ignore[arg-type]

    assert calls == ["fill_history"]
    assert result["applied"] is False
    assert result["note"] == "history_filled_for_review"
    assert result["status"] == "review"


def test_current_students_submit_transfer_form_falls_back_to_new_student_when_citizen_not_found(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    calls: list[str] = []

    class FakeNavigation:
        def __enter__(self) -> None:
            calls.append("expect_navigation")

        def __exit__(self, *_args: object) -> None:
            return None

    class FakeLocator:
        def click(self, *, timeout: int) -> None:
            calls.append(f"click:{timeout}")

    class FakePage:
        url = "https://portal.test/studentin/add_cif"

        def expect_navigation(self, *, wait_until: str, timeout: int) -> FakeNavigation:
            calls.append(f"wait_until:{wait_until}:{timeout}")
            return FakeNavigation()

        def locator(self, selector: str) -> FakeLocator:
            calls.append(f"locator:{selector}")
            return FakeLocator()

        def wait_for_timeout(self, ms: int) -> None:
            calls.append(f"timeout:{ms}")

        def wait_for_load_state(self, *, timeout: int) -> None:
            calls.append(f"load_state:{timeout}")

    page = FakePage()
    transfer_error = (
        "ไม่สามารถบันทึกได้เนื่องจากข้อมูลยังไม่ครบ | เลขประจำตัวประชาชน* "
        "ไม่พบเลขประจำตัวประชาชน ให้กรอกข้อมูลใหม่ในหน้าเพิ่มนักเรียน (2.7.3)"
    )
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20018",
        level_dtl_code="12",
        classroom="3",
        citizen_id="1810800164491",
        full_name="Test Student",
        dmc_form_values={"firstNameTh": "Test", "lastNameTh": "Student"},
    )

    monkeypatch.setattr(module, "_is_history_form", lambda _page: False)
    monkeypatch.setattr(module, "_extract_error_text", lambda _page: transfer_error)

    def fake_new_student_form(
        _page: FakePage,
        _record: DmcTransferInImportRecord,
        *,
        dry_run: bool,
        transfer_error_text: str,
    ) -> dict[str, object]:
        calls.append("new_student_form")
        assert dry_run is False
        assert transfer_error_text == transfer_error
        return {
            "applied": True,
            "note": "submitted",
            "status": "success",
            "message": "saved from studentnew/add",
            "fallback": "studentnew_add",
            "transfer_error": transfer_error_text,
        }

    monkeypatch.setattr(module, "_submit_new_student_form", fake_new_student_form)

    result = module._submit_transfer_form(page, record, dry_run=False)  # type: ignore[arg-type]

    assert "new_student_form" in calls
    assert result["applied"] is True
    assert result["status"] == "success"
    assert result["fallback"] == "studentnew_add"
    assert result["transfer_error"] == transfer_error


def test_current_students_submit_new_student_form_fills_history_form(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    calls: list[str] = []

    class FakePage:
        url = "https://portal.test/studentin/add_cif"

        def goto(self, url: str, *, wait_until: str, timeout: int) -> None:
            calls.append(f"goto:{url}:{wait_until}:{timeout}")
            self.url = url

        def wait_for_timeout(self, ms: int) -> None:
            calls.append(f"timeout:{ms}")

    page = FakePage()
    transfer_error = "ไม่พบเลขประจำตัวประชาชน ให้กรอกข้อมูลใหม่ในหน้าเพิ่มนักเรียน (2.7.3)"
    field_actions = [{"field_name": "fatherFirstNameTh", "action": "filled_blank", "existing_value": None, "incoming_value": "Father"}]
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20018",
        level_dtl_code="12",
        classroom="3",
        citizen_id="1810800164491",
        full_name="Test Student",
        dmc_form_values={"firstNameTh": "Test", "lastNameTh": "Student"},
    )

    monkeypatch.setattr(module, "_is_history_form", lambda _page: True)
    monkeypatch.setattr(module, "_fill_student_history_form", lambda _page, _record: field_actions)
    monkeypatch.setattr(
        module,
        "_submit_student_history_form",
        lambda _page: {
            "applied": True,
            "note": "submitted",
            "status": "success",
            "message": "DMC new-student form was saved.",
        },
    )

    result = module._submit_new_student_form(
        page,  # type: ignore[arg-type]
        record,
        dry_run=False,
        transfer_error_text=transfer_error,
    )

    assert calls == [
        "goto:https://portal.bopp-obec.info/obec69/studentnew/add:domcontentloaded:90000",
        "timeout:700",
    ]
    assert result["applied"] is True
    assert result["note"] == "submitted"
    assert result["fallback"] == "studentnew_add"
    assert result["transfer_error"] == transfer_error
    assert result["field_actions"] == field_actions


def test_current_students_new_student_defaults_fill_missing_parent_and_postal_fields() -> None:
    module = CurrentStudentsModule()
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20000",
        level_dtl_code="12",
        classroom="3",
        citizen_id="1819900918968",
        full_name="Test Student",
        dmc_form_values={
            "firstNameTh": "Test",
            "lastNameTh": "Student",
            "fatherFirstNameTh": "",
            "motherFirstNameTh": "Existing Mother",
            "motherLastNameTh": "Existing Last",
        },
    )

    next_record = module._with_new_student_defaults(record)
    values = next_record.dmc_form_values

    assert values["psPostalCode"] == "-"
    assert values["psHomeIdNo"] == "-"
    assert values["homeIdNo"] == "-"
    assert values["postalCode"] == "-"
    assert values["rubberDt"] == "10000.0"
    assert values["childIndex"] == "1"
    assert values["fatherCifNo"] == "-"
    assert values["fatherCifType"] == "O"
    assert values["fatherTitleCode"] == "003"
    assert values["fatherFirstNameTh"] == "-"
    assert values["fatherLastNameTh"] == "-"
    assert values["fatherOccupationCode"] == "5"
    assert values["motherTitleCode"] == "004"
    assert values["motherFirstNameTh"] == "Existing Mother"
    assert values["motherLastNameTh"] == "Existing Last"
    assert values["motherOccupationCode"] == "5"
    assert values["parentFamilyRelationCode"] == "02"
    assert values["parentTitleCode"] == "004"
    assert values["parentFirstNameTh"] == "Existing Mother"
    assert values["parentLastNameTh"] == "Existing Last"
    assert values["parentOccupationCode"] == "5"
    assert values["journeyTypeCode"] == "02"
    assert values["timeDt"] == "10.0"
    assert "psPostalCode" not in record.dmc_form_values


def test_current_students_new_student_defaults_mirror_partial_house_and_postal_values() -> None:
    module = CurrentStudentsModule()
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20005",
        level_dtl_code="10",
        classroom="3",
        citizen_id="1819900924402",
        full_name="เด็กหญิง พิรุฬกานต์ เพชรลูก",
        dmc_form_values={
            "firstNameTh": "พิรุฬกานต์",
            "lastNameTh": "เพชรลูก",
            "psHomeIdNo": "81010044089",
            "postalCode": "81130",
        },
    )

    values = module._with_new_student_defaults(record).dmc_form_values

    assert values["psHomeIdNo"] == "81010044089"
    assert values["homeIdNo"] == "81010044089"
    assert values["psPostalCode"] == "81130"
    assert values["postalCode"] == "81130"


def test_current_students_new_student_defaults_copy_parent_from_declared_relation() -> None:
    module = CurrentStudentsModule()
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20002",
        level_dtl_code="10",
        classroom="3",
        citizen_id="1819900923473",
        full_name="เด็กหญิง ปาริศา ไกรนรา",
        dmc_form_values={
            "parentFamilyRelationCode": "01",
            "fatherCifNo": "1234567890123",
            "fatherCifType": "I",
            "fatherTitleCode": "003",
            "fatherFirstNameTh": "สุริน",
            "fatherLastNameTh": "ไกรนรา",
            "fatherOccupationCode": "5",
            "fatherSalary": "0.0",
            "fatherTelNo": "-",
            "motherTitleCode": "004",
            "motherFirstNameTh": "ศิรากาด",
            "motherLastNameTh": "แป้นด้วง",
        },
    )

    values = module._with_new_student_defaults(record).dmc_form_values

    assert values["parentFamilyRelationCode"] == "01"
    assert values["parentCifNo"] == "1234567890123"
    assert values["parentCifType"] == "I"
    assert values["parentTitleCode"] == "003"
    assert values["parentFirstNameTh"] == "สุริน"
    assert values["parentLastNameTh"] == "ไกรนรา"
    assert values["parentOccupationCode"] == "5"


def test_current_students_new_student_defaults_fill_required_parent_when_all_parent_sources_missing() -> None:
    module = CurrentStudentsModule()
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20000",
        level_dtl_code="10",
        classroom="3",
        citizen_id="1819900918968",
        full_name="เด็กชาย นฤคณากร เกตุทอง",
        dmc_form_values={"firstNameTh": "นฤคณากร", "lastNameTh": "เกตุทอง"},
    )

    values = module._with_new_student_defaults(record).dmc_form_values

    assert values["parentFamilyRelationCode"] == "20"
    assert values["parentCifNo"] == "-"
    assert values["parentCifType"] == "O"
    assert values["parentTitleCode"] == "003"
    assert values["parentFirstNameTh"] == "-"
    assert values["parentLastNameTh"] == "-"
    assert values["parentOccupationCode"] == "5"


def test_current_students_new_student_defaults_normalize_title_by_level_and_gender() -> None:
    module = CurrentStudentsModule()
    m1_record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20000",
        level_dtl_code="10",
        classroom="3",
        citizen_id="1819900918968",
        full_name="Test Student",
        dmc_form_values={"genderCode": "F", "titleCode": "004"},
    )

    m1_values = module._with_new_student_defaults(m1_record).dmc_form_values

    assert m1_values["titleCode"] == "002"
    assert m1_values["genderCode"] == "F"

    m4_record = m1_record.model_copy(
        update={
            "level_dtl_code": "13",
            "dmc_form_values": {"genderCode": "M", "titleCode": "001"},
        }
    )

    m4_values = module._with_new_student_defaults(m4_record).dmc_form_values

    assert m4_values["titleCode"] == "003"
    assert m4_values["genderCode"] == "M"


def test_current_students_new_student_form_fills_title_after_generic_values(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace(
        url="https://portal.bopp-obec.info/obec69/studentnew/add",
        wait_for_timeout=lambda _ms: None,
    )
    filled_values: dict[str, object] = {}
    title_codes: list[str] = []
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20000",
        level_dtl_code="10",
        classroom="3",
        citizen_id="1819900918968",
        full_name="Test Student",
        dmc_form_values={"firstNameTh": "Test", "lastNameTh": "Student", "genderCode": "F", "titleCode": "004"},
    )

    monkeypatch.setattr(module, "_wait_for_history_form", lambda _page: None)
    monkeypatch.setattr(module, "_fill_post_date", lambda _page, _value: None)
    monkeypatch.setattr(module, "_fill_admission_date", lambda _page, _value: None)
    monkeypatch.setattr(module, "_fill_address_chain", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(module, "_preserve_existing_parent_name_fields", lambda _page, _values: [])
    monkeypatch.setattr(module, "_fill_form_values", lambda _page, values: filled_values.update(values))
    monkeypatch.setattr(module, "_fill_title_code", lambda _page, title_code: title_codes.append(title_code))

    module._fill_student_history_form(page, record)  # type: ignore[arg-type]

    assert "titleCode" not in filled_values
    assert filled_values["journeyTypeCode"] == "02"
    assert filled_values["timeDt"] == "10.0"
    assert filled_values["genderCode"] == "F"
    assert title_codes == ["002"]


def test_current_students_submit_student_history_form_clicks_final_save(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    calls: list[str] = []

    class FakeNavigation:
        def __enter__(self) -> None:
            calls.append("expect_navigation")

        def __exit__(self, *_args: object) -> None:
            return None

    class FakeLocator:
        def click(self, *, timeout: int) -> None:
            calls.append(f"click:{timeout}")

    class FakePage:
        url = "https://portal.test/studentin/"

        def expect_navigation(self, *, wait_until: str, timeout: int) -> FakeNavigation:
            calls.append(f"wait_until:{wait_until}:{timeout}")
            return FakeNavigation()

        def locator(self, selector: str) -> FakeLocator:
            calls.append(f"locator:{selector}")
            return FakeLocator()

        def wait_for_timeout(self, ms: int) -> None:
            calls.append(f"timeout:{ms}")

    page = FakePage()
    monkeypatch.setattr(module, "_extract_error_text", lambda _page: "")
    monkeypatch.setattr(module, "_extract_success_text", lambda _page: "")
    monkeypatch.setattr(module, "_is_history_form", lambda _page: False)

    result = module._submit_student_history_form(page)  # type: ignore[arg-type]

    assert calls == [
        "wait_until:domcontentloaded:15000",
        "expect_navigation",
        'locator:input[name="submit"]',
        "click:10000",
        "timeout:800",
    ]
    assert result["applied"] is True
    assert result["note"] == "submitted"
    assert result["status"] == "success"


def test_current_students_submit_student_history_form_accepts_success_message_on_same_page(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    calls: list[str] = []

    class FakeNavigation:
        def __enter__(self) -> None:
            calls.append("expect_navigation")

        def __exit__(self, *_args: object) -> None:
            return None

    class FakeLocator:
        def click(self, *, timeout: int) -> None:
            calls.append(f"click:{timeout}")

    class FakePage:
        url = "https://portal.test/studentin/add"

        def expect_navigation(self, *, wait_until: str, timeout: int) -> FakeNavigation:
            calls.append(f"wait_until:{wait_until}:{timeout}")
            return FakeNavigation()

        def locator(self, selector: str) -> FakeLocator:
            calls.append(f"locator:{selector}")
            return FakeLocator()

        def wait_for_timeout(self, ms: int) -> None:
            calls.append(f"timeout:{ms}")

    page = FakePage()
    monkeypatch.setattr(module, "_extract_error_text", lambda _page: "")
    monkeypatch.setattr(module, "_extract_success_text", lambda _page: "บันทึกข้อมูลเรียบร้อย")
    monkeypatch.setattr(module, "_is_history_form", lambda _page: True)

    result = module._submit_student_history_form(page)  # type: ignore[arg-type]

    assert calls == [
        "wait_until:domcontentloaded:15000",
        "expect_navigation",
        'locator:input[name="submit"]',
        "click:10000",
        "timeout:800",
    ]
    assert result["applied"] is True
    assert result["note"] == "submitted"
    assert result["status"] == "success"
    assert result["message"] == "บันทึกข้อมูลเรียบร้อย"


def test_current_students_address_chain_waits_before_selecting_children(monkeypatch: pytest.MonkeyPatch) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace()
    calls: list[tuple[str, ...]] = []

    def fake_set(_page: SimpleNamespace, name: str, value: str, *, trigger_change: bool = True) -> None:
        calls.append(("set", name, value, str(trigger_change)))

    def fake_wait(_page: SimpleNamespace, name: str, value: str) -> bool:
        calls.append(("wait", name, value))
        return True

    def fake_populate(_page: SimpleNamespace, *, child_name: str, remote_method: str, parent_code: str) -> bool:
        calls.append(("populate", child_name, remote_method, parent_code))
        return True

    monkeypatch.setattr(module, "_set_named_field_value", fake_set)
    monkeypatch.setattr(module, "_wait_for_select_option", fake_wait)
    monkeypatch.setattr(module, "_populate_address_child_options", fake_populate)

    module._fill_address_chain(  # type: ignore[arg-type]
        page,
        {
            "psProvinceCode": "81000000",
            "psAmphurCode": "81080000",
            "psTumbolCode": "81080300",
        },
        province_name="psProvinceCode",
        district_name="psAmphurCode",
        subdistrict_name="psTumbolCode",
    )

    assert calls == [
        ("set", "psProvinceCode", "81000000", "False"),
        ("populate", "psAmphurCode", "getAmphurListByProvinceCode", "81000000"),
        ("wait", "psAmphurCode", "81080000"),
        ("set", "psAmphurCode", "81080000", "False"),
        ("populate", "psTumbolCode", "getTumbolListByAmphurCode", "81080000"),
        ("wait", "psTumbolCode", "81080300"),
        ("set", "psTumbolCode", "81080300", "False"),
        ("set", "psAmphurCode", "81080000", "False"),
    ]


def test_current_students_address_chain_does_not_trigger_dmc_change_handlers(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace()
    calls: list[tuple[str, ...]] = []

    def fake_set(_page: SimpleNamespace, name: str, value: str, *, trigger_change: bool = True) -> None:
        calls.append(("set", name, value, str(trigger_change)))

    def fake_wait(_page: SimpleNamespace, name: str, value: str) -> bool:
        calls.append(("wait", name, value))
        return True

    monkeypatch.setattr(module, "_set_named_field_value", fake_set)
    monkeypatch.setattr(module, "_wait_for_select_option", fake_wait)
    monkeypatch.setattr(
        module,
        "_populate_address_child_options",
        lambda _page, *, child_name, remote_method, parent_code: calls.append(
            ("populate", child_name, remote_method, parent_code)
        )
        or True,
    )

    module._fill_address_chain(  # type: ignore[arg-type]
        page,
        {
            "psProvinceCode": "81000000",
            "psAmphurCode": "81040000",
            "psTumbolCode": "81040400",
        },
        province_name="psProvinceCode",
        district_name="psAmphurCode",
        subdistrict_name="psTumbolCode",
    )

    assert calls == [
        ("set", "psProvinceCode", "81000000", "False"),
        ("populate", "psAmphurCode", "getAmphurListByProvinceCode", "81000000"),
        ("wait", "psAmphurCode", "81040000"),
        ("set", "psAmphurCode", "81040000", "False"),
        ("populate", "psTumbolCode", "getTumbolListByAmphurCode", "81040000"),
        ("wait", "psTumbolCode", "81040400"),
        ("set", "psTumbolCode", "81040400", "False"),
        ("set", "psAmphurCode", "81040000", "False"),
    ]
    assert not any(call[0] == "set" and call[3] == "True" for call in calls)


def test_current_students_address_chain_stops_when_district_option_never_loads(monkeypatch: pytest.MonkeyPatch) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace(wait_for_timeout=lambda _ms: None)
    set_calls: list[tuple[str, str]] = []

    def fake_set(_page: SimpleNamespace, name: str, value: str, *, trigger_change: bool = True) -> None:
        set_calls.append((name, value))

    monkeypatch.setattr(module, "_set_named_field_value", fake_set)
    monkeypatch.setattr(module, "_wait_for_select_option", lambda *_args: False)
    monkeypatch.setattr(module, "_populate_address_child_options", lambda *_args, **_kwargs: True)

    module._fill_address_chain(  # type: ignore[arg-type]
        page,
        {
            "psProvinceCode": "81000000",
            "psAmphurCode": "81080000",
            "psTumbolCode": "81080300",
        },
        province_name="psProvinceCode",
        district_name="psAmphurCode",
        subdistrict_name="psTumbolCode",
    )

    assert set_calls == [
        ("psProvinceCode", "81000000"),
    ]


def test_current_students_address_chain_uses_existing_options_when_direct_load_fails(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace(wait_for_timeout=lambda _ms: None)
    set_calls: list[tuple[str, str, bool]] = []

    def fake_set(_page: SimpleNamespace, name: str, value: str, *, trigger_change: bool = True) -> None:
        set_calls.append((name, value, trigger_change))

    monkeypatch.setattr(module, "_set_named_field_value", fake_set)
    monkeypatch.setattr(module, "_wait_for_select_option", lambda *_args: True)
    monkeypatch.setattr(module, "_populate_address_child_options", lambda *_args, **_kwargs: False)

    module._fill_address_chain(  # type: ignore[arg-type]
        page,
        {
            "psProvinceCode": "81000000",
            "psAmphurCode": "81080000",
            "psTumbolCode": "81080300",
        },
        province_name="psProvinceCode",
        district_name="psAmphurCode",
        subdistrict_name="psTumbolCode",
    )

    assert set_calls == [
        ("psProvinceCode", "81000000", False),
        ("psAmphurCode", "81080000", False),
        ("psTumbolCode", "81080300", False),
        ("psAmphurCode", "81080000", False),
    ]


def test_current_students_today_buddhist_date_text_uses_buddhist_year() -> None:
    assert _today_buddhist_date_text(date(2026, 6, 3)) == "03/06/2569"


def test_current_students_fill_post_date_sets_post_date_field() -> None:
    module = CurrentStudentsModule()
    calls: list[tuple[str, str]] = []

    class FakePage:
        def evaluate(self, script: str, value: str) -> None:
            calls.append((script, value))

    module._fill_post_date(FakePage(), "03/06/2569")  # type: ignore[arg-type]

    assert len(calls) == 1
    assert '[name="postDate"]' in calls[0][0]
    assert calls[0][1] == "03/06/2569"


def test_current_students_history_form_preserves_existing_parent_names(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace(wait_for_timeout=lambda _ms: None)
    filled_values: dict[str, object] = {}
    post_date_values: list[str] = []
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20018",
        level_dtl_code="12",
        classroom="3",
        citizen_id="1810800164491",
        full_name="Test Student",
        dmc_form_values={
            "firstNameTh": "Test",
            "lastNameTh": "Student",
            "fatherFirstNameTh": "OCR Father",
            "fatherLastNameTh": "Existing Father Last",
            "motherFirstNameTh": "OCR Mother",
            "motherLastNameTh": "Existing Mother Last",
            "postDate": "01/01/2500",
        },
    )

    monkeypatch.setattr(module, "_wait_for_history_form", lambda _page: None)
    monkeypatch.setattr(module, "_fill_post_date", lambda _page, value: post_date_values.append(value))
    monkeypatch.setattr(module, "_fill_admission_date", lambda _page, _value: None)
    monkeypatch.setattr(module, "_fill_address_chain", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        module,
        "_read_named_field_values",
        lambda _page, _names: {
            "fatherFirstNameTh": "Existing Father",
            "fatherLastNameTh": "Existing Father Last",
            "motherFirstNameTh": "",
            "motherLastNameTh": "Existing Mother Last",
        },
    )

    def fake_fill_form_values(_page: SimpleNamespace, values: dict[str, object]) -> None:
        filled_values.update(values)

    monkeypatch.setattr(module, "_fill_form_values", fake_fill_form_values)

    actions = module._fill_student_history_form(page, record)  # type: ignore[arg-type]

    assert "fatherFirstNameTh" not in filled_values
    assert "fatherLastNameTh" not in filled_values
    assert filled_values["motherFirstNameTh"] == "OCR Mother"
    assert "motherLastNameTh" not in filled_values
    assert "postDate" not in filled_values
    assert len(post_date_values) == 1
    actions_by_field = {action["field_name"]: action for action in actions}
    assert actions_by_field["fatherFirstNameTh"]["action"] == "skipped_existing_conflict"
    assert actions_by_field["fatherFirstNameTh"]["existing_value"] == "Existing Father"
    assert actions_by_field["fatherFirstNameTh"]["incoming_value"] == "OCR Father"
    assert actions_by_field["fatherLastNameTh"]["action"] == "skipped_existing_same"
    assert actions_by_field["motherFirstNameTh"]["action"] == "filled_blank"
    assert actions_by_field["motherLastNameTh"]["action"] == "skipped_existing_same"


def test_current_students_submit_transfer_form_reports_existing_parent_name_conflicts(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = CurrentStudentsModule()
    page = SimpleNamespace(url="https://portal.test/studentin/add")
    field_actions = [
        {
            "field_name": "fatherFirstNameTh",
            "action": "skipped_existing_conflict",
            "existing_value": "Existing Father",
            "incoming_value": "OCR Father",
        }
    ]
    record = DmcTransferInImportRecord(
        row_index=3,
        record_id="record-1",
        student_no="20018",
        level_dtl_code="12",
        classroom="3",
        citizen_id="1810800164491",
        full_name="Test Student",
        dmc_form_values={"firstNameTh": "Test", "lastNameTh": "Student"},
    )

    monkeypatch.setattr(module, "_is_history_form", lambda _page: True)
    monkeypatch.setattr(module, "_fill_student_history_form", lambda _page, _record: field_actions)
    monkeypatch.setattr(
        module,
        "_submit_student_history_form",
        lambda _page: {
            "applied": True,
            "note": "submitted",
            "status": "success",
            "message": "DMC transfer-in history form was saved.",
        },
    )

    result = module._submit_transfer_form(page, record, dry_run=False)  # type: ignore[arg-type]

    assert result["applied"] is True
    assert result["status"] == "success"
    assert result["field_actions"] == field_actions
    assert result["field_conflicts"] == field_actions


def _write_ocr_markdown(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "หน้า 1",
                "เลขประจำตัวประชาชน* 1 8199 00905 15 7 เลขประจำตัวนักเรียน 19984 คำนำหน้านาม* ด.ญ. เพศ* หญิง ชื่อ* กัญญาณัฐ นามสกุล* กุลดี ชื่อ (อังกฤษ)* Kanyanat นามสกุล (อังกฤษ)* Kuldee วัน/เดือน/ปีเกิด* 13 กุมภาพันธ์ 2557 จังหวัดที่เกิด* กระบี่ สถานพยาบาลที่เกิด* ร.พ. เหนือคลอง กลุ่มเลือด* O เชื้อชาติ* ไทย สัญชาติ* ไทย ศาสนา* พุทธ",
                "๒. ที่อยู่อาศัย",
                "ตามทะเบียนบ้าน รหัสประจำบ้าน* 8101-004408-9 บ้านเลขที่* 28 หมู่ที่ (ถ้าไม่มีใส่ -) 4 ถนน (ถ้าไม่มีใส่ -) - จังหวัด* กระบี่ อำเภอ* เหนือคลอง ตำบล* คลองขนาน รหัสไปรษณีย์* 81130 หมายเลขโทรศัพท์บ้าน - มือถือ* -",
                "ที่อยู่ปัจจุบัน รหัสประจำบ้าน* 8101-004408-9 บ้านเลขที่* 28 หมู่ที่ (ถ้าไม่มีใส่ -) 4 ถนน (ถ้าไม่มีใส่ -) - จังหวัด* กระบี่ อำเภอ* เหนือคลอง ตำบล* คลองขนาน รหัสไปรษณีย์* 81130 หมายเลขโทรศัพท์บ้าน - มือถือ* -",
                "๓. รายละเอียดนักเรียน",
                "การเดินทางมาโรงเรียน* [ ] เดินเท้า [x] พาหนะไม่เสียค่าโดยสาร [ ] พาหนะเสียค่าโดยสาร ระยะทางจากบ้านมา ร.ร.* ทางน้ำ (กม.) - ถนนลูกรัง (กม.) - ถนนลาดยาง (กม.) 15 กม. รวมระยะเวลาการเดินทางมาโรงเรียน (นาที)* 20 นาที",
                "๔. สุขภาพ",
                "น้ำหนัก* 40 กิโลกรัม ส่วนสูง* 150 เซนติเมตร",
                "๕. ครอบครัว",
                "สถานภาพสมรสของบิดามารดา สถานภาพสมรส [ ] สมรส [x] แยกกันอยู่ ข้อมูลพี่น้อง จำนวนพี่ชาย* - คน จำนวนน้องชาย* - คน จำนวนพี่สาว* 1 คน จำนวนน้องสาว* - คน จำนวนพี่น้องที่ศึกษาอยู่ (ไม่รวมตัวนักเรียนเอง) 1 คน นักเรียนเป็นบุตรคนที่* 2",
                "ข้อมูลบิดา เลขบัตรประจำตัวประชาชนบิดา* 3410100748561 ชนิดบัตร* [x] บัตรประชาชน [ ] อื่นๆ ชื่อบิดา* นาย ชวลิต นามสกุล* ตุ้มดำ กลุ่มเลือดบิดา* - อาชีพ* รับจ้างทั่วไป รายได้ต่อเดือน(บาท)* 5000 - 6000 บาท หมายเลขโทรศัพท์บิดา* 063-839-5699",
                "ข้อมูลมารดา เลขบัตรประจำตัวประชาชนมารดา* - ชนิดบัตร* [ ] บัตรประชาชน [ ] อื่นๆ ชื่อมารดา* - นามสกุล* - กลุ่มเลือดมารดา* - อาชีพ* - รายได้ต่อเดือน(บาท)* - หมายเลขโทรศัพท์มารดา* -",
                "ข้อมูลผู้ปกครอง เลขบัตรประจำตัวประชาชนผู้ปกครอง* 3410100748561 ชนิดบัตร* [x] บัตรประชาชน [ ] อื่นๆ ชื่อผู้ปกครอง* นาย ชวลิต นามสกุล* ตุ้มดำ กลุ่มเลือดผู้ปกครอง* - อาชีพ* รับจ้างทั่วไป รายได้ต่อเดือน(บาท)* 5000 - 6000 บาท หมายเลขโทรศัพท์ผู้ปกครอง* 063-839-5699 ความเกี่ยวข้องของผู้ปกครองกับนักเรียน* บิดาและบุตร",
                "**หมายเหตุ",
            ]
        ),
        encoding="utf-8",
    )


def _write_civil_registration_markdown(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "สำเนาทะเบียนบ้าน",
                "เลขรหัสประจำบ้าน 8101-008408-9",
                "รายการที่อยู่ 28 หมู่ที่ 4 ตำบลคลองขนาน อำเภอเหนือคลอง จังหวัดกระบี่ ชื่อหมู่บ้าน -",
                "รายการบุคคลในบ้าน ชื่อ ด.ญ. กัญญาณัฐ กุลดี สัญชาติ ไทย เพศ หญิง",
                "เลขประจำตัวประชาชน 1-8199-00905-15-7 สถานภาพ ผู้อาศัย",
                "เกิดเมื่อ 13 ก.พ. 2557",
                "มารดาผู้ให้กำเนิด ชื่อ กรานณิภา ตุ้มดำ เลขประจำตัวประชาชน 1-7603-00002-25-6 สัญชาติ ไทย",
                "บิดาผู้ให้กำเนิด ชื่อ ชวลิต ตุ้มดำ เลขประจำตัวประชาชน 3-4101-00748-56-1 สัญชาติ ไทย",
            ]
        ),
        encoding="utf-8",
    )


def _write_structured_ocr_json(path: Path) -> None:
    path.write_text(
        json.dumps(
            {
                "schema_version": "dmc_ocr_structured.v1",
                "document_type": "dmc_form",
                "engine": "gemini",
                "model": "gemini-3.5-flash",
                "usage_metadata": {
                    "prompt_token_count": 1200,
                    "candidates_token_count": 300,
                    "total_token_count": 1500,
                    "input_tokens_per_page": 600.0,
                    "output_tokens_per_page": 150.0,
                    "total_tokens_per_page": 750.0,
                },
                "records": [
                    {
                        "record_type": "dmc_form",
                        "page_start": 1,
                        "page_end": 2,
                        "fields": {
                            "citizen_id": "1-8199-00905-15-7",
                            "student_no": "19984",
                            "weight_kg": "40",
                            "height_cm": "150",
                            "registered_address": {
                                "house_id": "8101-004408-9",
                                "postal_code": "81130",
                            },
                            "guardian": {
                                "phone": "063-839-5699",
                            },
                            "unknown_field": "ignored",
                        },
                        "needs_review": [],
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def _write_structured_civil_registration_json(path: Path) -> None:
    path.write_text(
        json.dumps(
            {
                "schema_version": "dmc_ocr_structured.v1",
                "document_type": "civil_registration",
                "records": [
                    {
                        "record_type": "civil_registration",
                        "page_start": 1,
                        "page_end": 1,
                        "fields": {
                            "citizen_id": "1-8199-00905-15-7",
                            "registered_address": {
                                "house_id": "8101-008408-9",
                                "house_no": "28",
                                "moo": "4",
                                "postal_code": "81130",
                            },
                            "current_address": {
                                "house_id": "8101-008408-9",
                                "house_no": "28",
                            },
                            "father": {
                                "citizen_id": "3-4101-00748-56-1",
                                "first_name": "Father",
                                "last_name": "Family",
                            },
                            "mother": {
                                "citizen_id": "1-7603-00002-25-6",
                                "first_name": "Mother",
                                "last_name": "Family",
                            },
                        },
                        "needs_review": [],
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def test_read_ocr_records_accepts_structured_json_and_ignores_civil_records(tmp_path: Path) -> None:
    ocr_path = tmp_path / "structured-ocr.json"
    _write_structured_ocr_json(ocr_path)
    payload = json.loads(ocr_path.read_text(encoding="utf-8"))
    payload["records"].append(
        {
            "record_type": "civil_registration",
            "fields": {
                "citizen_id": "1-8199-00905-15-7",
                "registered_address": {"house_id": "8101-008408-9"},
            },
        }
    )
    ocr_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    records, warnings = read_ocr_records(ocr_path)

    assert warnings == []
    assert len(records) == 1
    record = records[0]
    assert record.citizen_id == "1819900905157"
    assert record.student_no == "19984"
    assert record.row_index == 1
    assert record.fields["registered_address.house_id"].value == "8101-004408-9"
    assert record.fields["guardian.phone"].value == "063-839-5699"
    assert "unknown_field" not in record.fields


def test_read_structured_json_accepts_null_needs_review(tmp_path: Path) -> None:
    civil_path = tmp_path / "structured-civil.json"
    _write_structured_civil_registration_json(civil_path)
    payload = json.loads(civil_path.read_text(encoding="utf-8"))
    payload["records"][0]["needs_review"] = None
    civil_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    records, warnings = read_civil_registration_markdown_records(civil_path)

    assert warnings == []
    assert len(records) == 1
    assert records[0].citizen_id == "1819900905157"


def test_preview_dmc_form_json_accepts_structured_json_ocr(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    ocr_path = tmp_path / "structured-ocr.json"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    _write_structured_ocr_json(ocr_path)

    preview = preview_dmc_form_json(
        PreviewDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[],
            school_year=2569,
            grade_levels=[1],
        )
    )

    assert preview.summary.ocr_form_records == 1
    assert preview.summary.ocr_attached_records == 1
    record = preview.records[0]
    assert record.fields["citizen_id"] == "1819900905157"
    assert record.fields["weight_kg"] == "40"
    assert record.fields["guardian.phone"] == "063-839-5699"
    assert record.field_details["weight_kg"].source == "ocr_form"
    assert record.field_details["registered_address.house_id"].source == "ocr_form"
    assert record.dmc_form_values["weight"] == "40.0"


def test_civil_registration_structured_json_supplements_dmc_form_payload(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    ocr_path = tmp_path / "structured-ocr.json"
    civil_path = tmp_path / "structured-civil.json"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    _write_structured_ocr_json(ocr_path)
    _write_structured_civil_registration_json(civil_path)

    civil_records, civil_warnings = read_civil_registration_markdown_records(civil_path)
    preview = preview_dmc_form_json(
        PreviewDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[str(civil_path)],
            school_year=2569,
            grade_levels=[1],
        )
    )

    assert civil_warnings == []
    assert len(civil_records) == 1
    assert civil_records[0].fields["father.first_name"].value == "Father"
    record = preview.records[0]
    assert preview.summary.civil_registration_records == 1
    assert record.fields["current_address.house_id"] == "8101-008408-9"
    assert record.fields["father.first_name"] == "Father"
    assert record.field_details["current_address.house_id"].source == "civil_registration"
    assert record.field_details["father.first_name"].source == "civil_registration"


def test_read_ocr_markdown_accepts_typhoon_labels_without_asterisks(tmp_path: Path) -> None:
    ocr_path = tmp_path / "ocr-no-stars.md"
    ocr_path.write_text(
        "\n".join(
            [
                "หน้า 1",
                "เลขประจำตัวประชาชน 1928800047 068 เลขประจำตัวนักเรียน........................................ คำนำหน้านาม............................ เพศ ................ ชื่อ........................................................นามสกุล........................................ ชื่อ(อังกฤษ) Hussuna นามสกุล(อังกฤษ) Madta วันเดือนปีเกิด 20 มีนาคม 2556 จังหวัดที่เกิด ตรัง สถานพยาบาลที่เกิด รพ. วังวิเศษ กลุ่มเลือด — เชื้อชาติ ไทย สัญชาติ ไทย ศาสนา อิสลาม",
                "ตามทะเบียนบ้าน รหัสประจำบ้าน 8104 0015336 บ้านเลขที่ 231 หมู่ที่ (ถ้าไม่มีใส่ -) 1 ถนน (ถ้าไม่มีใส่ -) — จังหวัด กระบี่ อำเภอ คลองท่อม ตำบล ทรายขาว รหัสไปรษณีย์ 81170 หมายเลขโทรศัพท์บ้าน — มือถือ 065038 0104",
                "ที่อยู่ปัจจุบัน รหัสประจำบ้าน 81040015336 บ้านเลขที่ 231 หมู่ที่ (ถ้าไม่มีใส่ -) 1 ถนน (ถ้าไม่มีใส่ -) — จังหวัด กระบี่ อำเภอ คลองท่อม ตำบล ทรายขาว รหัสไปรษณีย์ 81170 หมายเลขโทรศัพท์บ้าน — มือถือ 0650380104",
                "๓. รายละเอียดนักเรียน การเดินทางมาโรงเรียน [ ] เดินเท้า [ ] พาหนะไม่เสียค่าโดยสาร [x] พาหนะเสียค่าโดยสาร อื่นๆ........................................ ระยะทางจากบ้านมา ร.ร. ทางน้ำ (กม.)............................ถนนลูกรัง(กม.)............................ถนนลาดยาง(กม.) 40 รวมระยะเวลาการเดินทางมาโรงเรียน (นาที) 30 นาที",
                "น้ำหนัก .....35..... กิโลกรัม ส่วนสูง ....150..... เซนติเมตร",
                "สถานภาพสมรส [ ] สมรส [ ] โสด [ ] หม้าย [x] หย่าร้าง [ ] อยู่ด้วยกัน [ ] แยกกันอยู่ ข้อมูลพี่น้อง จำนวนพี่ชาย .........—..........คน จำนวนน้องชาย .........1..........คน จำนวนพี่สาว .........—..........คน จำนวนน้องสาว .........1..........คน จำนวนพี่น้องที่ศึกษาอยู่ (ไม่รวมตัวนักเรียนเอง) ........1.......... คน นักเรียนเป็นบุตรคนที่ ........1..........",
                "ข้อมูลบิดา เลขบัตรประจำตัวประชาชนบิดา ...1909800344128... ชนิดบัตร [ ] บัตรประชาชน [ ] อื่นๆ ชื่อบิดา ......นายดาฟิก.......นามสกุล ..หมาดอาหา.......กลุ่มเลือดบิดา ..B... อาชีพ .......—.......... รายได้ต่อเดือน(บาท) .......—.......... หมายเลขโทรศัพท์บิดา .......—..........",
                "ข้อมูลมารดา เลขบัตรประจำตัวประชาชนมารดา ...1810400052734... ชนิดบัตร [x] บัตรประชาชน [ ] อื่นๆ ชื่อมารดา ....นางสาวรัตนา....นามสกุล ...ท้วยยอ.....กลุ่มเลือดมารดา ..O... อาชีพ ....รับจ้างทั่วไป.... รายได้ต่อเดือน(บาท) ..30,000.. หมายเลขโทรศัพท์มารดา ...0650380104...",
                "ข้อมูลผู้ปกครอง เลขบัตรประจำตัวประชาชนผู้ปกครอง...1810400052734... ชนิดบัตร [x] บัตรประชาชน [ ] อื่นๆ ชื่อผู้ปกครอง ...นางสาวรัตนา...นามสกุล ...ท้วยยอ....กลุ่มเลือดปกครอง ..O... อาชีพ ....รับจ้างทั่วไป.... รายได้ต่อเดือน(บาท) ..30,000.. หมายเลขโทรศัพท์ปกครอง ....0650380104....หมายเลขโทรศัพท์ผู้ปกครอง ..0650380104.. ความเกี่ยวข้องของผู้ปกครองกับนักเรียน ....มารดา..... หมายเหตุ",
            ]
        ),
        encoding="utf-8",
    )

    record, warnings = read_ocr_markdown(ocr_path)

    assert [warning.code for warning in warnings] == []
    assert record.citizen_id == "1928800047068"
    assert record.student_no is None
    assert record.fields["birth_date"].value == "20 มีนาคม 2556"
    assert record.fields["registered_address.house_no"].value == "231"
    assert "registered_address.road" not in record.fields
    assert record.fields["registered_address.subdistrict"].value == "ทรายขาว"
    assert record.fields["current_address.house_id"].value == "81040015336"
    assert record.fields["distance_paved_road_km"].value == "40"
    assert record.fields["commute_minutes"].value == "30"
    assert record.fields["younger_brothers"].value == "1"
    assert record.fields["younger_sisters"].value == "1"
    assert record.fields["father.first_name"].value == "นายดาฟิก"
    assert record.fields["father.last_name"].value == "หมาดอาหา"
    assert "father.phone" not in record.fields
    assert record.fields["mother.first_name"].value == "นางสาวรัตนา"
    assert record.fields["mother.phone"].value == "0650380104"
    assert record.fields["guardian.phone"].value == "0650380104"
    assert record.fields["guardian_relationship"].value == "มารดา"


def test_read_student_roster_accepts_header_based_student_list(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentlist-M1.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Worksheet"
    sheet.append(["ลำดับ", "รหัสนักเรียน", "ชื่อ - นามสกุล", "วันเกิด", "เลขบัตรประชาชน", "ชั้น", "ห้อง"])
    sheet.append([80, 20020, "เด็กหญิง อัซซูน่า หมาดตา", "20 มีนาคม 2556", 1928800047068, "ม.1", 3])
    workbook.save(roster_path)

    records, warnings = read_student_roster(roster_path, school_year=2569, grade_levels=[1])

    assert warnings == []
    assert len(records) == 1
    assert records[0].student_no == "20020"
    assert records[0].citizen_id == "1928800047068"
    assert records[0].birth_date == "20 มีนาคม 2556"
    assert records[0].grade == 1
    assert records[0].room == 3
    assert records[0].seat_no == 80
    assert records[0].prefix == "เด็กหญิง"
    assert records[0].first_name == "อัซซูน่า"
    assert records[0].last_name == "หมาดตา"


def test_dmc_form_json_uses_header_roster_citizen_id_to_attach_ocr(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentlist-M1.xlsx"
    ocr_path = tmp_path / "ocr.md"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Worksheet"
    sheet.append(["ลำดับ", "รหัสนักเรียน", "ชื่อ - นามสกุล", "วันเกิด", "เลขบัตรประชาชน", "ชั้น", "ห้อง"])
    sheet.append([80, 20020, "เด็กหญิง อัซซูน่า หมาดตา", "20 มีนาคม 2556", 1928800047068, "ม.1", 3])
    workbook.save(roster_path)
    ocr_path.write_text(
        "เลขประจำตัวประชาชน 1928800047 068 เลขประจำตัวนักเรียน........................................ "
        "คำนำหน้านาม............................ เพศ หญิง ชื่อ........................................................"
        "นามสกุล........................................ ชื่อ(อังกฤษ) Hussuna นามสกุล(อังกฤษ) Madta "
        "วันเดือนปีเกิด 20 มีนาคม 2556 จังหวัดที่เกิด ตรัง",
        encoding="utf-8",
    )

    result = preview_dmc_form_json(
        PreviewDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=None,
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[],
            school_year=2569,
            grade_levels=[1],
        )
    )

    assert result.summary.roster_records == 1
    assert result.summary.auto_matched == 1
    assert result.summary.ocr_attached_records == 1
    assert result.summary.ocr_unmatched_records == 0
    assert result.summary.review_queue_records == 0
    record = result.records[0]
    assert record.student_no == "20020"
    assert record.citizen_id == "1928800047068"
    assert record.grade == 1
    assert record.room == 3
    assert record.prefix == "เด็กหญิง"
    assert record.first_name == "อัซซูน่า"
    assert record.last_name == "หมาดตา"
    assert record.fields["birth_date"] == "20 มีนาคม 2556"
    assert record.dmc_form_values["studentNo"] == "20020"
    assert record.dmc_form_values["levelDtlCode"] == "10"
    assert record.dmc_form_values["classroom"] == "3"
    assert record.dmc_form_values["titleCode"] == "002"


def test_dmc_form_json_suggests_and_confirms_fuzzy_ocr_roster_match(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentlist-M1.xlsx"
    ocr_path = tmp_path / "structured-ocr.json"
    output_path = tmp_path / "dmc-form-data-2569.json"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Worksheet"
    sheet.append(["ลำดับ", "รหัสนักเรียน", "ชื่อ - นามสกุล", "วันเกิด", "เลขบัตรประชาชน", "ชั้น", "ห้อง"])
    sheet.append([6, 19988, "Mr Jirawat Wongwutikorn", "1 April 2557", 1819900965150, "M.1", 3])
    workbook.save(roster_path)
    ocr_path.write_text(
        json.dumps(
            {
                "schema_version": "dmc_ocr_structured.v1",
                "document_type": "dmc_form",
                "records": [
                    {
                        "record_type": "dmc_form",
                        "page_start": 1,
                        "page_end": 2,
                        "fields": {
                            "citizen_id": "1819900405150",
                            "prefix": "Mr",
                            "first_name": "Jirawat",
                            "last_name": "Wongwuthikorn",
                            "birth_date": "1 April 2557",
                            "weight_kg": "37.5",
                            "height_cm": "154.3",
                        },
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    preview = preview_dmc_form_json(
        PreviewDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=None,
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[],
            school_year=2569,
            grade_levels=[1],
        )
    )

    assert preview.summary.ocr_attached_records == 0
    assert preview.summary.ocr_unmatched_records == 1
    assert preview.summary.needs_review == 1
    unconfirmed = preview.records[0]
    assert unconfirmed.record_id == "ocr:structured-ocr.json:ocr-1"
    assert unconfirmed.student_no is None
    assert unconfirmed.citizen_id == "1819900405150"
    assert unconfirmed.suggestions[0].student_no == "19988"
    assert unconfirmed.suggestions[0].score >= 0.9

    skipped = export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=None,
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[],
            school_year=2569,
            grade_levels=[1],
            confirmed_matches=[],
            output_path=str(output_path),
        )
    )

    assert skipped.records_exported == 0
    assert skipped.records == []

    confirmed = export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=None,
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[],
            school_year=2569,
            grade_levels=[1],
            confirmed_matches=[
                DmcFormMatchConfirmation(record_id=unconfirmed.record_id, student_no="19988")
            ],
            output_path=str(output_path),
        )
    )

    assert confirmed.summary.ocr_attached_records == 1
    assert confirmed.summary.ocr_unmatched_records == 0
    record = confirmed.records[0]
    assert record.record_id == "roster:Worksheet:2:19988"
    assert record.student_no == "19988"
    assert record.citizen_id == "1819900965150"
    assert record.grade == 1
    assert record.room == 3
    assert record.dmc_form_values["studentNo"] == "19988"
    assert record.dmc_form_values["cifNo"] == "1819900965150"
    assert "ocr_conflicts_with_citizen_id" in record.review_reasons
    validation = validate_current_students_import_form(
        ValidateCurrentStudentsImportFormRequest(excel_path=str(output_path))
    )
    assert validation.summary.ready_rows == 1
    assert validation.summary.invalid_rows == 0
    import_records = load_dmc_transfer_in_import_records(output_path)
    assert import_records[0].student_no == "19988"
    assert import_records[0].citizen_id == "1819900965150"


def test_dmc_form_json_export_skips_all_unconfirmed_ocr_records_and_accepts_low_score_confirmation(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentlist-M1.xlsx"
    ocr_path = tmp_path / "structured-ocr.json"
    output_path = tmp_path / "dmc-form-data-2569.json"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Worksheet"
    sheet.append(["ลำดับ", "รหัสนักเรียน", "ชื่อ - นามสกุล", "วันเกิด", "เลขบัตรประชาชน", "ชั้น", "ห้อง"])
    sheet.append([1, 20001, "Miss Nicha Mamat", "1 May 2556", 1819900800380, "M.1", 3])
    sheet.append([2, 19988, "Mr Jirawat Wongwutikorn", "1 April 2557", 1819900965150, "M.1", 3])
    sheet.append([3, 20223, "Mr Songkran Panniam", "13 April 2556", 1819900905157, "M.1", 8])
    workbook.save(roster_path)
    ocr_path.write_text(
        json.dumps(
            {
                "schema_version": "dmc_ocr_structured.v1",
                "document_type": "dmc_form",
                "records": [
                    {
                        "record_type": "dmc_form",
                        "page_start": 1,
                        "page_end": 2,
                        "fields": {
                            "citizen_id": "1819900800380",
                            "prefix": "Miss",
                            "first_name": "Nicha",
                            "last_name": "Mamat",
                        },
                    },
                    {
                        "record_type": "dmc_form",
                        "page_start": 3,
                        "page_end": 4,
                        "fields": {
                            "citizen_id": "1819900405150",
                            "prefix": "Mr",
                            "first_name": "Jirawat",
                            "last_name": "Wongwuthikorn",
                        },
                    },
                    {
                        "record_type": "dmc_form",
                        "page_start": 5,
                        "page_end": 6,
                        "fields": {
                            "prefix": "Mr",
                            "first_name": "Songkran",
                            "last_name": "Munchi",
                        },
                    },
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    preview = preview_dmc_form_json(
        PreviewDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=None,
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[],
            school_year=2569,
            grade_levels=[1],
        )
    )

    assert preview.records_previewed == 3
    assert any("ocr_fuzzy_roster_match_candidate" in record.review_reasons for record in preview.records)
    low_score_record = next(record for record in preview.records if record.first_name == "Songkran")
    assert low_score_record.suggestions[0].student_no == "20223"
    assert low_score_record.suggestions[0].score < 0.9
    assert "ocr_fuzzy_roster_match_candidate" not in low_score_record.review_reasons

    exported = export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=None,
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[],
            school_year=2569,
            grade_levels=[1],
            confirmed_matches=[],
            output_path=str(output_path),
        )
    )

    assert exported.records_exported == 1
    assert exported.records[0].student_no == "20001"
    assert all("ocr_fuzzy_roster_match_candidate" not in record.review_reasons for record in exported.records)
    written = json.loads(output_path.read_text(encoding="utf-8"))
    assert [record["student_no"] for record in written["records"]] == ["20001"]

    confirmed_low_score = export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=None,
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[],
            school_year=2569,
            grade_levels=[1],
            confirmed_matches=[
                DmcFormMatchConfirmation(record_id=low_score_record.record_id, student_no="20223")
            ],
            output_path=str(output_path),
        )
    )

    assert confirmed_low_score.records_exported == 2
    assert {record.student_no for record in confirmed_low_score.records} == {"20001", "20223"}


def test_reconcile_current_students_builds_canonical_records_and_review_queue(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    ocr_path = tmp_path / "1-3ex.md"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    _write_ocr_markdown(ocr_path)

    result = reconcile_current_students(
        ReconcileCurrentStudentsRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            school_year=2569,
            grade_levels=[1],
        )
    )

    assert result.module == "currentStudents"
    assert result.summary.roster_records == 4
    assert result.summary.thai_id_scan_records == 6
    assert result.summary.ocr_form_records == 1
    assert result.summary.records_total == 5
    assert result.summary.auto_matched == 1
    assert result.summary.needs_review == 1
    assert result.summary.invalid_id_records == 1
    assert result.summary.duplicate_records == 1
    assert result.summary.duplicate_scan_records == 2
    assert result.summary.new_or_transfer_candidates == 1
    assert result.summary.ocr_attached_records == 1
    assert result.summary.review_queue_records == 4

    by_student_no = {record.student_no: record for record in result.records if record.student_no}
    matched = by_student_no["19984"]
    assert matched.match_status == "auto_matched"
    assert matched.citizen_id == "1819900905157"
    assert matched.dmc_fields["citizen_id"].source == "thai_id_scan"
    assert matched.dmc_fields["weight_kg"].source == "ocr_form"
    assert matched.dmc_fields["guardian.phone"].value == "063-839-5699"
    assert matched.dmc_fields["siblings_studying_count"].value == "1"
    assert matched.dmc_fields["child_order"].value == "2"
    assert any(conflict.field_name == "prefix" for conflict in matched.conflicts)

    assert by_student_no["19985"].match_status == "needs_review"
    assert by_student_no["19985"].citizen_id == "1819900915055"
    assert by_student_no["19985"].dmc_fields["citizen_id"].source == "thai_id_scan"
    assert "fuzzy_roster_match_candidate" in by_student_no["19985"].review_reasons
    assert "missing_thai_id_scan" not in by_student_no["19985"].review_reasons
    assert by_student_no["19986"].match_status == "invalid_id"
    assert by_student_no["19987"].match_status == "duplicate"

    fuzzy_record = next(record for record in result.records if record.citizen_id == "1819900915055")
    assert fuzzy_record.match_status == "needs_review"
    assert fuzzy_record.student_no == "19985"
    assert fuzzy_record.suggestions[0].student_no == "19985"
    assert fuzzy_record.suggestions[0].score >= 0.9

    new_candidate = next(record for record in result.records if record.citizen_id == "1819900943555")
    assert new_candidate.match_status == "new_or_transfer_candidate"
    assert "not_found_in_roster" in new_candidate.review_reasons


def test_validate_current_student_sources_rpc(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    server = RpcServer(emit_notification=lambda _payload: None)

    response = json.loads(
        server.handle_text(
            json.dumps(
                {
                    "jsonrpc": "2.0",
                    "id": "req-current-students",
                    "method": "validate_current_student_sources",
                    "params": {
                        "roster_excel_path": str(roster_path),
                        "thai_id_csv_path": str(thai_id_path),
                        "school_year": 2569,
                        "grade_levels": [1],
                    },
                }
            )
        )
    )

    assert response["result"]["module"] == "currentStudents"
    assert response["result"]["summary"]["roster_records"] == 4
    assert response["result"]["summary"]["thai_id_scan_records"] == 6
    assert response["result"]["summary"]["review_queue_records"] == 4


def test_export_blank_form_and_validate_completed_import_form(tmp_path: Path) -> None:
    output_path = tmp_path / "current-students-import-template.xlsx"

    blank_form = export_current_students_blank_form(
        ExportCurrentStudentsBlankFormRequest(output_path=str(output_path))
    )

    assert blank_form.module == "currentStudents"
    assert output_path.exists()
    assert "citizen_id" in blank_form.required_fields

    workbook = load_workbook(output_path)
    sheet = workbook.active
    field_columns = {str(cell.value): index for index, cell in enumerate(sheet[2], start=1) if cell.value}
    sheet.cell(row=3, column=field_columns["operation_type"], value="current")
    sheet.cell(row=3, column=field_columns["school_year"], value=2569)
    sheet.cell(row=3, column=field_columns["student_no"], value="19984")
    sheet.cell(row=3, column=field_columns["citizen_id"], value="1819900905157")
    sheet.cell(row=3, column=field_columns["prefix"], value="Miss")
    sheet.cell(row=3, column=field_columns["first_name"], value="Kanyanat")
    sheet.cell(row=3, column=field_columns["last_name"], value="Kuldee")
    workbook.save(output_path)

    validation = validate_current_students_import_form(
        ValidateCurrentStudentsImportFormRequest(excel_path=str(output_path))
    )

    assert validation.module == "currentStudents"
    assert validation.summary.rows_total == 1
    assert validation.summary.ready_rows == 1
    assert validation.summary.invalid_rows == 0
    assert validation.preview[0].student_no == "19984"


def test_export_current_student_import_excel_writes_workbook(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    ocr_path = tmp_path / "1-3ex.md"
    output_path = tmp_path / "current-students-import-2569.xlsx"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    _write_ocr_markdown(ocr_path)

    result = export_current_students_import_excel(
        ExportCurrentStudentsImportExcelRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            school_year=2569,
            grade_levels=[1],
            output_path=str(output_path),
        )
    )

    assert result.module == "formConverter"
    assert result.output_path == str(output_path)
    assert result.rows_exported == 1
    assert result.summary.records_total == 1
    conflict_fields = {conflict.field_name for conflict in result.conflicts}
    assert "prefix" not in conflict_fields
    assert "mother.first_name" in conflict_fields
    assert "mother.last_name" not in conflict_fields
    assert all(conflict.reason == "missing_after_all_sources" for conflict in result.conflicts)
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    sheet = workbook.active
    field_columns = {str(cell.value): index for index, cell in enumerate(sheet[2], start=1) if cell.value}
    data_rows = [row for row in sheet.iter_rows(min_row=3, values_only=True) if any(row)]
    assert len(data_rows) == 1
    citizen_values = [
        row[field_columns["citizen_id"] - 1]
        for row in data_rows
        if row[field_columns["citizen_id"] - 1]
    ]
    assert "1819900905157" in citizen_values
    assert data_rows[0][field_columns["mother.last_name"] - 1] == "ตุ้มดำ"


def test_dmc_form_values_apply_safe_defaults_for_missing_required_family_fields() -> None:
    record = CanonicalStudentRecord(
        record_id="record-1",
        operation_type="transfer_in",
        match_status="auto_matched",
        student_no="20001",
        citizen_id="1819900800380",
        grade=1,
        room=3,
        seat_no=1,
        prefix="เด็กหญิง",
        first_name="นิชา",
        last_name="มามาตย์",
        full_name="เด็กหญิง นิชา มามาตย์",
        dmc_fields={
            "father.first_name": _ocr_field("นายทดสอบ"),
            "father.last_name": _ocr_field("พ่อ"),
            "mother.first_name": _ocr_field("นางทดสอบ"),
            "mother.last_name": _ocr_field("แม่"),
            "guardian.first_name": _ocr_field("นายทดสอบ"),
            "guardian.last_name": _ocr_field("พ่อ"),
        },
    )

    values = _dmc_form_values(record, default_school_year=2569)

    assert values["marriageStatusCode"] == "01"
    assert values["numOfOlderBrothers"] == "0"
    assert values["numOfYoungerBrothers"] == "0"
    assert values["numOfOlderSisters"] == "0"
    assert values["numOfYoungerSisters"] == "0"
    assert values["numOfStudyingSiblings"] == "0"
    assert values["fatherOccupationCode"] == "5"
    assert values["motherOccupationCode"] == "5"
    assert values["parentOccupationCode"] == "5"
    assert values["journeyTypeCode"] == "02"
    assert values["timeDt"] == "10.0"
    assert values["raceCode"] == "099"


def test_dmc_form_values_map_common_family_status_and_occupations() -> None:
    record = CanonicalStudentRecord(
        record_id="record-1",
        operation_type="transfer_in",
        match_status="auto_matched",
        student_no="20002",
        citizen_id="1819900923473",
        grade=1,
        room=3,
        seat_no=2,
        prefix="เด็กหญิง",
        first_name="ปาริศา",
        last_name="ไกรนรา",
        full_name="เด็กหญิง ปาริศา ไกรนรา",
        dmc_fields={
            "parents_marital_status": _ocr_field("สมรส"),
            "older_brothers": _ocr_field(""),
            "father.first_name": _ocr_field("นายทดสอบ"),
            "father.last_name": _ocr_field("พ่อ"),
            "father.occupation": _ocr_field("ทำสวน"),
            "mother.first_name": _ocr_field("นางทดสอบ"),
            "mother.last_name": _ocr_field("แม่"),
            "mother.occupation": _ocr_field("พยาบาล"),
            "guardian.first_name": _ocr_field("นางทดสอบ"),
            "guardian.last_name": _ocr_field("แม่"),
            "guardian.occupation": _ocr_field("ช่างสัก"),
        },
    )

    values = _dmc_form_values(record, default_school_year=2569)

    assert values["marriageStatusCode"] == "01"
    assert values["numOfOlderBrothers"] == "0"
    assert values["fatherOccupationCode"] == "4"
    assert values["motherOccupationCode"] == "6"
    assert values["parentOccupationCode"] == "5"

    widow_record = record.model_copy(
        update={"dmc_fields": {**record.dmc_fields, "parents_marital_status": _ocr_field("หม้าย")}}
    )
    assert _dmc_form_values(widow_record, default_school_year=2569)["marriageStatusCode"] == "04"


def test_dmc_form_values_map_unmatched_nonblank_occupation_to_other() -> None:
    record = CanonicalStudentRecord(
        record_id="record-1",
        operation_type="transfer_in",
        match_status="auto_matched",
        student_no="20003",
        citizen_id="1907500206283",
        grade=1,
        room=3,
        seat_no=3,
        prefix="เด็กหญิง",
        first_name="พิมชนก",
        last_name="เห้งลิ่ม",
        full_name="เด็กหญิง พิมชนก เห้งลิ่ม",
        dmc_fields={
            "father.first_name": _ocr_field("นายทดสอบ"),
            "father.last_name": _ocr_field("พ่อ"),
            "father.occupation": _ocr_field("YouTuber"),
        },
    )

    values = _dmc_form_values(record, default_school_year=2569)

    assert values["fatherOccupationCode"] == "99"


def test_dmc_form_values_align_district_to_resolved_subdistrict_code() -> None:
    record = CanonicalStudentRecord(
        record_id="record-1",
        operation_type="transfer_in",
        match_status="auto_matched",
        student_no="20002",
        citizen_id="1819900923473",
        grade=1,
        room=3,
        seat_no=2,
        prefix="เด็กหญิง",
        first_name="ปาริศา",
        last_name="ไกรนรา",
        full_name="เด็กหญิง ปาริศา ไกรนรา",
        dmc_fields={
            "current_address.province": _ocr_field("กระบี่"),
            "current_address.district": _ocr_field("เหนือคลอง"),
            "current_address.subdistrict": _ocr_field("คลองหิน"),
        },
    )

    values = _dmc_form_values(record, default_school_year=2569)

    assert values["provinceCode"] == "81000000"
    assert values["amphurCode"] == "81050000"
    assert values["tumbolCode"] == "81050400"

    typo_record = record.model_copy(
        update={
            "dmc_fields": {
                "current_address.province": _ocr_field("กระบี่"),
                "current_address.district": _ocr_field("เหนือคลอง"),
                "current_address.subdistrict": _ocr_field("ปกาไส"),
            }
        }
    )
    typo_values = _dmc_form_values(typo_record, default_school_year=2569)
    assert typo_values["amphurCode"] == "81080000"
    assert typo_values["tumbolCode"] == "81080700"


def test_export_dmc_form_json_writes_operation_neutral_json(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    ocr_path = tmp_path / "1-3ex.md"
    output_path = tmp_path / "dmc-form-data-2569.json"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    _write_ocr_markdown(ocr_path)

    result = export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            school_year=2569,
            grade_levels=[1],
            admission_date="2026-05-16",
            output_path=str(output_path),
        )
    )

    assert result.module == "formConverter"
    assert result.output_path == str(output_path)
    assert result.records_exported == 1
    assert result.records[0].fields["citizen_id"] == "1819900905157"
    assert result.records[0].field_details["citizen_id"].source == "thai_id_scan"
    assert "operation_type" not in result.records[0].fields

    form_values = result.records[0].dmc_form_values
    assert form_values["educationYear"] == "2569"
    assert form_values["admissionDate"] == "16/05/2569"
    assert form_values["studentNo"] == "19984"
    assert form_values["levelDtlCode"] == "10"
    assert form_values["classroom"] == "1"
    assert form_values["cifNo"] == "1819900905157"
    assert form_values["cifNoChk"] == "1819900905157"
    assert form_values["cifType"] == "I"
    assert form_values["titleCode"] == "002"
    assert form_values["genderCode"] == "F"
    assert form_values["birthDate"] == "13/02/2557"
    assert form_values["birthProvinceCode"] == "81000000"
    assert form_values["psHomeIdNo"] == "81010044089"
    assert form_values["psProvinceCode"] == "81000000"
    assert form_values["psAmphurCode"] == "81080000"
    assert form_values["psTumbolCode"] == "81080300"
    assert form_values["fatherTitleCode"] == "003"
    assert form_values["fatherFirstNameTh"] == "ชวลิต"
    assert form_values["fatherLastNameTh"] == "ตุ้มดำ"
    assert form_values["fatherOccupationCode"] == "5"
    assert form_values["fatherSalary"] == "5500.0"
    assert form_values["parentFamilyRelationCode"] == "01"
    assert form_values["parentFirstNameTh"] == "ชวลิต"
    assert form_values["parentLastNameTh"] == "ตุ้มดำ"
    assert form_values["rubberDt"] == "15000.0"
    assert form_values["weight"] == "40.0"
    assert form_values["height"] == "150.0"

    payload = json.loads(output_path.read_text(encoding="utf-8"))
    assert payload["schema_version"] == "dmc_form_json.v1"
    assert payload["records_exported"] == 1
    assert "operation_type" not in payload
    assert "operation_type" not in payload["records"][0]
    assert "operation_type" not in payload["records"][0]["fields"]
    assert payload["records"][0]["dmc_form_values"]["cifNo"] == "1819900905157"
    assert payload["records"][0]["dmc_form_values"]["admissionDate"] == "16/05/2569"
    assert payload["records"][0]["dmc_form_values"]["parentFamilyRelationCode"] == "01"
    assert payload["records"][0]["fields"]["siblings_studying_count"] == "1"
    assert payload["records"][0]["fields"]["child_order"] == "2"


def test_dmc_form_json_uses_ocr_citizen_id_as_fallback_when_scan_is_absent(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    ocr_path = tmp_path / "1-3ex.md"
    _write_roster(roster_path)
    _write_ocr_markdown(ocr_path)

    result = export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=None,
            ocr_markdown_paths=[str(ocr_path)],
            school_year=2569,
            grade_levels=[1],
        )
    )

    assert result.records[0].fields["citizen_id"] == "1819900905157"
    assert result.records[0].field_details["citizen_id"].source == "ocr_form"


def test_dmc_form_json_maps_admin_codes_outside_krabi(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    ocr_path = tmp_path / "1-3ex.md"
    _write_roster(roster_path)
    _write_ocr_markdown(ocr_path)

    ocr_text = ocr_path.read_text(encoding="utf-8")
    ocr_text = ocr_text.replace("\u0e01\u0e23\u0e30\u0e1a\u0e35\u0e48", "\u0e40\u0e0a\u0e35\u0e22\u0e07\u0e43\u0e2b\u0e21\u0e48")
    ocr_text = ocr_text.replace(
        "\u0e40\u0e2b\u0e19\u0e37\u0e2d\u0e04\u0e25\u0e2d\u0e07",
        "\u0e40\u0e21\u0e37\u0e2d\u0e07\u0e40\u0e0a\u0e35\u0e22\u0e07\u0e43\u0e2b\u0e21\u0e48",
    )
    ocr_text = ocr_text.replace(
        "\u0e04\u0e25\u0e2d\u0e07\u0e02\u0e19\u0e32\u0e19",
        "\u0e0a\u0e49\u0e32\u0e07\u0e04\u0e25\u0e32\u0e19",
    )
    ocr_text = ocr_text.replace("81130", "50100")
    ocr_path.write_text(ocr_text, encoding="utf-8")

    result = export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=None,
            ocr_markdown_paths=[str(ocr_path)],
            school_year=2569,
            grade_levels=[1],
        )
    )

    form_values = result.records[0].dmc_form_values
    assert form_values["psProvinceCode"] == "50000000"
    assert form_values["psAmphurCode"] == "50010000"
    assert form_values["psTumbolCode"] == "50010500"
    assert form_values["provinceCode"] == "50000000"
    assert form_values["amphurCode"] == "50010000"
    assert form_values["tumbolCode"] == "50010500"


def test_dmc_form_json_derives_guardian_fields_from_parent_relationship(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    ocr_path = tmp_path / "1-3ex.md"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    _write_ocr_markdown(ocr_path)

    text = ocr_path.read_text(encoding="utf-8")
    assert "ชื่อผู้ปกครอง* นาย ชวลิต" in text
    assert "หมายเลขโทรศัพท์ผู้ปกครอง* 063-839-5699" in text
    ocr_path.write_text(
        text.replace("ชื่อผู้ปกครอง* นาย ชวลิต", "ชื่อผู้ปกครอง* นาย สมชาย").replace(
            "หมายเลขโทรศัพท์ผู้ปกครอง* 063-839-5699",
            "หมายเลขโทรศัพท์ผู้ปกครอง* 099-999-9999",
        ),
        encoding="utf-8",
    )

    result = preview_dmc_form_json(
        PreviewDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            school_year=2569,
            grade_levels=[1],
        )
    )

    record = result.records[0]
    assert record.fields["guardian_relationship"] == "บิดาและบุตร"
    assert record.fields["guardian.first_name"] == record.fields["father.first_name"]
    assert record.fields["guardian.last_name"] == record.fields["father.last_name"]
    assert record.fields["guardian.phone"] == record.fields["father.phone"]
    assert record.field_details["guardian.first_name"].source == "derived"
    assert record.field_details["guardian.first_name"].raw_value == "derived_from_father.first_name:ocr_form"
    assert record.dmc_form_values["parentFamilyRelationCode"] == "01"
    assert record.dmc_form_values["parentFirstNameTh"] == record.dmc_form_values["fatherFirstNameTh"]
    assert record.dmc_form_values["parentLastNameTh"] == record.dmc_form_values["fatherLastNameTh"]


def test_validate_current_students_accepts_dmc_form_json_for_transfer_in(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    ocr_path = tmp_path / "1-3ex.md"
    output_path = tmp_path / "dmc-form-data-2569.json"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    _write_ocr_markdown(ocr_path)

    export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            school_year=2569,
            grade_levels=[1],
            output_path=str(output_path),
        )
    )

    validation = validate_current_students_import_form(
        ValidateCurrentStudentsImportFormRequest(excel_path=str(output_path))
    )

    assert validation.module == "currentStudents"
    assert validation.excel_path == str(output_path)
    assert validation.summary.rows_total == 1
    assert validation.summary.ready_rows == 1
    assert validation.summary.invalid_rows == 0
    assert validation.preview[0].operation_type == "transfer_in"
    assert validation.preview[0].student_no == "19984"
    assert validation.preview[0].citizen_id == "1819900905157"
    import_records = load_dmc_transfer_in_import_records(output_path)
    assert import_records[0].student_no == "19984"
    assert import_records[0].level_dtl_code == "10"
    assert import_records[0].classroom == "1"
    assert import_records[0].dmc_form_values["firstNameTh"] == "กัญญาณัฐ"
    assert import_records[0].dmc_form_values["middleNameTh"] == ""
    assert import_records[0].dmc_form_values["parentFamilyRelationCode"] == "01"


def test_preview_dmc_form_json_reports_missing_required_data_read_only(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    ocr_path = tmp_path / "1-3ex.md"
    output_path = tmp_path / "dmc-form-data-2569.json"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    _write_ocr_markdown(ocr_path)

    preview = preview_dmc_form_json(
        PreviewDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            school_year=2569,
            grade_levels=[1],
        )
    )

    assert preview.records_previewed == 1
    assert preview.records[0].fields["citizen_id"] == "1819900905157"
    assert preview.records[0].fields["siblings_studying_count"] == "1"
    assert preview.records[0].fields["child_order"] == "2"
    preview_conflict_fields = {conflict.field_name for conflict in preview.conflicts}
    assert "mother.first_name" in preview_conflict_fields
    assert "mother.last_name" not in preview_conflict_fields
    assert preview.records[0].fields["mother.last_name"] == "ตุ้มดำ"
    assert preview.records[0].field_details["mother.last_name"].source == "derived"
    assert "operation_type" not in preview.records[0].fields

    result = export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            school_year=2569,
            grade_levels=[1],
            output_path=str(output_path),
        )
    )

    result_conflict_fields = {conflict.field_name for conflict in result.conflicts}
    assert "mother.first_name" in result_conflict_fields
    assert "mother.last_name" not in result_conflict_fields
    assert result.records[0].fields["mother.first_name"] is None
    assert result.records[0].fields["mother.last_name"] == "ตุ้มดำ"
    payload = json.loads(output_path.read_text(encoding="utf-8"))
    assert payload["records"][0]["fields"]["mother.last_name"] == "ตุ้มดำ"
    assert payload["records"][0]["field_details"]["mother.last_name"]["source"] == "derived"


def test_dmc_form_json_uses_civil_registration_ocr_as_supplement(tmp_path: Path) -> None:
    roster_path = tmp_path / "studentListM1-M4 2569.xlsx"
    thai_id_path = tmp_path / "ThaiID M1-2569.CSV"
    ocr_path = tmp_path / "1-3ex.md"
    civil_path = tmp_path / "CivilDoc.md"
    output_path = tmp_path / "dmc-form-data-2569.json"
    _write_roster(roster_path)
    _write_thai_id_csv(thai_id_path)
    _write_ocr_markdown(ocr_path)
    _write_civil_registration_markdown(civil_path)

    preview = preview_dmc_form_json(
        PreviewDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[str(civil_path)],
            school_year=2569,
            grade_levels=[1],
        )
    )

    record = preview.records[0]
    conflict_fields = {conflict.field_name for conflict in preview.conflicts}
    assert preview.summary.civil_registration_records == 1
    assert record.fields["registered_address.house_id"] == "8101-008408-9"
    assert record.fields["current_address.house_id"] == "8101-008408-9"
    assert record.fields["current_address.house_no"] == record.fields["registered_address.house_no"]
    assert record.fields["father.citizen_id"] == "3410100748561"
    assert record.fields["father.first_name"] == "ชวลิต"
    assert record.fields["mother.citizen_id"] == "1760300002256"
    assert record.fields["mother.first_name"] == "กรานณิภา"
    assert record.fields["mother.last_name"] == "ตุ้มดำ"
    assert record.field_details["registered_address.house_id"].source == "civil_registration"
    assert record.field_details["current_address.house_id"].source == "civil_registration"
    assert record.field_details["current_address.subdistrict"].source != "ocr_form"
    assert record.field_details["mother.first_name"].source == "civil_registration"
    assert "mother.first_name" not in conflict_fields
    assert "mother.last_name" not in conflict_fields

    result = export_dmc_form_json(
        ExportDmcFormJsonRequest(
            roster_excel_path=str(roster_path),
            thai_id_csv_path=str(thai_id_path),
            ocr_markdown_paths=[str(ocr_path)],
            civil_registration_markdown_paths=[str(civil_path)],
            school_year=2569,
            grade_levels=[1],
            output_path=str(output_path),
        )
    )

    assert result.records[0].fields["registered_address.house_id"] == "8101-008408-9"
    assert result.records[0].fields["current_address.house_id"] == "8101-008408-9"
    payload = json.loads(output_path.read_text(encoding="utf-8"))
    assert payload["records"][0]["field_details"]["current_address.house_id"]["source"] == "civil_registration"
    assert payload["records"][0]["field_details"]["mother.first_name"]["source"] == "civil_registration"


def test_civil_registration_markdown_splits_typhoon_pages_and_bounds_parent_fields(tmp_path: Path) -> None:
    civil_path = tmp_path / "multi-page-civil.md"
    civil_path.write_text(
        "\n".join(
            [
                "<!-- Page 1 confidence: n/a -->",
                "รายการเกี่ยวกับบ้าน",
                "เลขรหัสประจำบ้าน 8101-008408-9",
                "รายการที่อยู่ 23 หมู่ที่ 4 ตำบลคลองขนาน อำเภอเหนือคลอง จังหวัดกระบี่",
                "รายการบุคคลในบ้านของเลขรหัสประจำบ้าน 8101-008408-9",
                "ลำดับที่ 12",
                "เพศชาย",
                "ชื่อ ด.ช.อนุวัฒน์ ตุ้มคำ",
                "สัญชาติ ไทย",
                "เลขประจำตัวประชาชน 1-8108-00164-49-1",
                "สถานภาพ ผู้อาศัย",
                "เกิดเมื่อ 13 ก.พ. 2557",
                "มารดาผู้ให้กำเนิด ชื่อ กรานนิกา 1-7603-00002-25-6 สัญชาติ ไทย",
                "บิดาผู้ให้กำเนิด ชื่อ ชวลิต 3-8101-00748-86-1 สัญชาติ ไทย",
                "<!-- Page 2 confidence: n/a -->",
                "รายการเกี่ยวกับบ้าน",
                "เลขรหัสประจำบ้าน 8104-001533-6",
                "รายการที่อยู่ 23/1 หมู่ที่ 1 ตำบลกรายขาว อำเภอคลองก่อม จังหวัดกระบี่",
                "รายการบุคคลในบ้านของเลขรหัสประจำบ้าน 8104-001533-6 ลำดับที่ 14",
                "ชื่อ ค.ญ.ธัญชน่า หมาดตา",
                "สัญชาติ ไทย",
                "เพศ หญิง",
                "เลขประจำตัวประชาชน 1-9288-00047-06-8",
                "สถานภาพ ผู้อาศัย",
                "เกิดเมื่อ 20 มี.ค. 2556",
                "มารดาผู้ให้กำเนิด ชื่อ รัตนา 1-8104-00052-13-4 สัญชาติ ไทย",
                "บิดาผู้ให้กำเนิด ชื่อ อาฟิก 1-9098-00344-12-8 สัญชาติ ไทย",
                "<!-- Page 3 confidence: n/a -->",
                "รายการเกี่ยวกับบ้าน",
                "เลขรหัสประจำบ้าน: 8108-002888-9",
                "รายการที่อยู่: 138 หมู่ที่ 5 ตำบลห้วยยูง อำเภอเหนือคลอง จังหวัดกระบี่",
                "รายการบุคคลในบ้านของเลขรหัสประจำบ้าน 8108-002888-9",
                "ลำดับที่ 7",
                "ชื่อ: ด.ญ.พิรุฬกานต์ เพชรสูก",
                "สัญชาติ: ไทย",
                "เพศ: หญิง",
                "เลขประจำตัวประชาชน: 1-8199-00924-40-2",
                "สถานภาพผู้อาศัย: เกิดเมื่อ 19 มิ.ย. 2556",
                "มารดาผู้ให้กำเนิด: ชื่อ วันเพ็ญ 1-6199-00222-35-3 สัญชาติ ไทย",
                "บิดาผู้ให้กำเนิด: ชื่อ นพเดช 1-8102-00065-94-8 สัญชาติ ไทย",
            ]
        ),
        encoding="utf-8",
    )

    records, warnings = read_civil_registration_markdown_records(civil_path)

    assert [warning.code for warning in warnings] == []
    assert len(records) == 3
    assert records[0].fields["father.first_name"].value == "ชวลิต"
    assert "father.last_name" not in records[0].fields
    assert records[0].fields["mother.first_name"].value == "กรานนิกา"
    assert "Page 2" not in str(records[0].fields.get("mother.first_name"))
    assert records[1].fields["father.first_name"].value == "อาฟิก"
    assert records[1].fields["mother.first_name"].value == "รัตนา"
    assert records[2].citizen_id == "1819900924402"
    assert records[2].first_name == "พิรุฬกานต์"
    assert records[2].fields["father.first_name"].value == "นพเดช"
    assert records[2].fields["mother.first_name"].value == "วันเพ็ญ"


UPLOAD_TEST_DIR = Path(__file__).parents[3] / "uploadTest"


@pytest.mark.skipif(
    not all(
        path.exists()
        for path in (
            UPLOAD_TEST_DIR / "ThaiID M1-2569.CSV",
            UPLOAD_TEST_DIR / "studentListM1-M4 2569.xlsx",
            UPLOAD_TEST_DIR / "1-3ex.md",
        )
    ),
    reason="uploadTest sample files are local fixtures",
)
def test_reconcile_current_students_with_upload_test_samples() -> None:
    result = reconcile_current_students(
        ReconcileCurrentStudentsRequest(
            roster_excel_path=str(UPLOAD_TEST_DIR / "studentListM1-M4 2569.xlsx"),
            thai_id_csv_path=str(UPLOAD_TEST_DIR / "ThaiID M1-2569.CSV"),
            ocr_markdown_paths=[str(UPLOAD_TEST_DIR / "1-3ex.md")],
            school_year=2569,
            grade_levels=[1],
        )
    )

    assert result.summary.roster_records == 328
    assert result.summary.thai_id_scan_records == 263
    assert result.summary.auto_matched >= 245
    assert result.summary.review_queue_records > 0
    assert result.summary.invalid_id_records >= 1
    assert result.summary.duplicate_scan_records >= 2
